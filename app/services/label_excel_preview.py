"""Read a saved Excel layout as one physical 90x50 label per material, without writes."""
import base64
from copy import copy
from datetime import date
import io
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from print_fill import _Filler, _match_image_cell, _render_placeholder_image, _resolve_path, validate_template_file
from services.material_label import material_label_qr
from services.report_print_preview import _cell_css, _display_value, _number, _print_bounds


def _png_data(image):
    from PIL import Image
    with Image.open(io.BytesIO(image._data())) as pixels:
        output = io.BytesIO()
        pixels.save(output, format='PNG')
        return 'data:image/png;base64,' + base64.b64encode(output.getvalue()).decode('ascii')


def _sheet_layout(sheet):
    if sheet._charts:
        raise ValueError('标签模板含图表，请在模板设计中改为文字或图片后打印')
    bounds = _print_bounds(sheet)
    if len(bounds) > 1:
        raise ValueError('标签模板有多个打印区域，请在模板设计中保留一个标签区域')
    left, top, right, bottom = bounds[0] if bounds else (1, 1, sheet.max_column, sheet.max_row)
    if (right - left + 1) * (bottom - top + 1) > 20000:
        raise ValueError('标签模板打印区域过大，请精简模板')
    widths = {col: (0 if sheet.column_dimensions[get_column_letter(col)].hidden else
                   _number(sheet.column_dimensions[get_column_letter(col)].width, 13, 1, 80) * 7)
              for col in range(left, right + 1)}
    heights = {row: (0 if sheet.row_dimensions[row].hidden else
                    _number(sheet.row_dimensions[row].height, 15, 1, 400) * 96 / 72)
               for row in range(top, bottom + 1)}
    width, height = sum(widths.values()), sum(heights.values())
    if not width or not height:
        raise ValueError('标签模板打印区域没有可见单元格')
    xs, ys = {left: 0}, {top: 0}
    for col in widths:
        xs[col + 1] = xs[col] + widths[col]
    for row in heights:
        ys[row + 1] = ys[row] + heights[row]
    merges, skipped = {}, set()
    for merged in sheet.merged_cells.ranges:
        if merged.max_col < left or merged.min_col > right or merged.max_row < top or merged.min_row > bottom:
            continue
        if merged.min_col < left or merged.min_row < top or merged.max_col > right or merged.max_row > bottom:
            raise ValueError('标签打印区域截断了合并单元格，请调整打印区域')
        merges[merged.min_row, merged.min_col] = (merged.max_row, merged.max_col)
        skipped.update((row, col) for row in range(merged.min_row, merged.max_row + 1)
                       for col in range(merged.min_col, merged.max_col + 1)
                       if (row, col) != (merged.min_row, merged.min_col))
    cells = []
    for row in heights:
        for col in widths:
            if not heights[row] or not widths[col] or (row, col) in skipped:
                continue
            cell = sheet.cell(row, col)
            if cell.data_type == 'f':
                raise ValueError('标签模板含 Excel 公式，请在模板设计中改为确定的值后打印')
            end_row, end_col = merges.get((row, col), (row, col))
            cells.append((cell, dict(left=xs[col], top=ys[row],
                                    width=xs[end_col + 1] - xs[col], height=ys[end_row + 1] - ys[row])))
    images = []
    if len(sheet._images) > 40:
        raise ValueError('标签模板图片过多，请精简模板')
    for image in sheet._images:
        anchor = image.anchor
        marker = getattr(anchor, '_from', None)
        if marker is None:
            raise ValueError('标签图片定位无法识别，请在模板设计中重新放置图片')
        row, col = marker.row + 1, marker.col + 1
        if row not in ys or col not in xs or row > bottom or col > right:
            continue
        ext = getattr(anchor, 'ext', None)
        end = getattr(anchor, 'to', None)
        image_width, image_height = image.width, image.height
        if ext is not None:
            image_width, image_height = ext.cx / 9525, ext.cy / 9525
        elif end is not None:
            if end.col + 1 not in xs or end.row + 1 not in ys:
                raise ValueError('标签图片超出打印区域，请调整图片位置')
            image_width = xs[end.col + 1] - xs[col] + (end.colOff - marker.colOff) / 9525
            image_height = ys[end.row + 1] - ys[row] + (end.rowOff - marker.rowOff) / 9525
        images.append(dict(left=xs[col] + marker.colOff / 9525, top=ys[row] + marker.rowOff / 9525,
                           width=image_width, height=image_height, src=_png_data(image)))
    return dict(width=90, height=50, content_width=width, content_height=height,
                scale=min(86 * 96 / 25.4 / width, 46 * 96 / 25.4 / height), cells=cells, images=images)


def build_excel_label_preview(template_path, materials, qr_base_url=None):
    """Preserve merged multiline layouts; fill each label independently with shared tokens."""
    if len(materials) > 500:
        raise ValueError('单次最多预览500个标签，请分批打印')
    path = Path(template_path)
    if not path.is_file():
        raise ValueError('所选 Excel 模板文件不存在，请在模板管理中修复或重新选择')
    if path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError('标签模板文件过大，请精简模板')
    raw = path.read_bytes()
    error = validate_template_file(raw)
    if error:
        raise ValueError(error)
    book = load_workbook(io.BytesIO(raw), keep_links=False)
    try:
        visible = [sheet for sheet in book.worksheets if sheet.sheet_state == 'visible']
        if len(visible) != 1:
            raise ValueError('物料标签模板需保留一个可见工作表，请在模板设计中调整')
        layout = _sheet_layout(visible[0])
        if len(layout['cells']) * len(materials) > 200000:
            raise ValueError('本批标签预览内容过大，请减少物料数量后打印')
        labels = []
        order = SimpleNamespace(title='物料标签', operator=SimpleNamespace(username=''), total_amount=0)
        for material in materials:
            item = SimpleNamespace(**material)
            filler = _Filler(order, [item], date.today().isoformat())
            cells = []
            for source, geometry in layout['cells']:
                cell = copy(source)
                image = _match_image_cell(cell.value)
                result = dict(**geometry, css=_cell_css(cell), text='')
                if image:
                    kind, scope, attr = image
                    value = _resolve_path(item if scope == 'item' else order, attr.split('.'))
                    if kind == 'qrcode' and scope == 'item' and attr in ('code', 'barcode', 'qrcode'):
                        qr = material_label_qr(material['code'], qr_base_url)
                        result['image'] = 'data:image/png;base64,' + qr['qrcode_png']
                    else:
                        pixels = _render_placeholder_image(kind, value)
                        if pixels is None:
                            raise ValueError('标签条码或二维码生成失败，请检查模板字段与物料编码')
                        output = io.BytesIO()
                        pixels.save(output, format='PNG')
                        result['image'] = 'data:image/png;base64,' + base64.b64encode(output.getvalue()).decode('ascii')
                        pixels.close()
                else:
                    cell.value = filler._replace_cell(cell.value, lambda token: filler.item_value(token, item))
                    result['text'] = _display_value(cell)
                cells.append(result)
            labels.append({**layout, 'cells': cells})
        return labels
    finally:
        book.close()
