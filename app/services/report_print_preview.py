"""Read-only Excel report template rendering for browser preview and printing."""
from __future__ import annotations

import io
import math
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries

from doc_print_excel import generate_report_builtin_template
from print_fill import _Filler, _color_to_hex, _has_item_placeholder, validate_template_file

MAX_PREVIEW_ROWS = 5000
MAX_PREVIEW_CELLS = 150000
_MONEY = re.compile(r'price|amount|value|cost', re.I)


def visible_report_data(columns, rows, include_prices=True):
    """Hide price dimensions in both default and user-supplied template inputs."""
    if include_prices:
        return columns, rows
    columns = [column for column in columns if column.get('type') != 'money'
               and not _MONEY.search(column.get('field', ''))]
    rows = [{key: value for key, value in row.items() if not _MONEY.search(key)}
            for row in rows]
    return columns, rows


def _number(value, default=0, minimum=0, maximum=1000):
    try:
        result = float(value)
        return min(maximum, max(minimum, result)) if math.isfinite(result) else default
    except (ValueError, TypeError):
        return default


def _cell_css(cell):
    font, alignment = cell.font, cell.alignment
    # Font names are data from uploaded workbooks; allow names, never CSS syntax.
    font_name = re.sub(r'[^\w \-]', '', str(font.name or 'Arial'))[:80]
    parts = ['font-family:"%s",Arial,"Microsoft YaHei",sans-serif' % font_name,
             'font-size:%gpt' % _number(font.sz, 11, 6, 72),
             'font-weight:' + ('bold' if font.bold else 'normal'),
             'font-style:' + ('italic' if font.italic else 'normal')]
    if font.underline:
        parts.append('text-decoration:underline')
    for prop, color in (('color', font.color),
                        ('background-color', cell.fill.fgColor if cell.fill.patternType == 'solid' else None)):
        value = _color_to_hex(color)
        if value:
            parts.append('%s:%s' % (prop, value))
    horizontal = alignment.horizontal
    if horizontal not in ('left', 'center', 'right', 'justify'):
        horizontal = 'right' if isinstance(cell.value, (int, float)) else 'left'
    parts.append('text-align:' + horizontal)
    parts.append('vertical-align:' + {'center': 'middle', 'top': 'top', 'bottom': 'bottom'}.get(alignment.vertical, 'middle'))
    parts.append('white-space:' + ('pre-wrap' if alignment.wrap_text else 'pre'))
    for edge in ('top', 'bottom', 'left', 'right'):
        side = getattr(cell.border, edge)
        if side.style:
            style = 'dashed' if 'dash' in side.style.lower() else ('dotted' if side.style == 'dotted' else 'solid')
            width = 2 if side.style in ('medium', 'thick', 'double') else 1
            parts.append('border-%s:%dpx %s %s' % (edge, width, style, _color_to_hex(side.color) or '#333333'))
    return ';'.join(parts)


def _display_value(cell):
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return '是' if value else '否'
    if not isinstance(value, (int, float)):
        return str(value)
    if not math.isfinite(value):
        raise ValueError('模板数据包含无效数值，无法打印')
    fmt = cell.number_format or 'General'
    if fmt in ('General', '@'):
        return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    if is_date_format(fmt):
        raise ValueError('模板日期格式与数值数据不匹配，请修正模板')
    section = fmt.split(';')[1 if value < 0 and ';' in fmt else 0]
    section = re.sub(r'\[[^]]*\]', '', section)
    decimals = re.search(r'\.([0#]+)', section)
    if not re.search(r'[0#]', section) or re.search(r'[Ee][+-]|[?]|/', section):
        raise ValueError('模板包含浏览器暂不支持的数字格式，请使用 Excel 模板文件打印')
    number = value * (100 if '%' in section else 1)
    digits = len(decimals.group(1)) if decimals else 0
    rendered = format(abs(number) if '(' in section and value < 0 else number,
                      (',' if ',' in section else '') + '.' + str(digits) + 'f')
    if value < 0 and '(' in section:
        rendered = '(' + rendered + ')'
    if '%' in section:
        rendered += '%'
    if '¥' in section or '￥' in section:
        rendered = '¥' + rendered
    elif '$' in section:
        rendered = '$' + rendered
    return rendered


def _print_bounds(sheet):
    # openpyxl 3.0 returns a list, 3.1 returns a sheet-qualified string.
    area = sheet.print_area
    if not area:
        return []
    refs = re.findall(r'\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+', str(area))
    if not refs:
        raise ValueError('模板打印区域无法识别，请检查模板')
    return [range_boundaries(ref) for ref in refs]


def _fill_sheet(sheet, filler):
    """Retain merges and styles when the shared filler expands detail rows."""
    detail = next((row[0].row for row in sheet.iter_rows()
                   if any(_has_item_placeholder(cell.value) for cell in row)), None)
    merges = list(sheet.merged_cells.ranges)
    new_merges = []
    if detail is not None:
        boundary = filler._find_boundary(sheet, detail, sheet.max_row, sheet.max_column)
        capacity = boundary - detail if boundary else 1
        count = max(1, len(filler.items))
        delta = count - capacity
        source = [(cell.value, copy(cell._style)) for cell in sheet[detail]]
        for merged in merges:
            if merged.min_row < detail <= merged.max_row or (merged.min_row == detail and merged.max_row > detail):
                raise ValueError('明细行跨行合并暂不支持浏览器预览，请使用 Excel 模板文件打印')
            if merged.min_row == merged.max_row == detail:
                new_merges.extend((merged.min_col, detail + offset, merged.max_col, detail + offset)
                                  for offset in range(count))
            else:
                shift = delta if merged.min_row > detail else 0
                new_merges.append((merged.min_col, merged.min_row + shift,
                                   merged.max_col, merged.max_row + shift))
        for merged in merges:
            sheet.unmerge_cells(str(merged))
        # openpyxl does not update print areas on insert_rows; expand them explicitly.
        original_areas = _print_bounds(sheet)
        filler.fill(sheet)
        if not filler.items and capacity > 1:
            sheet.delete_rows(detail + 1, capacity - 1)
        for offset in range(count):
            for index, (value, style) in enumerate(source, 1):
                cell = sheet.cell(detail + offset, index)
                cell._style = copy(style)
                if not _has_item_placeholder(value):
                    cell.value = filler._replace_cell(value, filler.order_value)
        for left, top, right, bottom in new_merges:
            sheet.merge_cells(start_row=top, start_column=left, end_row=bottom, end_column=right)
        if original_areas:
            areas = []
            for left, top, right, bottom in original_areas:
                if top > detail:
                    top += delta
                if bottom >= detail:
                    bottom += delta
                areas.append('%s%d:%s%d' % (get_column_letter(left), top, get_column_letter(right), bottom))
            sheet.print_area = areas
    else:
        filler.fill(sheet)


def build_report_print_preview(title, columns, rows, template_path=None, include_prices=True):
    """Fill an existing workbook entirely in memory; return escaped-template render data."""
    if len(rows) > MAX_PREVIEW_ROWS:
        raise ValueError('当前结果超过浏览器打印上限（5000 行），请缩小筛选范围后打印')
    columns, rows = visible_report_data(columns, rows, include_prices)
    if template_path:
        path = Path(template_path)
        if not path.is_file():
            raise ValueError('所选打印模板文件不存在，请重新选择或修复模板')
        if path.stat().st_size > 10 * 1024 * 1024:
            raise ValueError('打印模板文件过大，请精简模板')
        raw = path.read_bytes()
    else:
        content = generate_report_builtin_template(title, columns)
        if content is None:
            raise ValueError('报表没有可打印的字段')
        raw = content.getvalue()
    error = validate_template_file(raw)
    if error:
        raise ValueError(error)
    workbook = load_workbook(io.BytesIO(raw), keep_links=False)
    try:
        if len(workbook.worksheets) > 10:
            raise ValueError('模板工作表过多，请精简模板')
        order = SimpleNamespace(title=title, operator=SimpleNamespace(username=''),
            total_amount=sum(float(row.get('amount') or 0) for row in rows) if include_prices else '')
        items = [SimpleNamespace(**{'quantity': row.get('quantity', 0), **row}) for row in rows]
        sheets = []
        for index, sheet in enumerate(workbook.worksheets):
            if sheet.sheet_state != 'visible':
                continue
            if sheet.max_row * sheet.max_column > MAX_PREVIEW_CELLS:
                raise ValueError('模板工作表规模过大，请精简模板')
            if (sheet.max_row + max(0, len(items) - 1)) * sheet.max_column > MAX_PREVIEW_CELLS:
                raise ValueError('展开后的打印内容过大，请缩小筛选范围后打印')
            if sheet._images or sheet._charts:
                raise ValueError('此模板含图片或图表，浏览器暂不能完整呈现，请使用 Excel 模板文件打印')
            for cells in sheet.iter_rows():
                for cell in cells:
                    if cell.data_type == 'f':
                        raise ValueError('此模板含 Excel 公式，浏览器无法重新计算；请下载 Excel 模板文件后打印')
                    if isinstance(cell.value, str) and re.search(r'\{img_(?:barcode|qrcode):', cell.value):
                        raise ValueError('此模板含条码或二维码图片，浏览器暂不能完整呈现，请使用 Excel 模板文件打印')
                    if not include_prices and isinstance(cell.value, str):
                        cell.value = re.sub(r'\{(?:item|order)\.[^{}]*(?:price|amount|value|cost)[^{}]*\}|\{total_amount\}', '', cell.value, flags=re.I)
            _fill_sheet(sheet, _Filler(order, items, date.today().isoformat()))
            if sheet.max_row * sheet.max_column > MAX_PREVIEW_CELLS:
                raise ValueError('打印内容过大，请缩小筛选范围后打印')
            # Print area is part of the selected template, never silently print excluded cells.
            bounds = (1, 1, sheet.max_column, sheet.max_row)
            if sheet.print_area:
                ranges = _print_bounds(sheet)
                if len(ranges) != 1:
                    raise ValueError('不连续打印区域暂不支持浏览器预览，请使用 Excel 模板文件打印')
                bounds = ranges[0]
            left, top, right, bottom = bounds
            if (right - left + 1) * (bottom - top + 1) > MAX_PREVIEW_CELLS:
                raise ValueError('模板打印区域过大，请精简打印区域')
            anchors, skipped = {}, set()
            for merged in sheet.merged_cells.ranges:
                anchors[merged.min_row, merged.min_col] = (merged.max_row - merged.min_row + 1, merged.max_col - merged.min_col + 1)
                skipped.update((r, c) for r in range(merged.min_row, merged.max_row + 1)
                               for c in range(merged.min_col, merged.max_col + 1)
                               if (r, c) != (merged.min_row, merged.min_col))
            rendered_rows = []
            for row in range(top, bottom + 1):
                if sheet.row_dimensions[row].hidden:
                    continue
                cells = []
                for col in range(left, right + 1):
                    if (row, col) in skipped or sheet.column_dimensions[get_column_letter(col)].hidden:
                        continue
                    cell = sheet.cell(row, col)
                    rowspan, colspan = anchors.get((row, col), (1, 1))
                    cells.append(dict(text=_display_value(cell), css=_cell_css(cell), rowspan=rowspan, colspan=colspan))
                rendered_rows.append(dict(cells=cells, height=_number(sheet.row_dimensions[row].height, 16, 8, 400)))
            margins = sheet.page_margins
            papers = {'1': 'Letter', '5': 'Legal', '8': 'A3', '9': 'A4', '11': 'A5'}
            paper = papers.get(str(sheet.page_setup.paperSize or '9'))
            if paper is None:
                raise ValueError('此模板纸张规格暂不支持浏览器打印，请使用 Excel 模板文件打印')
            page_width, page_height = {'Letter': (215.9, 279.4), 'Legal': (215.9, 355.6),
                                      'A3': (297, 420), 'A4': (210, 297), 'A5': (148, 210)}[paper]
            if sheet.page_setup.orientation == 'landscape':
                page_width, page_height = page_height, page_width
            sheets.append(dict(name=sheet.title, page_name='reportSheet%d' % index,
                page_width=page_width, page_height=page_height,
                paper=paper, orientation='landscape' if sheet.page_setup.orientation == 'landscape' else 'portrait',
                margins={key: '%gmm' % (_number(getattr(margins, key), .5, 0, 3) * 25.4)
                         for key in ('top', 'bottom', 'left', 'right')},
                widths=[_number(sheet.column_dimensions[get_column_letter(col)].width, 13, 1, 80) * 7
                        for col in range(left, right + 1) if not sheet.column_dimensions[get_column_letter(col)].hidden],
                rows=rendered_rows))
        if not sheets:
            raise ValueError('模板没有可见工作表')
        return sheets
    finally:
        workbook.close()
