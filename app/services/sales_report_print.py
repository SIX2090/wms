"""Shared sales-report page collectors and read-only template outputs."""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from types import SimpleNamespace

from flask import current_app, flash, jsonify, redirect, request, url_for

from db import db


def sales_reconciliation_context():
    from sqlalchemy import func
    from sqlalchemy.orm import selectinload
    from app import SalesOrder, SalesOrderItem, StockTransaction, Warehouse, _require_report_warehouse, STOCK_COMPARE_EPSILON, round_to_2_decimals
    warehouse_id = request.args.get('warehouse_id', type=int)
    selected_warehouse, _wh_err = _require_report_warehouse()
    if selected_warehouse:
        warehouse_id = selected_warehouse.id
    if not selected_warehouse:
        orders = []
    else:
        query = SalesOrder.query.filter(SalesOrder.status != 'cancelled')
        query = query.filter(db.or_(SalesOrder.warehouse_id == selected_warehouse.id, db.and_(SalesOrder.warehouse_id.is_(None), SalesOrder.warehouse == selected_warehouse.name)))
        orders = query.options(selectinload(SalesOrder.items), selectinload(SalesOrder.outbound_orders)).order_by(SalesOrder.date.desc(), SalesOrder.id.desc()).all()
    rows = []
    for order in orders:
        completed_orders = [outbound for outbound in order.outbound_orders if outbound.status == 'completed']
        completed_qty = 0
        completed_amount = 0
        inventory_qty = 0
        for outbound in completed_orders:
            for item in outbound.items:
                if item.source_sales_order_item_id:
                    source = db.session.get(SalesOrderItem, item.source_sales_order_item_id)
                    if not source or source.sales_order_id != order.id:
                        continue
                completed_qty += float(item.quantity or 0)
                completed_amount += float(item.amount or 0)
                inventory_qty += abs(float(db.session.query(func.coalesce(func.sum(StockTransaction.quantity), 0)).filter(StockTransaction.reference_type == 'out_order', StockTransaction.reference_id == outbound.id, StockTransaction.material_id == item.material_id, StockTransaction.transaction_type == 'out').scalar() or 0))
        shipped_qty = sum((float(item.shipped_quantity or 0) for item in order.items))
        shipped_amount = float(order.shipped_amount or 0)
        rows.append({'order': order, 'shipped_qty': round_to_2_decimals(shipped_qty), 'completed_qty': round_to_2_decimals(completed_qty), 'inventory_qty': round_to_2_decimals(inventory_qty), 'quantity_diff': round_to_2_decimals(shipped_qty - completed_qty), 'amount_diff': round_to_2_decimals(shipped_amount - completed_amount), 'ok': abs(shipped_qty - completed_qty) <= STOCK_COMPARE_EPSILON and abs(shipped_amount - completed_amount) <= 0.01})
    total_orders = len(rows)
    passed_orders = sum((1 for r in rows if r['ok']))
    failed_orders = total_orders - passed_orders
    total_shipped = sum((r['shipped_qty'] for r in rows))
    total_completed = sum((r['completed_qty'] for r in rows))
    total_diff = sum((abs(r['quantity_diff']) for r in rows))
    chart_data = {'summary': {'total': total_orders, 'passed': passed_orders, 'failed': failed_orders, 'passed_pct': round(passed_orders / total_orders * 100, 1) if total_orders > 0 else 0, 'total_shipped': round(total_shipped, 2), 'total_completed': round(total_completed, 2), 'total_diff': round(total_diff, 2)}, 'warehouse_stats': {}, 'monthly_trend': []}
    return dict(rows=rows, warehouses=Warehouse.query.filter_by(status='active').order_by(Warehouse.code.asc()).all(), warehouse_id=warehouse_id or '', chart_data=chart_data)


def sales_summary_context():
    from datetime import date, datetime
    from types import SimpleNamespace
    from sqlalchemy.orm import joinedload, selectinload
    from app import Customer, Employee, Material, SalesOrder, SalesOrderItem, Warehouse, _require_report_warehouse, round_to_2_decimals, sales_status_label
    date_start = request.args.get('date_start') or date.today().replace(day=1).isoformat()
    date_end = request.args.get('date_end') or date.today().isoformat()
    drill_customer_id = request.args.get('customer_id', type=int)
    drill_material_code = (request.args.get('material_code') or '').strip()
    salesperson_id = request.args.get('salesperson_id', type=int)
    status = (request.args.get('status') or '').strip()
    shipment_status = (request.args.get('shipment_status') or '').strip()
    project_no = (request.args.get('project_no') or '').strip()
    warehouse_id = request.args.get('warehouse_id', type=int)
    warehouse = (request.args.get('warehouse') or '').strip()
    selected_warehouse, _wh_err = _require_report_warehouse()
    if selected_warehouse:
        warehouse_id = selected_warehouse.id
        warehouse = selected_warehouse.name
    start = datetime.strptime(date_start, '%Y-%m-%d').date()
    end = datetime.strptime(date_end, '%Y-%m-%d').date()
    if not selected_warehouse:
        orders = []
    else:
        query = SalesOrder.query.options(joinedload(SalesOrder.customer), joinedload(SalesOrder.salesperson), selectinload(SalesOrder.items).joinedload(SalesOrderItem.material)).filter(SalesOrder.date >= start, SalesOrder.date <= end, SalesOrder.status != 'cancelled')
        if drill_customer_id:
            query = query.filter(SalesOrder.customer_id == drill_customer_id)
        if salesperson_id:
            query = query.filter(SalesOrder.salesperson_id == salesperson_id)
        if status:
            query = query.filter(SalesOrder.status == status)
        if shipment_status:
            query = query.filter(SalesOrder.shipment_status == shipment_status)
        if project_no:
            query = query.filter(SalesOrder.project_no.like(f'%{project_no}%'))
        query = query.filter(db.or_(SalesOrder.warehouse_id == selected_warehouse.id, db.and_(SalesOrder.warehouse_id.is_(None), SalesOrder.warehouse == selected_warehouse.name)))
        orders = query.order_by(SalesOrder.date.asc(), SalesOrder.id.asc()).all()
    drill_material_id = None
    if drill_material_code:
        material = Material.query.filter_by(code=drill_material_code).first()
        if material:
            drill_material_id = material.id
    by_customer = {}
    by_material = {}
    by_salesperson = {}
    total_untaxed = 0
    total_tax = 0
    for order in orders:
        customer_key = order.customer.name if order.customer else '未设置客户'
        customer_row = by_customer.setdefault(customer_key, {'name': customer_key, 'customer_id': order.customer_id, 'orders': 0, 'quantity': 0, 'amount': 0, 'untaxed_amount': 0, 'tax_amount': 0})
        customer_row['orders'] += 1
        customer_row['amount'] += order.total_amount or 0
        customer_row['untaxed_amount'] += order.untaxed_amount or 0
        customer_row['tax_amount'] += order.tax_amount or 0
        salesperson_key = order.salesperson.name if order.salesperson else '未设置业务员'
        salesperson_row = by_salesperson.setdefault(salesperson_key, {'name': salesperson_key, 'orders': 0, 'amount': 0, 'untaxed_amount': 0, 'tax_amount': 0})
        salesperson_row['orders'] += 1
        salesperson_row['amount'] += order.total_amount or 0
        salesperson_row['untaxed_amount'] += order.untaxed_amount or 0
        salesperson_row['tax_amount'] += order.tax_amount or 0
        for item in order.items:
            customer_row['quantity'] += item.quantity or 0
            material_key = item.material.code if item.material else '-'
            material_row = by_material.setdefault(material_key, {'code': material_key, 'name': item.material.name if item.material else '-', 'quantity': 0, 'amount': 0, 'untaxed_amount': 0, 'tax_amount': 0})
            material_row['quantity'] += item.quantity or 0
            material_row['amount'] += item.tax_included_amount or item.amount or 0
            material_row['untaxed_amount'] += item.untaxed_amount or 0
            material_row['tax_amount'] += item.tax_amount or 0
    shipped_quantity = round_to_2_decimals(sum((item.shipped_quantity or 0 for order in orders for item in order.items)))
    shipped_amount = round_to_2_decimals(sum(((item.shipped_quantity or 0) * (item.price or 0) for order in orders for item in order.items)))
    total_amount = round_to_2_decimals(sum((order.total_amount or 0 for order in orders)))
    total_untaxed = round_to_2_decimals(sum((order.untaxed_amount or 0 for order in orders)))
    total_tax = round_to_2_decimals(sum((order.tax_amount or 0 for order in orders)))
    drill_items_all = []
    if drill_material_id:
        for order in orders:
            for item in order.items:
                if item.material_id == drill_material_id:
                    drill_items_all.append({'order_no': order.order_no, 'order_id': order.id, 'date': order.date, 'customer': order.customer.name if order.customer else '-', 'salesperson': order.salesperson.name if order.salesperson else '-', 'quantity': item.quantity or 0, 'price': item.price or 0, 'tax_rate': item.tax_rate or 0, 'untaxed_amount': item.untaxed_amount or 0, 'tax_amount': item.tax_amount or 0, 'tax_included_amount': item.tax_included_amount or item.amount or 0, 'shipped_quantity': item.shipped_quantity or 0})
    drill_page = request.args.get('page', 1, type=int) or 1
    drill_per_page = max(1, min(request.args.get('per_page', 50, type=int) or 50, 200))
    drill_total = len(drill_items_all)
    drill_pages = max(1, (drill_total + drill_per_page - 1) // drill_per_page)
    drill_page = min(max(1, drill_page), drill_pages)
    _dstart = (drill_page - 1) * drill_per_page
    drill_items = drill_items_all[_dstart:_dstart + drill_per_page]
    drill_pagination = SimpleNamespace(total=drill_total, page=drill_page, per_page=drill_per_page, pages=drill_pages, has_prev=drill_page > 1, prev_num=drill_page - 1, has_next=drill_page < drill_pages, next_num=drill_page + 1)
    return dict(date_start=date_start, date_end=date_end, drill_customer_id=drill_customer_id, drill_material_code=drill_material_code, salesperson_id=salesperson_id, status=status, shipment_status=shipment_status, project_no=project_no, warehouse=warehouse, warehouse_id=warehouse_id or '', drill_material_name=db.session.get(Material, drill_material_id).name if drill_material_id else '', drill_items=drill_items, drill_pagination=drill_pagination, customers=Customer.query.order_by(Customer.code.asc()).all(), employees=Employee.query.order_by(Employee.id.asc()).all(), warehouses=Warehouse.query.filter_by(status='active').order_by(Warehouse.code.asc()).all(), orders=orders, by_customer=sorted(by_customer.values(), key=lambda row: row['amount'], reverse=True), by_material=sorted(by_material.values(), key=lambda row: row['amount'], reverse=True), by_salesperson=sorted(by_salesperson.values(), key=lambda row: row['amount'], reverse=True), total_amount=total_amount, total_untaxed=total_untaxed, total_tax=total_tax, shipped_amount=shipped_amount, pending_amount=round_to_2_decimals(total_amount - shipped_amount), shipped_quantity=shipped_quantity, total_orders=len(orders), status_label=sales_status_label, _drill_items_all=drill_items_all)


def sales_outflow_context():
    """销售出库明细表：基于 OutOrder(business_type='销售出库') 的实际出库记录。"""
    from datetime import date, datetime
    import re
    from app import OutOrder, SalesOrder, Warehouse, _require_report_warehouse, round_to_2_decimals, sales_status_label
    date_start = request.args.get('date_start') or date.today().replace(day=1).isoformat()
    date_end = request.args.get('date_end') or date.today().isoformat()
    search = (request.args.get('search') or '').strip()
    warehouse_id = request.args.get('warehouse_id', type=int)
    warehouse = (request.args.get('warehouse') or '').strip()
    selected_warehouse, _wh_err = _require_report_warehouse()
    if selected_warehouse:
        warehouse_id = selected_warehouse.id
        warehouse = selected_warehouse.name
    customer_name = (request.args.get('customer') or '').strip()
    try:
        start = datetime.strptime(date_start, '%Y-%m-%d').date()
        end = datetime.strptime(date_end, '%Y-%m-%d').date()
    except ValueError:
        return (jsonify({'status': 'error', 'msg': '日期格式不正确'}), 400)
    if not selected_warehouse:
        out_orders = []
    else:
        query = OutOrder.query.filter(OutOrder.business_type == '销售出库', OutOrder.date >= start, OutOrder.date <= end)
        if search:
            like = f'%{search}%'
            query = query.filter(db.or_(OutOrder.order_no.like(like), OutOrder.customer.like(like), OutOrder.purpose.like(like), OutOrder.remark.like(like)))
        query = query.filter(OutOrder.warehouse == warehouse)
        if customer_name:
            query = query.filter(OutOrder.customer.like(f'%{customer_name}%'))
        out_orders = query.order_by(OutOrder.date.desc(), OutOrder.id.desc()).all()
    rows = []
    total_quantity = 0
    total_amount = 0
    total_untaxed = 0
    total_tax = 0
    for oo in out_orders:
        sales_order = None
        if oo.source_sales_order_id:
            sales_order = SalesOrder.query.get(oo.source_sales_order_id)
        if not sales_order:
            source_match = re.search('来源销售订单\\s+([^\\s]+)', oo.purpose or '')
            if source_match:
                sales_order = SalesOrder.query.filter_by(order_no=source_match.group(1)).first()
        sales_item_map = {}
        if sales_order:
            sales_item_map = {item.material_id: item for item in sales_order.items}
        for item in oo.items or []:
            material = item.material
            line_amount = item.amount or round_to_2_decimals((item.quantity or 0) * (item.price or 0))
            sales_item = sales_item_map.get(item.material_id) if sales_item_map else None
            tax_rate = sales_item.tax_rate if sales_item and sales_item.tax_rate else 0.13
            untaxed_amount = round_to_2_decimals(line_amount / (1 + tax_rate)) if 1 + tax_rate > 0 else line_amount
            tax_amount = round_to_2_decimals(line_amount - untaxed_amount)
            rows.append({'out_date': oo.date, 'out_order_no': oo.order_no, 'out_order_id': oo.id, 'out_status': oo.status, 'customer': oo.customer or '', 'warehouse': oo.warehouse or '', 'material_code': material.code if material else '', 'material_name': material.name if material else '', 'spec': material.spec if material else '', 'unit': material.unit.name if material and material.unit else '', 'quantity': item.quantity or 0, 'price': item.price or 0, 'tax_rate': tax_rate, 'untaxed_amount': untaxed_amount, 'tax_amount': tax_amount, 'amount': line_amount, 'purpose': oo.purpose or '', 'remark': item.remark or oo.remark or ''})
            if oo.status == 'completed':
                total_quantity += item.quantity or 0
                total_amount += line_amount
                total_untaxed += untaxed_amount
                total_tax += tax_amount
    return dict(date_start=date_start, date_end=date_end, search=search, warehouse=warehouse, warehouse_id=warehouse_id or '', customer=customer_name, warehouses=Warehouse.query.filter_by(status='active').order_by(Warehouse.code.asc()).all(), rows=rows, total_quantity=round_to_2_decimals(total_quantity), total_amount=round_to_2_decimals(total_amount), total_untaxed=round_to_2_decimals(total_untaxed), total_tax=round_to_2_decimals(total_tax), total_rows=len(rows), status_label=sales_status_label)


def sales_trend_context():
    """销售趋势分析表：按月聚合销售订单金额、数量和发货进度。"""
    from datetime import date, timedelta
    from app import SalesOrder, Warehouse, _require_report_warehouse, round_to_2_decimals
    months_back = request.args.get('months', 12, type=int)
    months_back = max(1, min(months_back, 60))
    warehouse_id = request.args.get('warehouse_id', type=int)
    selected_warehouse, _wh_err = _require_report_warehouse()
    today = date.today()
    start_month = (today.replace(day=1) - timedelta(days=months_back * 31)).replace(day=1)
    if not selected_warehouse:
        orders = []
    else:
        orders = SalesOrder.query.filter(SalesOrder.date >= start_month, SalesOrder.status != 'cancelled', db.or_(SalesOrder.warehouse_id == selected_warehouse.id, db.and_(SalesOrder.warehouse_id.is_(None), SalesOrder.warehouse == selected_warehouse.name))).order_by(SalesOrder.date.asc()).all()
    by_month = {}
    for order in orders:
        if not order.date:
            continue
        key = order.date.strftime('%Y-%m')
        row = by_month.setdefault(key, {'month': key, 'orders': 0, 'untaxed_amount': 0, 'tax_amount': 0, 'amount': 0, 'quantity': 0, 'shipped_amount': 0, 'customers': set()})
        row['orders'] += 1
        row['untaxed_amount'] += order.untaxed_amount or 0
        row['tax_amount'] += order.tax_amount or 0
        row['amount'] += order.total_amount or 0
        row['shipped_amount'] += order.shipped_amount or 0
        if order.customer_id:
            row['customers'].add(order.customer_id)
        for item in order.items:
            row['quantity'] += item.quantity or 0
    months = sorted(by_month.keys())
    rows = []
    prev_amount = None
    for m in months:
        row = by_month[m]
        row['customers'] = len(row['customers'])
        row['untaxed_amount'] = round_to_2_decimals(row['untaxed_amount'])
        row['tax_amount'] = round_to_2_decimals(row['tax_amount'])
        row['amount'] = round_to_2_decimals(row['amount'])
        row['shipped_amount'] = round_to_2_decimals(row['shipped_amount'])
        if prev_amount and prev_amount > 0:
            row['growth'] = round_to_2_decimals((row['amount'] - prev_amount) / prev_amount * 100)
        else:
            row['growth'] = None
        rows.append(row)
        prev_amount = row['amount']
    return dict(months_back=months_back, period_start=start_month.isoformat(), warehouse_id=selected_warehouse.id if selected_warehouse else '', warehouse=selected_warehouse.name if selected_warehouse else '', warehouses=Warehouse.query.filter_by(status='active').order_by(Warehouse.code.asc()).all(), rows=rows, total_orders=sum((r['orders'] for r in rows)), total_amount=round_to_2_decimals(sum((r['amount'] for r in rows))), total_untaxed=round_to_2_decimals(sum((r['untaxed_amount'] for r in rows))), total_tax=round_to_2_decimals(sum((r['tax_amount'] for r in rows))), total_shipped=round_to_2_decimals(sum((r['shipped_amount'] for r in rows))))


def sales_execution_context():
    from app import Customer, SalesOrder, Warehouse, _sales_report_filters_context, _sales_report_orders, round_to_2_decimals, sales_status_label
    orders = _sales_report_orders()
    rows = []
    for order in orders:
        total_qty = sum((item.quantity or 0 for item in order.items))
        shipped_qty = sum((item.shipped_quantity or 0 for item in order.items))
        remaining_qty = max(total_qty - shipped_qty, 0)
        amount = float(order.total_amount or 0)
        shipped_amount = float(order.shipped_amount or 0)
        rows.append({'order': order, 'quantity': round_to_2_decimals(total_qty), 'shipped_quantity': round_to_2_decimals(shipped_qty), 'remaining_quantity': round_to_2_decimals(remaining_qty), 'amount': round_to_2_decimals(amount), 'shipped_amount': round_to_2_decimals(shipped_amount), 'remaining_amount': round_to_2_decimals(max(amount - shipped_amount, 0)), 'execution_rate': round_to_2_decimals(shipped_amount / amount * 100) if amount else 0})
    return dict(rows=rows, filters=_sales_report_filters_context(), customers=Customer.query.order_by(Customer.code.asc(), Customer.id.asc()).all(), warehouses=Warehouse.query.filter_by(status='active').order_by(Warehouse.code.asc()).all(), status_label=sales_status_label, total_orders=len(rows), total_amount=round_to_2_decimals(sum((row['amount'] for row in rows))), shipped_amount=round_to_2_decimals(sum((row['shipped_amount'] for row in rows))), remaining_amount=round_to_2_decimals(sum((row['remaining_amount'] for row in rows))))


def sales_price_context():
    from app import Customer, SalesOrder, Warehouse, _sales_report_filters_context, _sales_report_orders, price_dimension_enabled, round_to_2_decimals
    if not price_dimension_enabled():
        flash('价格维度未启用，销售价格分析已下线；如需启用请把系统设置 price_dimension_enabled 设为 1。', 'warning')
        return redirect(url_for('sales_report'))
    aggregates = {}
    for order in _sales_report_orders():
        customer_id = order.customer_id
        for item in order.items:
            material = item.material
            if not material:
                continue
            key = material.id
            row = aggregates.setdefault(key, {'code': material.code, 'name': material.name, 'spec': material.spec or '', 'quantity': 0, 'amount': 0, 'prices': [], 'customers': set(), 'lines': 0})
            quantity = float(item.quantity or 0)
            price = float(item.price or 0)
            row['quantity'] += quantity
            row['amount'] += float(item.tax_included_amount or item.amount or 0)
            row['prices'].append(price)
            row['customers'].add(customer_id)
            row['lines'] += 1
    rows = []
    for row in aggregates.values():
        prices = row.pop('prices')
        row['avg_price'] = round_to_2_decimals(sum(prices) / len(prices)) if prices else 0
        row['min_price'] = round_to_2_decimals(min(prices)) if prices else 0
        row['max_price'] = round_to_2_decimals(max(prices)) if prices else 0
        row['quantity'] = round_to_2_decimals(row['quantity'])
        row['amount'] = round_to_2_decimals(row['amount'])
        row['customers'] = len({value for value in row['customers'] if value})
        rows.append(row)
    rows.sort(key=lambda row: row['amount'], reverse=True)
    return dict(rows=rows, filters=_sales_report_filters_context(), customers=Customer.query.order_by(Customer.code.asc(), Customer.id.asc()).all(), warehouses=Warehouse.query.filter_by(status='active').order_by(Warehouse.code.asc()).all(), total_quantity=round_to_2_decimals(sum((row['quantity'] for row in rows))), total_amount=round_to_2_decimals(sum((row['amount'] for row in rows))))


SALES_CONTEXT_BUILDERS = {
    'summary': sales_summary_context, 'reconciliation': sales_reconciliation_context,
    'outflow': sales_outflow_context, 'trend': sales_trend_context,
    'execution': sales_execution_context, 'price_analysis': sales_price_context,
}


def collect_sales_report_print(report_kind):
    """Select a page grain explicitly, collecting every row independently of pagination."""
    from app import (_require_report_warehouse, price_dimension_enabled,
                     sales_status_label, sales_shipment_status_label)
    from doc_print_excel import SALES_REPORT_PRINT_SPECS
    if report_kind not in SALES_CONTEXT_BUILDERS:
        raise ValueError('不支持的销售报表')
    scope = (request.args.get('scope') or '').strip()
    if report_kind == 'summary' and scope not in ('customer', 'material', 'salesperson', 'detail'):
        raise ValueError('请选择销售报表打印范围：按客户、按物料、按业务员或明细')
    if report_kind == 'price_analysis' and not price_dimension_enabled():
        raise ValueError('价格维度未启用，销售价格分析已下线')
    warehouse, error = _require_report_warehouse()
    if warehouse is None:
        raise ValueError(error or '请选择仓库')
    key = 'summary_' + scope if report_kind == 'summary' else report_kind
    title, specs = SALES_REPORT_PRINT_SPECS[key]
    context = SALES_CONTEXT_BUILDERS[report_kind]()
    if not isinstance(context, dict):
        raise ValueError('报表筛选条件无效，请检查日期和仓库')
    if report_kind == 'summary':
        if scope == 'detail':
            material_code = (request.args.get('material_code') or '').strip()
            rows = []
            for order in context['orders']:
                for item in order.items:
                    material = item.material
                    if material_code and (not material or material.code != material_code):
                        continue
                    rows.append(dict(order_no=order.order_no, date=order.date,
                        customer=order.customer.name if order.customer else '-',
                        salesperson=order.salesperson.name if order.salesperson else '-',
                        material_code=material.code if material else '-',
                        material_name=material.name if material else '-',
                        spec=material.spec if material else '',
                        unit=material.unit.name if material and material.unit else '',
                        quantity=item.quantity or 0, shipped_quantity=item.shipped_quantity or 0,
                        price=item.price or 0, tax_rate=item.tax_rate or 0,
                        untaxed_amount=item.untaxed_amount or 0, tax_amount=item.tax_amount or 0,
                        amount=item.tax_included_amount or item.amount or 0,
                        contract_no=order.contract_no or '', project_name=order.project_name or '',
                        status=sales_status_label(order.status),
                        shipment_status=sales_shipment_status_label(order.shipment_status)))
        else:
            rows = context['by_' + scope]
    elif report_kind in ('execution', 'reconciliation'):
        rows = []
        for original in context['rows']:
            order = original['order']
            row = {key: value for key, value in original.items() if key != 'order'}
            row.update(order_no=order.order_no, date=order.date, warehouse=order.warehouse or '',
                customer=order.customer.name if order.customer else '-',
                contract_no=order.contract_no or '', project_name=order.project_name or '',
                status=sales_status_label(order.status))
            if report_kind == 'reconciliation':
                row['result'] = '通过' if row['ok'] else '不一致'
            rows.append(row)
    else:
        rows = [dict(row) for row in context['rows']]
        if report_kind == 'outflow':
            for row in rows:
                row['out_status'] = sales_status_label(row['out_status'])
        elif report_kind == 'trend':
            for row in rows:
                row['shipment_rate'] = round(row['shipped_amount'] / row['amount'] * 100, 1) if row['amount'] > 0 else '-'
                if row['growth'] is None:
                    row['growth'] = '-'
    filters = context.get('filters', context)
    scope_parts = [f'仓库：{warehouse.name}（{warehouse.code}）']
    if report_kind == 'trend':
        scope_parts.append(f"日期：{context['period_start']}起（近{context['months_back']}个月筛选）")
    else:
        scope_parts.append(f"日期：{filters.get('date_start') or '不限'} 至 {filters.get('date_end') or '不限'}")
    return dict(title=title, target_code='report_sales_' + key, print_scope=' · '.join(scope_parts),
                columns=[dict(title=label, field=field) for label, field, width in specs], rows=rows)


def resolve_sales_report_print_template(target_code):
    """Resolve only this report's selected template, keeping stale defaults harmless."""
    from app import ExcelPrintTemplate
    from print_fill import template_file_abspath
    query = ExcelPrintTemplate.query.filter_by(target_type='report',
        target_code=target_code, template_type='excel')
    selected = (request.args.get('template_id') or '').strip()
    if selected:
        if not selected.isdecimal() or len(selected) > 18 or int(selected) <= 0:
            raise ValueError('所选打印模板无效，请重新选择')
        template = query.filter_by(id=int(selected)).first()
        if template is None:
            raise ValueError('所选模板不存在或不属于当前报表范围，请重新选择')
    else:
        template = query.filter_by(is_default=True).first()
    if template is not None:
        candidate = template_file_abspath(template.excel_template_path, current_app.static_folder)
        path = Path(candidate).resolve() if candidate else None
        if path and path.is_relative_to(Path(current_app.static_folder).resolve()) and path.is_file():
            return path, template.name
        if selected:
            raise ValueError('所选打印模板文件不存在或不可用，请修复或重新选择模板')
    return None, '系统内置模板'


def render_sales_report_excel(title, columns, rows, template_path=None, print_scope=''):
    """Fill the chosen workbook in memory; never register or save templates during GET."""
    from openpyxl import load_workbook
    from doc_print_excel import generate_report_builtin_template
    from print_fill import _Filler, validate_template_file
    from services.report_print_preview import MAX_PREVIEW_ROWS, MAX_PREVIEW_CELLS, _fill_sheet
    if len(rows) > MAX_PREVIEW_ROWS:
        raise ValueError('当前结果超过模板打印上限（5000 行），请缩小筛选范围')
    if template_path:
        path = Path(template_path)
        if not path.is_file():
            raise ValueError('所选打印模板文件不存在')
        if path.stat().st_size > 10 * 1024 * 1024:
            raise ValueError('打印模板文件过大，请精简模板')
        raw = path.read_bytes()
    else:
        builtin = generate_report_builtin_template(title, columns)
        if builtin is None:
            raise ValueError('当前报表没有可打印的字段')
        raw = builtin.getvalue()
    error = validate_template_file(raw)
    if error:
        raise ValueError(error)
    workbook = load_workbook(io.BytesIO(raw), keep_links=False)
    try:
        if len(workbook.worksheets) > 10:
            raise ValueError('模板工作表过多，请精简模板')
        order = SimpleNamespace(title=title, operator=SimpleNamespace(username=''),
            total_amount=sum(float(row.get('amount') or 0) for row in rows))
        items = [SimpleNamespace(**{'quantity': row.get('quantity', 0), **row}) for row in rows]
        for sheet in workbook.worksheets:
            if (sheet.max_row + max(0, len(items) - 1)) * sheet.max_column > MAX_PREVIEW_CELLS:
                raise ValueError('展开后的打印内容过大，请缩小筛选范围')
            _fill_sheet(sheet, _Filler(order, items, date.today().isoformat()))
            if print_scope:
                for header in (sheet.oddHeader, sheet.evenHeader, sheet.firstHeader):
                    existing = header.center.text or ''
                    header.center.text = (existing + '\n' if existing else '') + print_scope.replace('&', '&&')
                    if not existing:
                        header.center.size = 8
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output
    finally:
        workbook.close()
