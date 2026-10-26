"""Read-only stock print inputs shared by Excel downloads and browser preview."""
from pathlib import Path


def collect_stock_print_data(args):
    """Use the screen's full filtered query and warehouse ledger, never its page limit."""
    from app import get_warehouse_stock_quantities, resolve_request_warehouse
    from routes.stock_query import build_stock_material_query, get_stock_query_filters
    warehouse, error = resolve_request_warehouse(args)
    if error:
        raise ValueError(error)
    quantities = get_warehouse_stock_quantities(warehouse)
    materials = build_stock_material_query(get_stock_query_filters(args), quantities).all()
    return warehouse, materials, quantities


def resolve_stock_print_template(args, static_folder):
    """Validate an explicit choice without silently replacing the requested template."""
    from app import ExcelPrintTemplate
    from print_fill import template_file_abspath
    raw_id = (args.get('template_id') or '').strip()
    query = ExcelPrintTemplate.query.filter_by(
        target_type='list', target_code='stock_query', template_type='excel')
    if raw_id:
        if not raw_id.isdecimal() or len(raw_id) > 18 or int(raw_id) <= 0:
            raise ValueError('所选打印模板无效，请重新选择')
        template = query.filter_by(id=int(raw_id)).first()
        if template is None:
            raise ValueError('所选模板不存在或不属于库存查询，请重新选择')
    else:
        template = query.filter_by(is_default=True).first() or query.order_by(
            ExcelPrintTemplate.updated_at.desc()).first()
    if template is None:
        return None, None
    candidate = template_file_abspath(template.excel_template_path, static_folder)
    candidate = Path(candidate).resolve() if candidate else None
    if candidate and candidate.is_relative_to(Path(static_folder).resolve()) and candidate.is_file():
        return template, candidate
    if raw_id:
        raise ValueError('所选打印模板文件不存在或不可用，请修复或重新选择模板')
    return None, None


def stock_print_rows(warehouse, materials, quantities, include_prices=True):
    """Keep legacy stock template keys, withholding price values when disabled."""
    rows = []
    for material in materials:
        stock = quantities.get(material.id, 0)
        price = material.price or 0
        rows.append(dict(warehouse=warehouse.name or '', code=material.code or '',
            brand=material.brand or '', name=material.name or '', spec=material.spec or '',
            unit=material.unit.name if material.unit else '', stock=stock, quantity=stock,
            price=price if include_prices else '', amount=stock * price if include_prices else ''))
    return rows


def stock_print_columns():
    """Keep the existing stock namespace column keys for the in-memory default."""
    from doc_print_excel import TABLE_EXCEL_PRINT_TYPES
    return [{'field': key, 'title': title}
        for title, key, _width in TABLE_EXCEL_PRINT_TYPES['stock_query']['sheets'][0]['columns']]


def build_stock_excel_print(title, rows, include_prices=True, template_path=None):
    """Fill selected or built-in stock layouts in memory, matching browser row expansion."""
    import io
    from datetime import date
    from types import SimpleNamespace
    from openpyxl import load_workbook
    from doc_print_excel import generate_report_builtin_template
    from print_fill import _Filler, validate_template_file
    from services.report_print_preview import (MAX_PREVIEW_ROWS, MAX_PREVIEW_CELLS,
        _fill_sheet, visible_report_data)
    if len(rows) > MAX_PREVIEW_ROWS:
        raise ValueError('当前结果超过模板打印上限（5000 行），请缩小筛选范围')
    columns, rows = visible_report_data(stock_print_columns(), rows, include_prices)
    if template_path:
        path = Path(template_path)
        if not path.is_file():
            raise ValueError('所选打印模板文件不存在')
        if path.stat().st_size > 10 * 1024 * 1024:
            raise ValueError('打印模板文件过大，请精简模板')
        raw = path.read_bytes()
    else:
        raw = generate_report_builtin_template(title, columns).getvalue()
    error = validate_template_file(raw)
    if error:
        raise ValueError(error)
    workbook = load_workbook(io.BytesIO(raw), keep_links=False)
    order = SimpleNamespace(title=title, operator=SimpleNamespace(username=''),
        total_amount=sum(float(row.get('amount') or 0) for row in rows) if include_prices else '')
    try:
        if len(workbook.worksheets) > 10:
            raise ValueError('模板工作表过多，请精简模板')
        items = [SimpleNamespace(**row) for row in rows]
        for sheet in workbook.worksheets:
            if sheet.sheet_state != 'visible':
                continue
            if (sheet.max_row + max(0, len(items) - 1)) * sheet.max_column > MAX_PREVIEW_CELLS:
                raise ValueError('展开后的打印内容过大，请缩小筛选范围')
            _fill_sheet(sheet, _Filler(order, items, date.today().isoformat()))
            if sheet.max_row * sheet.max_column > MAX_PREVIEW_CELLS:
                raise ValueError('打印内容过大，请缩小筛选范围')
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output
    finally:
        workbook.close()
