import os,sys,json,io,tempfile,subprocess
from pathlib import Path
ROOT=Path(__file__).resolve().parents[2]
os.environ.update(DATABASE_URL='sqlite:///:memory:',WMS_DATABASE_URI='sqlite:///:memory:',WMS_NO_DB_TOUCH='1',WMS_SKIP_AUTO_UPDATE='1',WMS_ALLOW_INSECURE_COOKIE='1',LOG_FILE=str(Path(tempfile.gettempdir())/'wms-audit-probe.log'))
sys.path.insert(0,str(ROOT/'app'))
import app as a
from openpyxl import Workbook,load_workbook
from flask_login import login_user
a.app.config.update(TESTING=True,WTF_CSRF_ENABLED=False)
with a.app.app_context():
 assert a.app.config['SQLALCHEMY_DATABASE_URI']=='sqlite:///:memory:', 'Audit requires isolated in-memory database'
 a.db.create_all()
 u=a.User(username='audit',password_hash='disabled-fixture',role='admin',must_change_password=False)
 wa=a.Warehouse(code='A',name='仓库A',is_default=True,status='active');wb=a.Warehouse(code='B',name='仓库B',status='active');unit=a.Unit(code='U',name='个')
 a.db.session.add_all([u,wa,wb,unit]);a.db.session.flush()
 m=a.Material(code='AUD-M',name='审计物料',spec='X',unit_id=unit.id,stock=20)
 a.db.session.add(m);a.db.session.flush();a.db.session.add(a.StockTransaction(material_id=m.id,warehouse_id=wa.id,transaction_type='in',quantity=20,location=wa.name));a.db.session.commit(); uid=u.id;wid=wa.id;bid=wb.id;mid=m.id
client=a.app.test_client()
with client.session_transaction() as s:s['_user_id']=str(uid);s['_fresh']=True
result={}
source=(ROOT/'app/templates/document_table_form.html').read_text(encoding='utf-8')
fn='function collectHeader() {'+source.split('function collectHeader() {',1)[1].split('\nfunction submitForm',1)[0]
js='const docType="requisition";const document={querySelector(s){let n=s.match(/name="([^"]+)"/)[1];return {value:({warehouse_id:"1",picker:"张三"})[n]||""};}};'+fn+';process.stdout.write(JSON.stringify(collectHeader()));'
header=json.loads(subprocess.check_output(['node','-e',js],text=True,encoding='utf-8'))
response=client.post('/requisition/save_table',json={'order_no':'AUD-REQ','date':'2026-10-19','header':header,'items':[{'code':'AUD-M','quantity':1}]})
with a.app.app_context():
 req=a.ProductionRequisition.query.filter_by(req_no='AUD-REQ').first();result['picker']={'ui_value':'张三','payload_has_picker':'picker'in header,'saved_picker':req.picker if req else None,'http':response.status_code,'response':response.get_json()}
book=Workbook();sheet=book.active;sheet.append(['单据编号','仓库','物料编码','物料名称','数量']);sheet.append(['AUD-IMPORT','仓库B','AUD-M','审计物料',2]);buf=io.BytesIO();book.save(buf);buf.seek(0)
r=client.post('/requisition/import',data={'file':(buf,'audit.xlsx')},content_type='multipart/form-data')
with a.app.app_context():
 req=a.ProductionRequisition.query.filter_by(req_no='AUD-IMPORT').first();result['import_warehouse']={'specified':'仓库B','saved':req.warehouse if req else None,'http':r.status_code}
 order=a.InOrder(order_no='AUD-PENDING',warehouse='仓库A',status='pending');a.db.session.add(order);a.db.session.flush();a.db.session.add(a.InOrderItem(in_order_id=order.id,material_id=mid,quantity=7,price=0,amount=0));a.db.session.commit()
r=client.get('/report/inout/print',query_string={'warehouse_id':wid,'status':'completed','search':'NONMATCH'})
rows=list(load_workbook(io.BytesIO(r.data)).worksheets[0].values)
result['report_filter']={'http':r.status_code,'pending_in_completed_filter':any(row[0]=='AUD-PENDING' for row in rows),'rows':len(rows)-1}
a.app.config['WTF_CSRF_ENABLED']=True
r=client.post('/mobile/api/scan_submit',json={'mode':'in','code':'AUD-M','warehouse_id':wid,'quantity':2})
with a.app.app_context():result['csrf_session_scan']={'http':r.status_code,'status':(r.get_json() or {}).get('status'),'stock_after_no_csrf':a.db.session.get(a.Material,mid).stock}
anon=a.app.test_client();codes=[anon.post('/api/login',json={'username':'audit-nonexistent-'+str(i),'password':'invalid'}).status_code for i in range(65)]
result['api_login_rate']={'requests':len(codes),'status_counts':{str(x):codes.count(x) for x in set(codes)}}
with a.app.test_request_context('/'):
 login_user(a.db.session.get(a.User,uid))
 a.location_management_enabled=lambda:True
 mat=a.Material(code='AUD-LOC',name='库位审计物料',spec='Y',unit_id=1,stock=0);a.db.session.add(mat);a.db.session.flush()
 wh=a.db.session.get(a.Warehouse,wid)
 opening,delta=a._apply_opening_stock_balance(None,mat,10,0,0,'audit',wh,location='L-A');a.db.session.commit()
 opening,delta=a._apply_opening_stock_balance(opening,mat,10,0,0,'audit',wh,location='L-B');a.db.session.commit()
 result['opening_location_change']={'line_location':opening.location,'delta':delta,'location_rows':[{'location':x.location,'quantity':x.quantity} for x in a.LocationInventory.query.filter_by(material_id=mat.id).all()]}
a.app.config['WTF_CSRF_ENABLED']=False
with a.app.app_context():
 a.db.session.add(a.InventoryCheck(check_no='AUD-CHECK',warehouse='仓库A',status='pending'))
 a.db.session.add(a.AdjustmentOrder(adjustment_no='AUD-ADJUST',adjustment_type='profit',warehouse='仓库A',status='pending'))
 a.db.session.commit()
r=client.post('/warehouse/'+str(wid)+'/edit',data={'code':'A','name':'仓库A新名','status':'active'})
with a.app.app_context():
 result['warehouse_rename']={'http':r.status_code,'warehouse_name':a.db.session.get(a.Warehouse,wid).name,'check_warehouse':a.InventoryCheck.query.filter_by(check_no='AUD-CHECK').first().warehouse,'adjustment_warehouse':a.AdjustmentOrder.query.filter_by(adjustment_no='AUD-ADJUST').first().warehouse}
print(json.dumps(result,ensure_ascii=False,indent=2))
Path(tempfile.gettempdir(),'wms-audit-20261020-proofs.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
