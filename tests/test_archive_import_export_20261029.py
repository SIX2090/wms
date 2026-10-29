"""Regressions for BTN29 header/status/hierarchy/export observations."""
import io
from urllib.parse import parse_qs, urlsplit

import pytest
from openpyxl import load_workbook

import app as m
from tests.test_browser_audit_20261029_repairs import audit_case, excel, Elements  # noqa: F401


ARCHIVES = [
    ('unit', m.Unit, ['单位编号', '单位名称']),
    ('supplier', m.Supplier, ['供应商编号', '供应商名称']),
    ('category', m.MaterialCategory, ['分类编码', '分类名称']),
    ('customer', m.Customer, ['客户编号', '客户名称']),
    ('warehouse', m.Warehouse, ['仓库编码', '仓库名称']),
]


def test_master_import_rows():
    from master_import import MasterImportError, master_import_rows
    workbook = load_workbook(excel(['单位名称', '单位编号'], [['零号单位', 0]]))
    assert list(master_import_rows(workbook.active, ('单位编号', '单位名称'))) == [
        {'单位名称': '零号单位', '单位编号': '0'}]
    workbook = load_workbook(excel(['单位编号', '单位名称', '单位编号'], [['A', '甲', 'B']]))
    with pytest.raises(MasterImportError, match='重复列'):
        master_import_rows(workbook.active, ('单位编号', '单位名称'))


def test_normalize_master_status():
    from master_import import MasterImportError, normalize_master_status
    assert normalize_master_status(' 停用 ') == 'inactive'
    assert normalize_master_status('ACTIVE') == 'active'
    with pytest.raises(MasterImportError, match='状态无效'):
        normalize_master_status('disabled')


def test_build_category_tree_rows(audit_case):
    with m.app.app_context():
        first = m.MaterialCategory(code='Z', name='先输入')
        second = m.MaterialCategory(code='A', name='后输入')
        m.db.session.add_all([first, second])
        m.db.session.flush()
        assert [row['category'].code for row in m.build_category_tree_rows([first, second])] == ['A', 'Z']
        assert [row['category'].code for row in m.build_category_tree_rows([first, second], preserve_order=True)] == ['Z', 'A']


def upload(client, page, headers, rows, xhr=True):
    return client.post(f'/{page}/import', data={'file': (excel(headers, rows), 'audit.xlsx')},
                       headers={'X-Requested-With': 'XMLHttpRequest'} if xhr else {})


@pytest.mark.parametrize('page,model,headers', ARCHIVES)
@pytest.mark.parametrize('kind', ['wrong', 'missing', 'duplicate', 'empty'])
def test_invalid_archive_headers_write_nothing(audit_case, page, model, headers, kind):
    client, _ = audit_case
    bad = {'wrong': ['错误编号', '错误名称'], 'missing': headers[:1],
           'duplicate': headers + headers[:1], 'empty': []}[kind]
    with m.app.app_context():
        before = model.query.count()
    response = upload(client, page, bad, [['FIX29-WRONG', '禁止写入', 'OTHER']])
    assert response.status_code == 400, response.get_json()
    assert '表头' in response.get_json()['msg']
    with m.app.app_context():
        assert model.query.count() == before
        assert not model.query.filter_by(code='FIX29-WRONG').first()


@pytest.mark.parametrize('page,model,headers', ARCHIVES)
def test_master_import_rows_reordered_archive_headers_and_optional_columns(audit_case, page, model, headers):
    client, _ = audit_case
    response = upload(client, page, ['未知扩展列', headers[1], ' ' + headers[0] + ' '],
                      [['不作编码', '重排档案', 'FIX29-REORDER']])
    assert response.status_code == 200, response.get_json()
    with m.app.app_context():
        row = model.query.filter_by(code='FIX29-REORDER').one()
        assert row.name == '重排档案'
    duplicate = upload(client, page, headers, [['FIX29-REORDER', '不能覆盖原档案']])
    assert '导入 0 条' in duplicate.get_json()['msg']
    with m.app.app_context():
        assert model.query.filter_by(code='FIX29-REORDER').one().name == '重排档案'


def test_supplier_non_ajax_header_error_flashes_without_writes(audit_case):
    client, _ = audit_case
    response = upload(client, 'supplier', ['客户编号', '客户名称'], [['FIX29-WRONG', '不应写入']], False)
    assert response.status_code == 302
    with client.session_transaction() as session:
        assert any('Excel表头缺少必要列' in message for _, message in session['_flashes'])
    with m.app.app_context():
        assert not m.Supplier.query.filter_by(code='FIX29-WRONG').first()


@pytest.mark.parametrize('raw,expected', [('启用', 'active'), ('停用', 'inactive'),
    (' ACTIVE ', 'active'), ('INACTIVE', 'inactive'), ('', 'active')])
def test_normalize_master_status_for_warehouse_import(audit_case, raw, expected):
    client, ids = audit_case
    response = upload(client, 'warehouse', ['状态', '仓库名称', '仓库编码', '仓库位置', '备注', '仓库类型'],
                      [[raw, '状态回归仓', 'FIX29-STATUS', '甲区', '状态测试', '普通']])
    assert response.status_code == 200, response.get_json()
    with m.app.app_context():
        row = m.Warehouse.query.filter_by(code='FIX29-STATUS').one()
        assert (row.status, row.location, row.remark, row.type) == (expected, '甲区', '状态测试', '普通')
        assert not row.is_default
        assert m.Warehouse.query.filter_by(is_default=True).one().id == ids['first']
    for state in ('active', 'inactive'):
        text = client.get('/warehouse', query_string={'status': state, 'search': 'FIX29-STATUS'}).get_data(as_text=True)
        assert ('FIX29-STATUS' in text.split('<tbody>')[1].split('</tbody>')[0]) == (state == expected)
    exported = client.get('/warehouse/export?search=FIX29-STATUS')
    rows = list(load_workbook(io.BytesIO(exported.data)).active.values)
    assert rows[1][4] == ('启用' if expected == 'active' else '停用')


def test_warehouse_invalid_status_rolls_back_all_rows(audit_case):
    client, _ = audit_case
    response = upload(client, 'warehouse', ['仓库编码', '仓库名称', '状态'],
                      [['FIX29-FIRST', '合法先行仓', 'active'], ['FIX29-BAD', '非法状态仓', '随便']])
    assert response.status_code == 400
    assert '状态无效' in response.get_json()['msg']
    with m.app.app_context():
        assert m.Warehouse.query.filter(m.Warehouse.code.like('FIX29-%')).count() == 0


@pytest.mark.parametrize('raw,state', [('启用', 'active'), ('停用', 'inactive'), ('unknown', '')])
def test_legacy_warehouse_status_is_explicit_and_never_rewritten(audit_case, raw, state):
    client, ids = audit_case
    with m.app.app_context():
        row = m.Warehouse(code='FIX29-LEGACY', name='历史状态仓', status=raw)
        m.db.session.add(row)
        m.db.session.commit()
    query = {'search': 'FIX29-LEGACY', 'status': state}
    text = client.get('/warehouse', query_string=query).get_data(as_text=True)
    assert 'FIX29-LEGACY' in text.split('<tbody>')[1].split('</tbody>')[0]
    assert ('历史状态，请编辑确认' if state else '状态异常，请编辑确认') in text
    assert 'setDefaultWarehouse(' not in text.split('<tbody>')[1].split('</tbody>')[0]
    rows = list(load_workbook(io.BytesIO(client.get('/warehouse/export', query_string=query).data)).active.values)
    assert rows[1][4] == raw
    with m.app.app_context():
        assert m.Warehouse.query.filter_by(code='FIX29-LEGACY').one().status == raw
        assert m.Warehouse.query.filter_by(is_default=True).one().id == ids['first']


def test_category_import_resolves_reverse_multilevel_and_existing_parent(audit_case):
    client, _ = audit_case
    headers = ['分类编码', '分类名称', '上级分类编码']
    rows = [['FIX29-C', '层级孙', 'FIX29-B'], ['FIX29-B', '层级子', 'FIX29-A'], ['FIX29-A', '层级根', '']]
    response = upload(client, 'category', headers, rows)
    assert '导入 3 条' in response.get_json()['msg']
    with m.app.app_context():
        a, b, c = [m.MaterialCategory.query.filter_by(code='FIX29-'+code).one() for code in 'ABC']
        assert (b.parent_id, c.parent_id) == (a.id, b.id)
    response = upload(client, 'category', headers, [['FIX29-D', '现有父子级', 'FIX29-A']])
    assert '导入 1 条' in response.get_json()['msg']


def test_category_export_roundtrip_preserves_every_relationship(audit_case):
    client, _ = audit_case
    upload(client, 'category', ['分类编码', '分类名称', '上级分类编码'],
           [['FIX29-Z', '导出父', ''], ['FIX29-A', '导出子', 'FIX29-Z']])
    rows = list(load_workbook(io.BytesIO(client.get('/category/export?search=FIX29-').data)).active.values)
    assert {row[0] for row in rows[1:]} == {'FIX29-A', 'FIX29-Z'}
    # Deliberately reverse the hierarchy, reproducing historical exported files.
    renamed = [[str(value).replace('FIX29-', 'RT29-') if value else '' for value in row]
               for row in sorted(rows[1:], key=lambda row: row[0])]
    for row in renamed:
        row[1] = '往返' + row[1]
    response = upload(client, 'category', list(rows[0]), renamed)
    assert '导入 2 条' in response.get_json()['msg']
    with m.app.app_context():
        child = m.MaterialCategory.query.filter_by(code='RT29-A').one()
        assert child.parent.code == 'RT29-Z'


def test_category_invalid_graph_rows_are_skipped_without_orphans(audit_case):
    client, _ = audit_case
    rows = [['FIX29-A', '循环甲', 'FIX29-B'], ['FIX29-B', '循环乙', 'FIX29-A'],
            ['FIX29-S', '自引用', 'FIX29-S'], ['FIX29-M', '缺父', 'missing'],
            ['FIX29-CHILD', '无效下级', 'FIX29-M'], ['FIX29-OK', '有效根', ''],
            ['FIX29-OK', '重复编码', ''], ['FIX29-DUP', '有效根', '']]
    response = upload(client, 'category', ['分类编码', '分类名称', '上级分类编码'], rows)
    assert '导入 1 条' in response.get_json()['msg']
    assert '跳过 7 条' in response.get_json()['msg']
    with m.app.app_context():
        assert [(r.code, r.parent_id) for r in m.MaterialCategory.query.all()] == [('FIX29-OK', None)]


def test_reverse_deep_category_import_is_iterative(audit_case):
    client, _ = audit_case
    rows = [[f'DEEP29-{i}', f'层级{i}', f'DEEP29-{i-1}' if i else '']
            for i in reversed(range(180))]
    response = upload(client, 'category', ['分类编码', '分类名称', '上级分类编码'], rows)
    assert '导入 180 条' in response.get_json()['msg']
    with m.app.app_context():
        assert m.MaterialCategory.query.count() == 180
        assert m.MaterialCategory.query.filter_by(code='DEEP29-179').one().parent.code == 'DEEP29-178'


@pytest.mark.parametrize('page,model', [('supplier', m.Supplier), ('customer', m.Customer),
    ('unit', m.Unit), ('warehouse', m.Warehouse), ('department', m.Department), ('employee', m.Employee)])
@pytest.mark.parametrize('order', ['asc', 'desc'])
def test_actual_page_export_link_preserves_filtered_sort(audit_case, page, model, order):
    client, ids = audit_case
    with m.app.app_context():
        extras = {'department_id': ids['department']} if page == 'employee' else {}
        m.db.session.add_all([model(code='FIX29-A', name='FIX29-Zed', **extras),
                              model(code='FIX29-Z', name='FIX29-Alpha', **extras)])
        m.db.session.commit()
    params = dict(search='FIX29-', sort='name', order=order, page=1, per_page=10)
    if page in ('warehouse', 'department'):
        params['status'] = 'active'
    if page in ('supplier', 'customer'):
        params.update(search_field='name', has_business='no', date_from='2020-01-01', date_to='2099-12-31')
    if page == 'employee':
        params['department_id'] = ids['department']
    text = client.get('/'+page, query_string=params).get_data(as_text=True)
    link = next(attrs['href'] for tag, attrs in Elements(text).elements
                if tag == 'a' and attrs.get('href', '').startswith('/'+page+'/export'))
    query = parse_qs(urlsplit(link).query)
    for key, value in params.items():
        if key not in ('page', 'per_page'):
            assert query[key] == [str(value)], (link, key)
    response = client.get(link)
    assert response.status_code == 200
    rows = list(load_workbook(io.BytesIO(response.data)).active.values)
    expected = ['FIX29-Z', 'FIX29-A'] if order == 'asc' else ['FIX29-A', 'FIX29-Z']
    assert [row[0] for row in rows[1:]] == expected
    tbody = text.split('<tbody>')[1].split('</tbody>')[0]
    assert tbody.index(expected[0]) < tbody.index(expected[1])


@pytest.mark.parametrize('order', ['asc', 'desc'])
@pytest.mark.parametrize('search', ['CAT29-', 'CAT29-CHILD'])
def test_build_category_tree_rows_sibling_sort_export_matches_tree_and_filter(audit_case, order, search):
    client, _ = audit_case
    with m.app.app_context():
        root = m.MaterialCategory(code='CAT29-ROOT', name='根分类')
        m.db.session.add(root)
        m.db.session.flush()
        m.db.session.add_all([m.MaterialCategory(code='CAT29-CHILD-A', name='Z子', parent_id=root.id),
                              m.MaterialCategory(code='CAT29-CHILD-Z', name='A子', parent_id=root.id)])
        m.db.session.commit()
    text = client.get('/category', query_string=dict(search=search, sort='name', order=order)).get_data(as_text=True)
    link = next(attrs['href'] for tag, attrs in Elements(text).elements
                if tag == 'a' and attrs.get('href', '').startswith('/category/export'))
    assert parse_qs(urlsplit(link).query)['order'] == [order]
    rows = list(load_workbook(io.BytesIO(client.get(link).data)).active.values)
    expected = ['CAT29-CHILD-Z', 'CAT29-CHILD-A'] if order == 'asc' else ['CAT29-CHILD-A', 'CAT29-CHILD-Z']
    if search == 'CAT29-':
        expected.insert(0, 'CAT29-ROOT')
    assert [row[0] for row in rows[1:]] == expected
    tbody = text.split('<tbody>')[1].split('</tbody>')[0]
    assert all(tbody.index(a) < tbody.index(b) for a, b in zip(expected, expected[1:]))
    assert '2 级' in tbody
