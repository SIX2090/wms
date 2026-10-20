"""BUG-2026-10-20-003: statistics and filtered document printing contracts."""
import io
import json
import shutil
import subprocess
from datetime import date
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pytest
from openpyxl import load_workbook

import app as m


@pytest.fixture
def exports(monkeypatch):
    monkeypatch.setitem(m.app.config, 'TESTING', True)
    monkeypatch.setitem(m.app.config, 'WTF_CSRF_ENABLED', False)
    with m.app.app_context():
        m.db.session.remove()
        m.db.drop_all()
        m.db.create_all()
        user = m.User(username='export-audit', role='admin', password_hash='fixture',
                      must_change_password=False)
        a = m.Warehouse(code='WHA', name='仓库A', status='active', is_default=True)
        b = m.Warehouse(code='WHB', name='仓库B', status='active')
        target = m.Material(code='TARGET', name='轴承', brand='鑫牌', stock=0)
        other = m.Material(code='OTHER', name='螺母', stock=0)
        supplier = m.Supplier(code='SUP-X', name='鑫达供货')
        second = m.Supplier(code='SUP-Y', name='其他供应商')
        department = m.Department(code='D1', name='生产车间')
        m.db.session.add_all([user, a, b, target, other, supplier, second, department])
        m.db.session.flush()
        for prefix, model, item_model, fk in (
            ('IN', m.InOrder, m.InOrderItem, 'in_order_id'),
            ('OUT', m.OutOrder, m.OutOrderItem, 'out_order_id'),
        ):
            for suffix, warehouse, status, when, qty in (
                ('C', a.name, 'completed', date(2026, 10, 10), 3),
                ('D', a.name, 'pending', date(2026, 10, 10), 7),
                ('CODE', a.code, 'completed', date(2026, 10, 10), 2),
                ('B', b.name, 'completed', date(2026, 10, 10), 100),
                ('UNKNOWN', '历史失效仓', 'completed', date(2026, 10, 10), 1000),
                ('EMPTY', '', 'completed', date(2026, 10, 10), 1000),
                ('OLD', a.name, 'completed', date(2026, 9, 1), 11),
            ):
                fields = dict(order_no=f'{prefix}-{suffix}', warehouse=warehouse,
                              status=status, date=when, contract_no='C1',
                              project_name='工程一', purpose='备料')
                if prefix == 'IN':
                    fields.update(supplier_id=(second.id if suffix == 'OLD' else supplier.id),
                                  business_type=('其他入库' if suffix == 'OLD' else '采购入库'))
                else:
                    fields.update(department_id=department.id, customer='生产车间',
                                  business_type='领料单')
                order = model(**fields)
                m.db.session.add(order)
                m.db.session.flush()
                m.db.session.add(item_model(**{fk: order.id}, material_id=target.id,
                    quantity=qty, price=2, amount=qty * 2,
                    contract_no=None if suffix == 'CODE' else 'C1',
                    project_name=None if suffix == 'CODE' else '工程一'))
                if suffix == 'C':
                    m.db.session.add(item_model(**{fk: order.id}, material_id=other.id,
                        quantity=9, price=10, amount=90, contract_no='C2', project_name='工程二'))
            if prefix == 'OUT':
                for suffix, business_type in [('SALE', '销售出库'), ('OTHER', '其他出库'),
                                               ('RETURN', '采购退货出库')]:
                    order = model(order_no=f'OUT-{suffix}', warehouse=a.name,
                                  status='completed', date=date(2026, 10, 10),
                                  business_type=business_type)
                    m.db.session.add(order)
                    m.db.session.flush()
                    m.db.session.add(item_model(out_order_id=order.id, material_id=target.id,
                                                quantity=50, price=2, amount=100))
        m.db.session.commit()
        ids = {'user': user.id, 'a': a.id, 'b': b.id, 'supplier': supplier.id}
    client = m.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(ids['user'])
        session['_fresh'] = True
    yield client, ids
    with m.app.app_context():
        m.db.session.remove()
        m.db.drop_all()
        m.db.create_all()


def _workbook(exports, **filters):
    client, ids = exports
    response = client.get('/report/inout/print', query_string={'warehouse_id': ids['a'], **filters})
    assert response.status_code == 200, response.get_data(as_text=True)[:500] if response.is_json else response.status_code
    return load_workbook(io.BytesIO(response.data))


def _rows(book, prefix):
    return list(book['入库统计' if prefix == 'IN' else '领料统计'].values)[1:]


@pytest.mark.parametrize('prefix', ['IN', 'OUT'])
def test_default_statistics_exclude_drafts_and_other_warehouses(exports, prefix):
    rows = _rows(_workbook(exports), prefix)
    numbers = {r[0] for r in rows}
    assert f'{prefix}-D' not in numbers
    assert f'{prefix}-B' not in numbers and f'{prefix}-UNKNOWN' not in numbers
    assert f'{prefix}-EMPTY' not in numbers
    assert f'{prefix}-CODE' in numbers
    assert all(r[9] == '已完成' for r in rows)


def test_explicit_status_and_unmatched_search_do_not_export_drafts(exports):
    book = _workbook(exports, status='completed', search='根本不存在的物料')
    assert not _rows(book, 'IN') and not _rows(book, 'OUT')


@pytest.mark.parametrize('prefix,document_type', [('IN', 'in_order'), ('OUT', 'out_order')])
def test_list_print_pending_is_explicit_and_only_exports_current_document_type(exports, prefix, document_type):
    book = _workbook(exports, document_type=document_type, status='pending')
    rows = _rows(book, prefix)
    assert [r[0] for r in rows] == [prefix + '-D']
    assert rows[0][9] == '待完成'
    assert not _rows(book, 'OUT' if prefix == 'IN' else 'IN')


@pytest.mark.parametrize('prefix,document_type', [('IN', 'in_order'), ('OUT', 'out_order')])
@pytest.mark.parametrize('filters', [dict(search='轴承'), dict(contract_no='C1'),
                                     dict(project_name='工程一')])
def test_filters_select_matching_detail_rows_and_keep_legacy_header_fallback(exports, prefix, document_type, filters):
    book = _workbook(exports, document_type=document_type, status='completed',
                     start_date='2026-10-01', **filters)
    rows = _rows(book, prefix)
    assert [(r[0], r[5]) for r in sorted(rows)] == [(prefix + '-C', 'TARGET'), (prefix + '-CODE', 'TARGET')]
    assert sum(r[7] for r in rows) == 5
    assert all(r[3] == 'C1' and r[4] == '工程一' for r in rows)


@pytest.mark.parametrize('filters', [dict(supplier_name='SUP-X'), dict(search='鑫牌'),
                                     dict(type='purchase_in')])
def test_inbound_partner_brand_and_business_type_filters(exports, filters):
    book = _workbook(exports, document_type='in_order', status='completed',
                     start_date='2026-10-01', **filters)
    rows = _rows(book, 'IN')
    assert {r[0] for r in rows} == {'IN-C', 'IN-CODE'}


def test_supplier_id_has_precedence_over_free_text(exports):
    book = _workbook(exports, document_type='in_order', supplier_id=exports[1]['supplier'],
                     supplier_name='不存在', status='completed')
    assert {r[0] for r in _rows(book, 'IN')} == {'IN-C', 'IN-CODE'}


@pytest.mark.parametrize('alias,expected', [('sale', 'SALE'), ('other_out', 'OTHER'),
                                         ('purchase_return', 'RETURN')])
def test_outbound_business_type_matches_list_scope(exports, alias, expected):
    rows = _rows(_workbook(exports, document_type='out_order', type=alias), 'OUT')
    assert [r[0] for r in rows] == ['OUT-' + expected]


def test_list_print_includes_all_filtered_pages_in_requested_order(exports):
    rows = _rows(_workbook(exports, document_type='in_order', status='completed',
                          contract_no='C1', sort='total_amount', order='asc', page=99, per_page=1), 'IN')
    assert [r[8] for r in rows] == [4, 6, 22]
    assert len(rows) == 3 and sum(r[7] for r in rows) == 16


@pytest.mark.parametrize('filters', [dict(start_date='2026-99-88'), dict(end_date='bad'),
                                     dict(document_type='unknown'), dict(status='invalid')])
def test_invalid_filters_fail_closed(exports, filters):
    response = exports[0].get('/report/inout/export', query_string={'warehouse_id': exports[1]['a'], **filters})
    assert response.status_code == 400


def test_export_alias_and_date_aliases_match_print(exports):
    client, ids = exports
    params = dict(warehouse_id=ids['a'], document_type='in_order', status='completed',
                  search='轴承', date_start='2026-10-01', date_end='2026-10-31')
    response = client.get('/report/inout/export', query_string=params)
    assert response.status_code == 200
    assert [r[0] for r in _rows(load_workbook(io.BytesIO(response.data)), 'IN')] == ['IN-CODE', 'IN-C']


def test_unfiltered_list_print_keeps_drafts_and_empty_headers(exports):
    with m.app.app_context():
        m.db.session.add(m.InOrder(order_no='IN-EMPTY-DRAFT', warehouse='仓库A',
                                  status='pending', date=date(2026, 10, 10)))
        m.db.session.commit()
    rows = _rows(_workbook(exports, document_type='in_order'), 'IN')
    by_number = {r[0]: r for r in rows}
    assert by_number['IN-D'][9] == '待完成'
    assert by_number['IN-EMPTY-DRAFT'][7] == 0
    assert by_number['IN-EMPTY-DRAFT'][9] == '待完成'


def test_warehouse_totals_are_exact_for_all_matching_details(exports):
    a_rows = _rows(_workbook(exports, status='completed'), 'IN')
    b_rows = _rows(_workbook(exports, warehouse_id=exports[1]['b'], status='completed'), 'IN')
    assert sum(r[7] for r in a_rows) == 25
    assert sum(r[7] for r in b_rows) == 100
    assert sum(r[7] for r in a_rows + b_rows) == 125


def _js_function(source, name):
    start = source.index('function ' + name + '(')
    end = source.index('\nfunction ', start + 1)
    return source[start:end]


@pytest.mark.parametrize('path,key,expected', [('/in_order', 'in_order', 'in_order'),
    ('/out_order', 'out_order', 'out_order'), ('/other_in_order', 'in_order', 'in_order'),
    ('/other_out_order', 'out_order', 'out_order')])
def test_actual_toolbar_print_passes_filters_to_backend(exports, path, key, expected):
    node = shutil.which('node')
    if not node:
        pytest.skip('Node.js required for toolbar integration')
    source = (Path(__file__).resolve().parents[1] / 'app/static/js/app.js').read_text(encoding='utf-8')
    mapping = source[source.index('var WMS_REPORT_PRINT_PARAM_MAP ='):source.index('function buildReportPrintUrl(')]
    query = '?warehouse_id=1&status=completed&search=轴承&contract_no=C1&project_name=工程一&supplier_name=鑫达&date_start=2026-10-01&sort=total_amount&order=asc&page=99&embedded=1'
    script = 'const window={location:{origin:"http://localhost",pathname:' + json.dumps(path) + ',search:' + json.dumps(query) + '},open(url){process.stdout.write(JSON.stringify(url))}};function getCurrentRecordId(){return null;}\n'
    script += mapping + _js_function(source, 'buildReportPrintUrl') + '\n' + _js_function(source, 'printCurrent')
    script += '\nprintCurrent({key:' + json.dumps(key) + ',printListUrl:"/report/inout/print"});'
    result = subprocess.run([node, '-e', script], capture_output=True, text=True, encoding='utf-8', check=True)
    url = json.loads(result.stdout)
    params = parse_qs(urlparse(url).query)
    assert params.get('document_type') == [expected]
    for name, value in [('status', 'completed'), ('search', '轴承'), ('contract_no', 'C1'),
                        ('project_name', '工程一'), ('sort', 'total_amount'), ('order', 'asc')]:
        assert params.get(name) == [value]
    assert 'page' not in params and 'embedded' not in params
    assert params['start_date'] == ['2026-10-01']
    if path.startswith('/other_'):
        assert params.get('type') == ['other_in' if key == 'in_order' else 'other_out']
    response = exports[0].get(url)
    assert response.status_code == 200
    rows = _rows(load_workbook(io.BytesIO(response.data)), 'IN' if key == 'in_order' else 'OUT')
    assert [r[0] for r in rows] == ([] if path.startswith('/other_') else [key.split('_')[0].upper() + '-CODE', key.split('_')[0].upper() + '-C'])
