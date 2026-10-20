"""BUG-2026-10-20-001: Excel warehouse input must survive draft import."""
import io
import json
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import app as wms


@pytest.fixture
def import_case(monkeypatch):
    monkeypatch.setitem(wms.app.config, 'TESTING', True)
    monkeypatch.setitem(wms.app.config, 'WTF_CSRF_ENABLED', False)
    with wms.app.app_context():
        assert wms.db.engine.url.database in (None, ':memory:')
        wms.db.drop_all()
        wms.db.create_all()
        user = wms.User(username='import-fixture', password_hash='unused',
                        role='admin', must_change_password=False)
        a = wms.Warehouse(code='WA', name='仓库A', is_default=True, status='active')
        b = wms.Warehouse(code='WB', name='仓库B', status='active')
        off = wms.Warehouse(code='OFF', name='停用仓', status='inactive')
        unit = wms.Unit(code='EA', name='个')
        wms.db.session.add_all([user, a, b, off, unit])
        wms.db.session.flush()
        material = wms.Material(code='M1', name='测试物料', spec='A',
                                unit_id=unit.id, stock=13)
        wms.db.session.add(material)
        wms.db.session.flush()
        for warehouse, quantity in [(a, 5), (b, 8)]:
            wms.db.session.add(wms.StockTransaction(
                material_id=material.id, warehouse_id=warehouse.id,
                transaction_type='opening', quantity=quantity, location=warehouse.name))
        wms.db.session.commit()
        ids = {'user': user.id, 'a': a.id, 'b': b.id, 'off': off.id}
    client = wms.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(ids['user'])
        session['_fresh'] = True
    yield client, ids
    with wms.app.app_context():
        wms.db.session.remove()


def upload(client, rows, warehouse_header='仓库'):
    book = Workbook()
    sheet = book.active
    sheet.append(['单据编号', warehouse_header, '物料编码', '物料名称', '数量', '库位'])
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    return client.post('/requisition/import', data={'file': (output, 'input.xlsx')},
                       content_type='multipart/form-data')


def assert_stock_untouched():
    assert wms.Material.query.filter_by(code='M1').one().stock == 13
    assert wms.StockTransaction.query.count() == 2
    assert sum(x.quantity for x in wms.StockTransaction.query.all()) == 13


def test_download_template_has_warehouse_and_roundtrips(import_case):
    client, _ = import_case
    response = client.get('/export/template/requisition')
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.data))
    sheet = book.active
    headers = [cell.value for cell in sheet[1]]
    assert '仓库' in headers and '库位' in headers
    values = {'单据编号': 'REQ-TEMPLATE', '物料编码': 'M1', '数量': 2, '仓库': '仓库B'}
    for cell, header in zip(sheet[2], headers):
        cell.value = values.get(header, '')
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    response = client.post('/requisition/import', data={'file': (output, 'template.xlsx')},
                           content_type='multipart/form-data')
    assert response.status_code == 200, response.get_json()
    with wms.app.app_context():
        assert wms.ProductionRequisition.query.one().warehouse == '仓库B'
        assert_stock_untouched()


def test_fallback_template_has_warehouse_columns():
    node = shutil.which('node')
    if not node:
        pytest.skip('Node.js required to execute the fallback template exporter')
    source = (Path(__file__).resolve().parents[1] / 'app/static/js/app.js').read_text(encoding='utf-8')
    function = 'function downloadFallbackTemplate(module) {' + source.split(
        'function downloadFallbackTemplate(module) {', 1)[1].split('\nfunction showImportModalForModule', 1)[0]
    script = "function csvEscape(s){return s;} function getModuleFilename(){return 'template.csv';}"
    script += 'function downloadBlob(csv){process.stdout.write(JSON.stringify(csv));}'
    script += function + "\ndownloadFallbackTemplate({key:'requisition'});"
    result = subprocess.run([node, '-e', script], text=True, encoding='utf-8',
                            capture_output=True, check=True)
    headers = json.loads(result.stdout).lstrip('\ufeff').splitlines()[0].split(',')
    assert '仓库' in headers and '库位' in headers


@pytest.mark.parametrize('kind', ['name', 'code', 'id'])
def test_import_uses_explicit_warehouse(import_case, kind):
    client, ids = import_case
    value = {'name': '仓库B', 'code': 'WB', 'id': ids['b']}[kind]
    header = '仓库ID' if kind == 'id' else '仓库'
    response = upload(client, [['REQ-B', value, 'M1', '测试物料', 2, '']], header)
    assert response.status_code == 200, response.get_json()
    with wms.app.app_context():
        order = wms.ProductionRequisition.query.one()
        assert order.warehouse == '仓库B'
        assert order.status == 'pending'
        assert len(order.items) == 1
        assert_stock_untouched()


def test_import_missing_warehouse_uses_active_default(import_case):
    client, _ = import_case
    response = upload(client, [['REQ-A', '', 'M1', '测试物料', 2, '']])
    assert response.status_code == 200
    with wms.app.app_context():
        assert wms.ProductionRequisition.query.one().warehouse == '仓库A'


@pytest.mark.parametrize('id_first', [True, False])
@pytest.mark.parametrize('id_kind', ['blank', 'valid', 'invalid'])
def test_id_and_name_columns_have_unambiguous_priority(import_case, id_first, id_kind):
    client, ids = import_case
    warehouse_id = {'blank': '', 'valid': ids['b'], 'invalid': 9999}[id_kind]
    book = Workbook()
    headers = ['单据编号', '仓库ID', '仓库', '物料编码', '数量']
    values = ['REQ-BOTH', warehouse_id, '仓库B', 'M1', 1]
    if not id_first:
        headers[1], headers[2] = headers[2], headers[1]
        values[1], values[2] = values[2], values[1]
    book.active.append(headers)
    book.active.append(values)
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    response = client.post('/requisition/import', data={'file': (output, 'both.xlsx')},
                           content_type='multipart/form-data')
    assert response.status_code == (400 if id_kind == 'invalid' else 200), response.get_json()
    with wms.app.app_context():
        if id_kind == 'invalid':
            assert wms.ProductionRequisition.query.count() == 0
        else:
            assert wms.ProductionRequisition.query.one().warehouse == '仓库B'


@pytest.mark.parametrize('value', ['不存在', 'OFF'])
def test_invalid_explicit_warehouse_rejects_entire_import(import_case, value):
    client, _ = import_case
    response = upload(client, [
        ['REQ-GOOD', 'WB', 'M1', '测试物料', 2, ''],
        ['REQ-BAD', value, 'NEW-MATERIAL', '新物料', 1, ''],
    ])
    assert response.status_code == 400, response.get_json()
    assert '仓库' in response.get_json()['msg']
    with wms.app.app_context():
        assert wms.ProductionRequisition.query.count() == 0
        assert wms.Material.query.count() == 1
        assert_stock_untouched()


@pytest.mark.parametrize('inactive', [False, True])
def test_no_active_default_rejects_blank_warehouse(import_case, inactive):
    client, ids = import_case
    with wms.app.app_context():
        warehouse = wms.db.session.get(wms.Warehouse, ids['a'])
        warehouse.is_default = inactive
        if inactive:
            warehouse.status = 'inactive'
        wms.db.session.commit()
    response = upload(client, [['REQ-A', '', 'M1', '测试物料', 2, '']])
    assert response.status_code == 400, response.get_json()
    with wms.app.app_context():
        assert wms.ProductionRequisition.query.count() == 0


def test_same_document_cross_warehouse_rejects_and_rolls_back(import_case):
    client, _ = import_case
    response = upload(client, [
        ['REQ-X', 'WA', 'M1', '测试物料', 2, ''],
        ['REQ-X', 'WB', 'M1', '测试物料', 3, ''],
    ])
    assert response.status_code == 400, response.get_json()
    with wms.app.app_context():
        assert wms.ProductionRequisition.query.count() == 0
        assert_stock_untouched()


def test_same_document_blank_warehouse_inherits_selected_warehouse(import_case):
    client, _ = import_case
    response = upload(client, [
        ['REQ-X', 'WB', 'M1', '测试物料', 2, ''],
        ['REQ-X', '', 'M1', '测试物料', 3, ''],
    ])
    assert response.status_code == 200, response.get_json()
    with wms.app.app_context():
        order = wms.ProductionRequisition.query.one()
        assert order.warehouse == '仓库B'
        assert sum(item.quantity for item in order.items) == 5
        assert_stock_untouched()


def test_two_warehouses_and_legacy_stock_rows_are_not_rewritten(import_case):
    client, _ = import_case
    with wms.app.app_context():
        material = wms.Material.query.filter_by(code='M1').one()
        legacy = wms.StockTransaction(material_id=material.id, warehouse_id=None,
                                      transaction_type='in', quantity=0, location='')
        wms.db.session.add(legacy)
        wms.db.session.commit()
        legacy_id = legacy.id
    response = upload(client, [
        ['REQ-A', 'WA', 'M1', '测试物料', 2, ''],
        ['REQ-B', 'WB', 'M1', '测试物料', 3, ''],
    ])
    assert response.status_code == 200, response.get_json()
    with wms.app.app_context():
        orders = wms.ProductionRequisition.query.order_by(wms.ProductionRequisition.req_no).all()
        assert [order.warehouse for order in orders] == ['仓库A', '仓库B']
        assert sum(item.quantity for order in orders for item in order.items) == 5
        assert all(order.status == 'pending' for order in orders)
        assert wms.StockTransaction.query.count() == 3
        assert wms.db.session.get(wms.StockTransaction, legacy_id).warehouse_id is None
        assert wms.Material.query.filter_by(code='M1').one().stock == 13


@pytest.mark.parametrize('location', ['', 'L-B'])
def test_location_required_when_enabled(import_case, location):
    client, _ = import_case
    with wms.app.app_context():
        wms.set_system_setting('location_management_enabled', '1')
        wms.db.session.commit()
    response = upload(client, [['REQ-B', 'WB', 'M1', '测试物料', 2, location]])
    assert response.status_code == (200 if location else 400), response.get_json()
    with wms.app.app_context():
        if location:
            assert wms.ProductionRequisition.query.one().location == location
        else:
            assert wms.ProductionRequisition.query.count() == 0
        assert_stock_untouched()
