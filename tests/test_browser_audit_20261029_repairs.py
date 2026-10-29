"""Runtime regressions for the remaining isolated-browser audit findings."""
import io
import json
import re
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

import pytest
from openpyxl import Workbook, load_workbook

import app as m
from tests.test_sales_report_print import sales_case  # noqa: F401

ROOT = Path(__file__).resolve().parents[1]


class Elements(HTMLParser):
    def __init__(self, text):
        super().__init__()
        self.elements = []
        self.feed(text)

    def handle_starttag(self, tag, attrs):
        self.elements.append((tag, dict(attrs)))


def excel(headers, rows):
    workbook = Workbook()
    workbook.active.append(headers)
    for row in rows:
        workbook.active.append(row)
    result = io.BytesIO()
    workbook.save(result)
    result.seek(0)
    return result


@pytest.fixture
def audit_case(monkeypatch):
    monkeypatch.setitem(m.app.config, 'TESTING', True)
    monkeypatch.setitem(m.app.config, 'WTF_CSRF_ENABLED', False)
    with m.app.app_context():
        m.db.session.remove()
        m.db.drop_all()
        m.db.create_all()
        user = m.User(username='audit29', role='admin', password_hash='fixture-only',
                      must_change_password=False)
        first = m.Warehouse(code='A29', name='验收甲仓', status='active', is_default=True)
        second = m.Warehouse(code='B29', name='验收乙仓', status='active')
        unit = m.Unit(code='U29', name='件')
        supplier = m.Supplier(code='S29', name='隔离加工厂')
        department = m.Department(code='D29', name='验收部门', status='active')
        m.db.session.add_all([user, first, second, unit, supplier, department])
        m.db.session.flush()
        material = m.Material(code='M29', name='验收产品', unit_id=unit.id, stock=12, price=3)
        m.db.session.add(material)
        m.db.session.flush()
        order = m.SubcontractOrder(order_no='SC' + datetime.now().strftime('%y%m') + '00001',
            supplier_id=supplier.id, warehouse=second.name, status='pending',
            contact=None, phone=None, deadline=None)
        m.db.session.add(order)
        m.db.session.flush()
        m.db.session.add(m.SubcontractItem(subcontract_order_id=order.id,
            material_id=material.id, quantity=2, unit_id=unit.id))
        m.set_system_setting('location_management_enabled', '0')
        m.db.session.commit()
        ids = dict(user=user.id, first=first.id, second=second.id, unit=unit.id,
                   supplier=supplier.id, material=material.id, order=order.id,
                   department=department.id)
    client = m.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'], session['_fresh'] = str(ids['user']), True
    yield client, ids
    with m.app.app_context():
        m.db.session.remove()


def test_second_subcontract_can_be_created(audit_case):
    client, ids = audit_case
    response = client.post('/subcontract/add', data={
        'supplier_id': ids['supplier'], 'warehouse_id': ids['second']})
    assert response.status_code == 200, response.get_json()
    with m.app.app_context():
        orders = m.SubcontractOrder.query.order_by(m.SubcontractOrder.id).all()
        assert len(orders) == 2 and orders[-1].order_no.endswith('00002')
        assert orders[-1].warehouse == '验收乙仓'


def test_repeated_copy_keeps_warehouse_unit_and_unique_number(audit_case):
    client, ids = audit_case
    for _ in range(2):
        response = client.post('/subcontract/%s/copy' % ids['order'])
        assert response.status_code == 200, response.get_json()
    with m.app.app_context():
        orders = m.SubcontractOrder.query.order_by(m.SubcontractOrder.id).all()
        assert len({order.order_no for order in orders}) == 3
        assert all(order.warehouse == '验收乙仓' for order in orders)
        assert all(order.items[0].unit_id == ids['unit'] for order in orders)
        assert m.StockTransaction.query.count() == 0


@pytest.mark.parametrize('warehouse', ['验收乙仓', 'B29', None])
def test_whole_subcontract_import_has_explicit_or_default_warehouse(audit_case, warehouse):
    client, _ = audit_case
    headers = ['委外单号', '加工厂商', '物料编码', '数量']
    row = ['SC-IMPORT29', '隔离加工厂', 'M29', 2]
    if warehouse:
        headers.append('仓库')
        row.append(warehouse)
    response = client.post('/subcontract/import', data={
        'file': (excel(headers, [row]), 'audit.xlsx')})
    assert response.status_code == 200, response.get_json()
    with m.app.app_context():
        order = m.SubcontractOrder.query.filter_by(order_no='SC-IMPORT29').one()
        assert order.warehouse == ('验收乙仓' if warehouse else '验收甲仓')
        assert m.StockTransaction.query.count() == 0


def test_bad_import_warehouse_creates_nothing(audit_case):
    client, _ = audit_case
    response = client.post('/subcontract/import', data={'file': (excel(
        ['委外单号', '加工厂商', '物料编码', '数量', '仓库'],
        [['SC-BAD29', '隔离加工厂', 'M29', 2, '不存在的仓库']]), 'audit.xlsx')})
    assert response.status_code == 400 or response.get_json().get('count') == 0
    with m.app.app_context():
        assert m.SubcontractOrder.query.count() == 1


def test_empty_legacy_warehouse_does_not_display_global_stock(audit_case):
    client, ids = audit_case
    with m.app.app_context():
        m.db.session.get(m.SubcontractOrder, ids['order']).warehouse = ''
        m.db.session.commit()
    response = client.get('/subcontract/%s?embedded=1' % ids['order'])
    text = response.get_data(as_text=True)
    materials = re.search(r'const subcontractMaterials\s*=\s*(\[.*?\]);', text)
    assert materials, 'rendered material data must be inspectable'
    assert json.loads(materials[1])[0]['stock'] == 0


@pytest.mark.parametrize('warehouse_name,expected', [('验收甲仓', 4), ('B29', 8), ('已删除仓库', 0)])
def test_subcontract_stock_isolated_even_with_legacy_warehouse_values(audit_case, warehouse_name, expected):
    client, ids = audit_case
    with m.app.app_context():
        m.db.session.get(m.SubcontractOrder, ids['order']).warehouse = warehouse_name
        for key, quantity in [('first', 4), ('second', 8)]:
            m.db.session.add(m.StockTransaction(material_id=ids['material'], warehouse_id=ids[key],
                quantity=quantity, transaction_type='in'))
        m.db.session.commit()
    text = client.get('/subcontract/%s?embedded=1' % ids['order']).get_data(as_text=True)
    materials = re.search(r'const subcontractMaterials\s*=\s*(\[.*?\]);', text)
    assert json.loads(materials[1])[0]['stock'] == expected


@pytest.mark.parametrize('kind', ['issue', 'receive'])
@pytest.mark.parametrize('explicit,expected', [('', '验收乙仓'), ('A29', '验收甲仓'), ('错误仓', None)])
def test_resolve_subcontract_import_warehouse(audit_case, kind, explicit, expected):
    client, ids = audit_case
    with m.app.app_context():
        number = m.db.session.get(m.SubcontractOrder, ids['order']).order_no
    response = client.post('/subcontract/%s/import' % kind, data={'file': (excel(
        ['单据编号', '委外单号', '物料编码', '数量', '仓库'],
        [['IMPORT29', number, 'M29', 2, explicit]]), 'audit.xlsx')})
    assert response.status_code == (200 if expected else 400), response.get_json()
    with m.app.app_context():
        model = m.SubcontractIssue if kind == 'issue' else m.SubcontractReceive
        rows = model.query.all()
        assert len(rows) == (1 if expected else 0)
        if expected:
            assert rows[0].warehouse == expected
        assert m.StockTransaction.query.count() == 0


@pytest.mark.parametrize('path', ['/subcontract/import', '/subcontract/issue/import', '/subcontract/receive/import'])
def test_no_default_warehouse_rejects_file_without_warehouse(audit_case, path):
    client, _ = audit_case
    with m.app.app_context():
        m.Warehouse.query.update({'is_default': False})
        m.db.session.commit()
    response = client.post(path, data={'file': (excel(
        ['单据编号', '加工厂商', '物料编码', '数量'],
        [['NO-WH29', '隔离加工厂', 'M29', 2]]), 'audit.xlsx')})
    assert response.status_code == 400
    with m.app.app_context():
        assert m.SubcontractOrder.query.count() == 1
        assert m.SubcontractIssue.query.count() == m.SubcontractReceive.query.count() == 0


@pytest.mark.parametrize('path', ['/subcontract', '/subcontract_issue', '/subcontract_receive'])
def test_subcontract_file_templates_and_exports_include_warehouse(audit_case, path):
    client, _ = audit_case
    for suffix in ['/download_template', '/export']:
        response = client.get(path + suffix)
        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.data)).active
        assert '仓库' in [cell.value for cell in sheet[1]]


def test_pending_order_can_repair_warehouse_but_not_after_issue(audit_case):
    client, ids = audit_case
    response = client.post('/subcontract/%s/edit' % ids['order'], json={
        'warehouse_id': ids['first']})
    assert response.status_code == 200
    with m.app.app_context():
        order = m.db.session.get(m.SubcontractOrder, ids['order'])
        assert order.warehouse == '验收甲仓'
        m.db.session.add(m.SubcontractIssue(issue_no='SF29', subcontract_order_id=order.id,
            warehouse='验收甲仓', status='pending'))
        m.db.session.commit()
    response = client.post('/subcontract/%s/edit' % ids['order'], json={
        'warehouse_id': ids['second']})
    assert response.status_code == 400
    with m.app.app_context():
        assert m.db.session.get(m.SubcontractOrder, ids['order']).warehouse == '验收甲仓'


@pytest.mark.parametrize('quantity,code,status', [(3, 'M29', 200), (-1, 'M29', 400),
    ('abc', 'M29', 400), (float('inf'), 'M29', 400), (1, 'MISSING29', 400)])
def test_product_file_import_is_atomic_and_never_posts_stock(audit_case, quantity, code, status):
    client, ids = audit_case
    response = client.post('/subcontract/%s/items/import' % ids['order'], data={
        'file': (excel(['产品编码', '产品名称', '单位', '数量'],
                       [[code, '验收产品', '件', quantity]]), 'audit.xlsx')})
    assert response.status_code == status, response.get_json()
    with m.app.app_context():
        assert m.SubcontractItem.query.count() == (2 if status == 200 else 1)
        assert m.StockTransaction.query.count() == 0
        assert m.db.session.get(m.Material, ids['material']).stock == 12


def test_product_file_rejects_completed_order_and_mixed_invalid_rows(audit_case):
    client, ids = audit_case
    for state in ['pending', 'completed']:
        with m.app.app_context():
            m.db.session.get(m.SubcontractOrder, ids['order']).status = state
            m.db.session.commit()
        response = client.post('/subcontract/%s/items/import' % ids['order'], data={
            'file': (excel(['产品编码', '数量'], [['M29', 2], ['MISSING29', 1]]), 'audit.xlsx')})
        assert response.status_code == 400
        with m.app.app_context():
            assert m.SubcontractItem.query.count() == 1


def test_product_file_authorization_and_validation(audit_case):
    client, ids = audit_case
    path = '/subcontract/%s/items/import' % ids['order']
    assert m.app.test_client().post(path).status_code in (302, 401)
    assert client.post(path).status_code == 400
    for headers, row in [(['产品名称'], ['验收产品']),
                         (['产品编码', '数量', '单位'], ['M29', 2, '吨'])]:
        response = client.post(path, data={'file': (excel(headers, [row]), 'audit.xlsx')})
        assert response.status_code == 400
    with m.app.app_context():
        m.db.session.get(m.User, ids['user']).role = 'viewer'
        m.db.session.commit()
    assert client.post(path, headers={'Accept': 'application/json'}).status_code == 403
    with m.app.app_context():
        assert m.SubcontractItem.query.count() == 1


def test_subcontract_empty_values_are_not_python_none(audit_case):
    client, _ = audit_case
    text = client.get('/subcontract?embedded=1').get_data(as_text=True)
    assert not re.search(r'<td>\s*None\s*</td>', text)
    assert 'name="warehouse_id"' in text


def test_new_tab_document_print_return_has_a_real_destination(audit_case):
    client, ids = audit_case
    text = client.get('/subcontract/%s/print' % ids['order']).get_data(as_text=True)
    links = [attrs.get('href') for tag, attrs in Elements(text).elements if tag == 'a']
    assert '/subcontract/%s' % ids['order'] in links
    assert 'history.back()' not in text


def test_department_bad_headers_rejected_before_writes(audit_case):
    client, _ = audit_case
    response = client.post('/department/import', data={'file': (excel(
        ['未知列A', '未知列B', '未知列C', '未知列D'],
        [['D-BAD29', '错误表头', 'active', '']]), 'audit.xlsx')})
    assert response.status_code == 400
    with m.app.app_context():
        assert m.Department.query.count() == 1


@pytest.mark.parametrize('status,expected', [('停用', 'inactive'), ('启用', 'active'),
                                           ('INACTIVE', 'inactive')])
def test_department_import_maps_reordered_header_and_status(audit_case, status, expected):
    client, _ = audit_case
    response = client.post('/department/import', data={'file': (excel(
        ['备注', '状态', '部门名称', '部门编码'],
        [['验收', status, '状态验收部门', 'D-NEW29']]), 'audit.xlsx')})
    assert response.status_code == 200, response.get_json()
    with m.app.app_context():
        row = m.Department.query.filter_by(code='D-NEW29').one()
        assert row.status == expected and row.name == '状态验收部门'


def test_legacy_chinese_department_status_remains_filterable(audit_case):
    client, _ = audit_case
    with m.app.app_context():
        m.db.session.add(m.Department(code='D-LEGACY29', name='旧停用部门', status='停用'))
        m.db.session.commit()
    text = client.get('/department?status=inactive&embedded=1').get_data(as_text=True)
    assert 'D-LEGACY29' in text


def test_legacy_active_department_has_consistent_display_edit_and_options(audit_case):
    client, _ = audit_case
    with m.app.app_context():
        row = m.Department(code='D-OLD29', name='旧启用部门', status='启用')
        m.db.session.add(row)
        m.db.session.commit()
        identifier = row.id
    text = client.get('/department?status=active&embedded=1&search=D-OLD29').get_data(as_text=True)
    assert 'D-OLD29' in text and 'badge bg-success' in text
    assert client.get('/department/%s' % identifier).get_json()['department']['status'] == 'active'
    for path in ['/api/departments', '/department/api/list']:
        assert 'D-OLD29' in str(client.get(path).get_json())


def test_unknown_employee_department_does_not_clear_existing_link(audit_case):
    client, ids = audit_case
    with m.app.app_context():
        m.db.session.add(m.Employee(code='E29', name='原员工', department_id=ids['department']))
        m.db.session.commit()
    response = client.post('/employee/import', data={'file': (excel(
        ['员工编码', '姓名', '部门编码'], [['E29', '错误更新', 'D-NOT-EXIST29']]), 'audit.xlsx')})
    assert response.status_code == 400
    assert '部门编码' in response.get_json()['msg']
    with m.app.app_context():
        employee = m.Employee.query.filter_by(code='E29').one()
        assert employee.name == '原员工' and employee.department_id == ids['department']


def test_broken_employee_workbook_has_readable_error(audit_case):
    client, _ = audit_case
    response = client.post('/employee/import', data={'file': (io.BytesIO(b'not-a-workbook'), 'bad.xlsx')})
    assert response.status_code == 400
    message = response.get_json()['msg']
    assert 'Excel' in message and 'zip' not in message


def test_sales_drill_links_keep_all_filters(sales_case):
    client, ids = sales_case['client'], sales_case['ids']
    response = client.get('/sales/report', query_string=dict(warehouse_id=ids['warehouse'],
        date_start=date.today().replace(day=1).isoformat(), date_end=date.today().isoformat(),
        status='confirmed', shipment_status='partial', project_no='P-A', salesperson_id=ids['employee']))
    links = [a['href'] for tag, a in Elements(response.get_data(as_text=True)).elements
             if tag == 'a' and '/sales/report?' in a.get('href', '')]
    drills = [parse_qs(urlsplit(url).query) for url in links
              if 'material_code=' in url or 'customer_id=' in url]
    assert len(drills) >= 2
    for query in drills:
        assert query['warehouse_id'] == [str(ids['warehouse'])]
        assert query['status'] == ['confirmed']
        assert query['shipment_status'] == ['partial']
        assert query['project_no'] == ['P-A']


@pytest.mark.parametrize('scope', ['customer', 'material', 'salesperson', 'detail'])
def test_sales_prints_show_warehouse_and_period(sales_case, scope):
    query = dict(warehouse_id=sales_case['ids']['warehouse'], scope=scope,
                 date_start='2020-01-01', date_end='2099-12-31')
    client = sales_case['client']
    text = client.get('/sales/summary/print_preview', query_string=query).get_data(as_text=True)
    assert '销售甲仓' in text and '2020-01-01' in text and '2099-12-31' in text
    response = client.get('/sales/summary/print_excel', query_string=query)
    workbook = load_workbook(io.BytesIO(response.data))
    assert '销售甲仓' in str(workbook.active.oddHeader.center.text)
    assert '2020-01-01' in str(workbook.active.oddHeader.center.text)


@pytest.mark.parametrize('kind', ['reconciliation', 'outflow', 'trend', 'execution', 'price_analysis'])
def test_other_sales_print_scope_matches_actual_collector(sales_case, kind):
    params = dict(warehouse_id=sales_case['ids']['warehouse'], date_start='2020-01-01', date_end='2099-12-31', months=6)
    client = sales_case['client']
    text = client.get('/sales/%s/print_preview' % kind, query_string=params).get_data(as_text=True)
    assert '销售甲仓' in text and '销售乙仓' not in text
    response = client.get('/sales/%s/print_excel' % kind, query_string=params)
    header = load_workbook(io.BytesIO(response.data)).active.oddHeader.center.text
    assert '销售甲仓' in header
    if kind in ('reconciliation', 'trend'):
        assert '2020-01-01' not in header, 'unused date parameters must not mislabel the printed range'
    else:
        assert '2020-01-01' in header and '2099-12-31' in header


def test_sales_scope_preserves_custom_headers_and_template_file(sales_case):
    from services.sales_report_print import render_sales_report_excel
    path = sales_case['root'] / 'header.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = '{item.name}'
    workbook.active.oddHeader.center.text = '自定义页眉'
    workbook.save(path)
    original = path.read_bytes()
    result = render_sales_report_excel('验收', [{'title': '名称', 'field': 'name'}],
        [{'name': '样品'}], template_path=path, print_scope='仓库：A&B')
    header = load_workbook(result).active.oddHeader.center.text
    assert '自定义页眉' in header and 'A&&B' in header
    assert path.read_bytes() == original


@pytest.mark.parametrize('path', ['/sales/report', '/sales/execution_report'])
def test_price_entry_hidden_when_price_dimension_disabled(sales_case, path):
    with m.app.app_context():
        m.set_system_setting('price_dimension_enabled', '0')
        m.db.session.commit()
    text = sales_case['client'].get(path, query_string={
        'warehouse_id': sales_case['ids']['warehouse']}).get_data(as_text=True)
    assert 'href="/sales/price_analysis"' not in text


def test_all_report_workbench_script_urls_have_version(sales_case):
    client, ids = sales_case['client'], sales_case['ids']
    for path in ['/sales/report', '/sales/execution_report', '/sales/trend_report',
                 '/sales/outflow_report', '/sales/reconciliation', '/stock_query']:
        text = client.get(path, query_string={'warehouse_id': ids['warehouse']}).get_data(as_text=True)
        scripts = [a['src'] for tag, a in Elements(text).elements if tag == 'script'
                   and 'report-workbench.js' in a.get('src', '')]
        assert scripts, path
        assert all(parse_qs(urlsplit(src).query).get('v') for src in scripts), scripts


def test_bom_selection_label_is_readable_and_share_is_not_advertised(bom_case):
    client, _ = bom_case
    text = client.get('/bom?embedded=1').get_data(as_text=True)
    headers = [a for tag, a in Elements(text).elements if tag == 'th']
    assert headers[0].get('data-default-label') == '选择'
    assert '智能分享' not in text
