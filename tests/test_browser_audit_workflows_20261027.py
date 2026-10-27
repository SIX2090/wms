"""Second browser-audit batch: real import files and warehouse boundaries."""
from io import BytesIO

from flask import template_rendered
from openpyxl import Workbook, load_workbook
import pytest


def _xlsx(headers, rows):
    wb = Workbook()
    wb.active.append(headers)
    for row in rows:
        wb.active.append(row)
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _warehouses():
    from app import db, Warehouse
    a = Warehouse(code='WA', name='导入A仓', status='active', is_default=True)
    b = Warehouse(code='WB', name='导入B仓', status='active')
    db.session.add_all([a, b])
    db.session.commit()
    return a, b


def test_outbound_official_template_contains_warehouse(bom_case):
    client, _ = bom_case
    response = client.get('/export/template/out_order')
    assert response.status_code == 200
    headers = next(load_workbook(BytesIO(response.data)).active.values)
    assert '仓库' in headers
    assert '库位' in headers


def test_legacy_outbound_template_uses_default_warehouse(bom_case):
    from app import OutOrder, StockTransaction
    client, _ = bom_case
    a, _ = _warehouses()
    r = client.post('/import/out_order', data={'file':(_xlsx(
        ['单据编号','物料编码','数量','单价'], [['OLD-WH','B-M',1,0]]), 'legacy.xlsx')})
    assert r.status_code == 200
    assert OutOrder.query.one().warehouse == a.name
    assert OutOrder.query.one().status == 'pending'
    assert StockTransaction.query.count() == 0


def test_outbound_file_can_choose_other_warehouse_and_picker(bom_case):
    from app import OutOrder
    client, _ = bom_case
    _, b = _warehouses()
    r = client.post('/import/out_order', data={'file':(_xlsx(
        ['单据编号','物料编码','数量','仓库','领料人'], [['OTHER-WH','B-M',1,b.code,'验收领用人']]), 'warehouse.xlsx')})
    assert r.status_code == 200
    order = OutOrder.query.one()
    assert order.warehouse == b.name
    assert order.picker == '验收领用人'


def test_outbound_missing_default_does_not_create_invisible_draft(bom_case):
    from app import OutOrder
    client, _ = bom_case
    response = client.post('/import/out_order', data={'file':(_xlsx(
        ['单据编号','物料编码','数量'], [['NO-WH','B-M',1]]), 'missing.xlsx')})
    assert OutOrder.query.count() == 0
    assert '仓库' in str(response.get_json())


def test_inbound_detail_excel_appends_to_current_draft_only(bom_case):
    from app import db, InOrder, InOrderItem, Material, StockTransaction
    client, ids = bom_case
    a, _ = _warehouses()
    order = InOrder(order_no='CURRENT-DRAFT', business_type='其他入库', warehouse=a.name, status='pending')
    db.session.add(order)
    db.session.flush()
    db.session.add(InOrderItem(in_order_id=order.id, material_id=ids['material'],quantity=1,price=2,amount=2))
    db.session.commit()
    response = client.post(f'/in_order/{order.id}/batch_add_items', data={'file':(_xlsx(
        ['物料编码','数量','单价','备注'], [['B-M',2,3,'追加明细']]), 'details.xlsx')})
    assert response.status_code == 200, response.get_json()
    assert response.get_json()['count'] == 1
    assert InOrder.query.count() == 1
    assert sum(i.quantity for i in order.items) == 3
    assert db.session.get(Material, ids['material']).stock == 10
    assert StockTransaction.query.count() == 0


def test_material_cache_is_warehouse_scoped(bom_case):
    from app import db, StockTransaction
    client, ids = bom_case
    a, b = _warehouses()
    db.session.add_all([StockTransaction(material_id=ids['material'], warehouse_id=a.id,
        location=a.name,transaction_type='in',quantity=6),
        StockTransaction(material_id=ids['material'], warehouse_id=b.id,
        location=b.name,transaction_type='in',quantity=4)])
    db.session.commit()
    for wh, qty in ((a,6),(b,4)):
        result = client.get('/material/api/all', query_string={'warehouse':wh.name}).get_json()
        assert result['materials'][0]['stock'] == qty


@pytest.mark.parametrize('kind', ['in_order', 'out_order', 'after_sale_out'])
def test_document_details_display_only_their_warehouse_stock(bom_case, kind):
    from app import (db, StockTransaction, InOrder, InOrderItem, OutOrder, OutOrderItem,
                     AfterSaleOutOrder, AfterSaleOutOrderItem)
    from tests.test_browser_audit_repairs_20261027 import _context
    client, ids = bom_case
    a, b = _warehouses()
    db.session.add_all([StockTransaction(material_id=ids['material'],warehouse_id=a.id,
        location=a.name,transaction_type='in',quantity=6),
        StockTransaction(material_id=ids['material'],warehouse_id=b.id,
        location=b.name,transaction_type='in',quantity=4)])
    model, item_model, foreign_key = {
        'in_order':(InOrder,InOrderItem,'in_order_id'),
        'out_order':(OutOrder,OutOrderItem,'out_order_id'),
        'after_sale_out':(AfterSaleOutOrder,AfterSaleOutOrderItem,'after_sale_out_order_id'),
    }[kind]
    order = model(order_no='DETAIL-STOCK', warehouse=a.name, status='pending')
    db.session.add(order)
    db.session.flush()
    db.session.add(item_model(**{foreign_key:order.id}, material_id=ids['material'],quantity=1,price=1,amount=1))
    db.session.commit()
    _, context = _context(client, f'/{kind}/{order.id}?embedded=1')
    assert context['warehouse_stock'][ids['material']] == 6


def test_new_check_ignores_client_global_system_stock(bom_case):
    from app import db, StockTransaction, InventoryCheckItem
    client, ids = bom_case
    a, b = _warehouses()
    db.session.add_all([StockTransaction(material_id=ids['material'], warehouse_id=a.id,
        location=a.name,transaction_type='in',quantity=6),
        StockTransaction(material_id=ids['material'], warehouse_id=b.id,
        location=b.name,transaction_type='in',quantity=4)])
    db.session.commit()
    r = client.post('/check/save_table', json={'header':{'warehouse':a.name},
        'items':[{'material_id':ids['material'],'actual_stock':6,'system_stock':10}]})
    assert r.status_code == 200
    item = InventoryCheckItem.query.one()
    assert item.system_stock == 6 and item.difference == 0


def test_after_sale_source_options_searches_and_paginates(bom_case):
    from app import db, SalesOrder, Customer
    client, _ = bom_case
    customer = Customer(code='SOURCE-CUSTOMER', name='来源客户')
    db.session.add(customer)
    db.session.flush()
    db.session.add_all([SalesOrder(order_no=f'SOURCE-{i:02d}',status='confirmed',customer_id=customer.id) for i in range(23)])
    db.session.commit()
    response = client.get('/after_sale_out/source_options?type=sales&page=2&search=SOURCE')
    assert response.status_code == 200
    data = response.get_json()
    assert data['total'] == 23 and data['total_pages'] == 2 and data['page'] == 2
    assert len(data['items']) == 3
