# -*- coding: utf-8 -*-
"""BUG-2026-09-05-006 回归：盘点单导出补「盘点人/盘点时间」列。

问题：盘点单详情页早有行级盘点归属回查（INV-BATCH-001-D，counted_by/
counted_at 列），但 /check/<id>/export 导出只有账面/实盘/差异——想在
Excel 里筛"哪些行还没盘"（漏盘核对）做不到；且末列表头写「备注」，
实际写入的一直是 item.reason（差异原因），表头与内容不符。

修复：导出表头补「盘点人」「盘点时间」，末列正名「差异原因」（与导入
别名对齐，导出文件可直接回导）。

覆盖：
T1. 表头 12 列含差异原因/盘点人/盘点时间，无「备注」
T2. 已盘行导出盘点人用户名与时间；未盘行两列为空（= 漏盘行可筛）
T3. 差异原因列导出 item.reason 内容
"""
from __future__ import annotations

import io
import os
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
APP_DIR = ROOT / "app"
sys.path.insert(0, str(APP_DIR))
os.chdir(APP_DIR)

os.environ.setdefault("WMS_BOOTSTRAP_PASSWORD", "admin")
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["WMS_DATABASE_URI"] = "sqlite:///:memory:"
os.environ.setdefault("WMS_DEBUG", "0")

import app as app_module  # noqa: E402
from app import db  # noqa: E402

app_module.app.config["TESTING"] = True
app_module.app.config["WTF_CSRF_ENABLED"] = False

_ctx = app_module.app.app_context()

import pytest as _pytest  # noqa: E402


@_pytest.fixture(autouse=True, scope="module")
def _release_app_ctx_after_module():
    _ctx.push()
    yield
    try:
        _ctx.pop()
    except Exception:
        pass


def _seed_scene():
    from werkzeug.security import generate_password_hash
    from app import InventoryCheck, InventoryCheckItem, Material, User, Warehouse
    db.drop_all()
    db.create_all()
    admin = User(username="admin", password_hash=generate_password_hash("admin"),
                 role="admin", must_change_password=False)
    counter = User(username="zhangsan", password_hash=generate_password_hash("x"),
                   role="warehouse", must_change_password=False)
    wh = Warehouse(code="WA", name="A仓", status="active", is_default=True)
    db.session.add_all([admin, counter, wh])
    db.session.commit()
    m1 = Material(code="M001", name="电线1", stock=10)
    m2 = Material(code="M002", name="电线2", stock=5)
    db.session.add_all([m1, m2])
    db.session.commit()
    check = InventoryCheck(check_no="CK-TEST-001", warehouse="A仓",
                           status="pending", operator_id=admin.id)
    db.session.add(check)
    db.session.flush()
    counted = InventoryCheckItem(
        inventory_check_id=check.id, material_id=m1.id,
        system_stock=10, actual_stock=8, difference=-2,
        reason="破损2件", counted_by=counter.id,
        counted_at=datetime(2026, 9, 5, 14, 30))
    pending = InventoryCheckItem(
        inventory_check_id=check.id, material_id=m2.id,
        system_stock=5, actual_stock=5, difference=0)
    db.session.add_all([counted, pending])
    db.session.commit()
    return check


def _export_rows(client, check_id):
    from openpyxl import load_workbook
    r = client.get(f"/check/{check_id}/export")
    assert r.status_code == 200, r.get_data(as_text=True)
    wb = load_workbook(io.BytesIO(r.data))
    return list(wb.active.iter_rows(values_only=True))


def test_export_headers_and_counted_columns():
    check = _seed_scene()
    client = app_module.app.test_client()
    r = client.post("/login", data={"username": "admin", "password": "admin"})
    assert r.status_code in (302, 303)

    rows = _export_rows(client, check.id)
    header = list(rows[0])
    # T1：12 列、含新列、无旧「备注」
    assert "差异原因" in header and "盘点人" in header and "盘点时间" in header
    assert "备注" not in header
    assert len(header) == 12

    idx_reason = header.index("差异原因")
    idx_who = header.index("盘点人")
    idx_when = header.index("盘点时间")
    idx_code = header.index("物料编码")
    by_code = {row[idx_code]: row for row in rows[1:]}

    # T2+T3：已盘行（M001）
    m1 = by_code["M001"]
    assert m1[idx_reason] == "破损2件"
    assert m1[idx_who] == "zhangsan"
    assert m1[idx_when] == "2026-09-05 14:30"

    # T2：未盘行（M002）两列为空 → Excel 可直接筛漏盘
    m2 = by_code["M002"]
    assert (m2[idx_who] or "") == ""
    assert (m2[idx_when] or "") == ""
