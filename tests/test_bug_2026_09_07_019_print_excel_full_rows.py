# -*- coding: utf-8 -*-
"""BUG-2026-09-07-019 回归：模板 Excel 打印全量行 + 导出文件内截断提示。

背景（数据完整性回归）：前端「模板 Excel 打印」（templatePrintBtn →
/report/<type>/print_excel）构造参数时显式 delete('export')，请求不带
export=excel；而 SQL 分页 builder（BUG-2026-09-07-004 起 10 类报表）只在
export == 'excel' 时返回全量行，否则返回当页——print_excel 用 payload 的
rows 渲染模板，导致模板打印从「全量」退化为「第一页（默认 20 行）」。
内存路径报表的 all_rows 一直为全量，故该回归仅在 SQL 分页报表上出现。

修复：print_excel 路由强制 filters['export'] = 'excel'（打印语义 = 当前
筛选全量，与页面分页无关）；导出超限时 Excel 文件尾部追加截断提示行
（此前截断提示只在网页展示，下载后的文件脱离系统完全无感知）。

断言：
  T1. print_excel（无 export 参数）渲染全部 25 行数据（回归核心：修复前仅 20）。
  T2. export=excel 全量 25 行，未超限时文件内无截断提示行。
  T3. 导出超上限（monkeypatch REPORT_ROW_LIMIT=10）：文件仅含前 10 行数据，
      尾部有截断提示行（含真实行数 25 与已含行数 10）。
  T4. 页面分页路径不受导出上限影响：REPORT_ROW_LIMIT=10 时 page=2 仍返回
      剩余 5 行，truncated=False（页面无截断）。
  T5. print_excel 无仓库 → 400（AGENTS.md 仓库必填规则保持）。
  T6. print_excel 首行数据为默认排序（日期倒序）的最新单 IN-E。
"""
from __future__ import annotations

import io
import os
import sys
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
APP_DIR = ROOT / "app"
sys.path.insert(0, str(APP_DIR))
os.chdir(APP_DIR)

os.environ.setdefault("WMS_BOOTSTRAP_PASSWORD", "admin")
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["WMS_DATABASE_URI"] = "sqlite:///:memory:"
os.environ.setdefault("WMS_DEBUG", "0")

import app as app_module  # noqa: E402
from app import (InOrder, InOrderItem, Material, Unit, User,  # noqa: E402
                 Warehouse, db)

app_module.app.config["TESTING"] = True
app_module.app.config["WTF_CSRF_ENABLED"] = False

D1 = date(2026, 9, 1)


def _xlsx_rows(resp_bytes):
    """解析下载的 xlsx → 非空行列表（字符串化）。"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp_bytes))
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(v is not None and str(v).strip() != '' for v in row):
            rows.append(['' if v is None else str(v) for v in row])
    return rows


def _split_data_rows(rows):
    """模板打印含标题行（如「入库明细报表」）与页脚行（「制单：」开头），
    定位「日期」表头行，返回表头行与其后的数据行（排除页脚）。"""
    for i, row in enumerate(rows):
        if row[0] == '日期':
            data = [r for r in rows[i + 1:] if not r[0].startswith('制单')]
            return row, data
    raise AssertionError(f"未找到表头行（首列『日期』），实际前 3 行：{rows[:3]}")


class TestBug20260907019:
    def setup_method(self):
        with app_module.app.app_context():
            db.drop_all()
            db.create_all()
            from werkzeug.security import generate_password_hash
            wh = Warehouse(code="WHA", name="仓库A", is_default=True, status="active")
            unit = Unit(code="PCS", name="个")
            user = User(username="admin",
                        password_hash=generate_password_hash("admin"),
                        role="admin", must_change_password=False)
            db.session.add_all([wh, unit, user])
            db.session.flush()
            material = Material(code="M001", name="电缆", spec="2.5mm",
                                unit_id=unit.id, price=2.0, stock=0.0)
            db.session.add(material)
            db.session.flush()
            self.wh_id = wh.id
            # 5 单 × 5 行 = 25 行明细（> 默认 page_size 20）；数量 1..5 便于聚合校验
            for i in range(5):
                order = InOrder(order_no=f"IN-{chr(ord('A') + i)}",
                                warehouse="仓库A", status="completed",
                                date=D1 + timedelta(days=i))
                db.session.add(order)
                db.session.flush()
                for _ in range(5):
                    db.session.add(InOrderItem(
                        in_order_id=order.id, material_id=material.id,
                        quantity=float(i + 1), price=2.0, amount=2.0 * (i + 1)))
            db.session.commit()

    def _client(self):
        client = app_module.app.test_client()
        client.post("/login", data={"username": "admin", "password": "admin"})
        return client

    def test_T1_print_excel_contains_all_25_rows(self):
        """核心回归：print_excel 不带 export 参数也渲染全部 25 行（修复前仅 20）。"""
        r = self._client().get(f"/report/in_detail/print_excel?warehouse_id={self.wh_id}")
        assert r.status_code == 200
        assert 'spreadsheetml' in (r.mimetype or ''), f"应返回 xlsx，实际 {r.mimetype}"
        rows = _xlsx_rows(r.data)
        header, data_rows = _split_data_rows(rows)
        assert header[0] == '日期', f"首行应为表头，实际 {header[:3]}"
        assert len(data_rows) == 25, (
            f"模板打印应含全部 25 行数据（修复前只有第一页 20 行），实际 {len(data_rows)}")

    def test_T2_export_excel_full_rows_no_warning(self):
        """export=excel 全量 25 行；未超限时文件内无截断提示行。"""
        r = self._client().get(
            f"/report/api/in_detail?export=excel&warehouse_id={self.wh_id}")
        assert r.status_code == 200
        rows = _xlsx_rows(r.data)
        data_rows = rows[1:]
        assert len(data_rows) == 25
        joined = '\n'.join(cell for row in rows for cell in row)
        assert '导出截断' not in joined, "未超限时不应出现截断提示行"

    def test_T3_export_excel_over_limit_has_truncation_row(self):
        """超上限：文件仅含前 10 行数据 + 尾部截断提示行（真实 25 / 已含 10）。"""
        old = app_module.REPORT_ROW_LIMIT
        app_module.REPORT_ROW_LIMIT = 10
        try:
            r = self._client().get(
                f"/report/api/in_detail?export=excel&warehouse_id={self.wh_id}")
        finally:
            app_module.REPORT_ROW_LIMIT = old
        assert r.status_code == 200
        rows = _xlsx_rows(r.data)
        warning_rows = [row for row in rows if '导出截断提示' in row[0]]
        assert len(warning_rows) == 1, "截断时文件尾必须有且仅有一条提示行"
        text = warning_rows[0][0]
        assert '25' in text, f"提示行应含真实行数 25，实际：{text}"
        assert '10' in text, f"提示行应含已导出行数 10，实际：{text}"
        data_rows = [row for row in rows[1:] if '导出截断提示' not in row[0]]
        assert len(data_rows) == 10, f"超限时数据行应截断为 10，实际 {len(data_rows)}"

    def test_T4_page_path_unaffected_by_export_limit(self):
        """页面分页不受导出上限影响：REPORT_ROW_LIMIT=10 时 page=2 仍返回 5 行。"""
        old = app_module.REPORT_ROW_LIMIT
        app_module.REPORT_ROW_LIMIT = 10
        try:
            r = self._client().get(
                f"/report/api/in_detail?warehouse_id={self.wh_id}&page=2&page_size=20")
        finally:
            app_module.REPORT_ROW_LIMIT = old
        assert r.status_code == 200
        data = r.get_json()
        assert data['status'] == 'success'
        assert data['total'] == 25
        assert len(data['data']) == 5, "页面分页第二页应返回剩余 5 行"
        assert data['truncated'] is False, "页面路径不应被导出上限标记为截断"

    def test_T5_print_excel_requires_warehouse(self):
        """无仓库（且无默认仓库）时 print_excel 仍拒绝（AGENTS.md 仓库必填规则）。"""
        with app_module.app.app_context():
            wh = db.session.get(Warehouse, self.wh_id)
            wh.is_default = False
            db.session.commit()
        r = self._client().get("/report/in_detail/print_excel")
        assert r.status_code == 400

    def test_T6_print_excel_first_row_is_latest_order(self):
        """打印行序为默认排序（日期倒序）：首行数据为最新单 IN-E。"""
        r = self._client().get(f"/report/in_detail/print_excel?warehouse_id={self.wh_id}")
        assert r.status_code == 200
        rows = _xlsx_rows(r.data)
        header, data_rows = _split_data_rows(rows)
        first = data_rows[0]
        no_idx = header.index('入库单号')
        assert first[no_idx] == 'IN-E', (
            f"首行应为日期最新的 IN-E，实际 {first[no_idx]}")
