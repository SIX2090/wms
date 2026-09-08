# -*- coding: utf-8 -*-
"""BUG-2026-09-07-020 回归：分批全量导出（突破单次 5 万行上限）。

背景：SQL 分页报表导出此前走 builder 的 export=excel 全量分支，受
REPORT_ROW_LIMIT(50000) 单次截断——超限导出缺行（019 只补了文件内提示，
「按当前筛选导出全部」仍未实现，评测报告 0.4 后半）。

修复：导出/模板打印改走分批流式——SQL 分页报表循环 builder 分页分支逐批
取回（每批 REPORT_EXPORT_BATCH_SIZE=5000），拼接 = 筛选全集，不再有 5 万
上限；write_only Workbook 逐批写、内存峰值恒为单批。内存路径报表（台账/
月报，结存/倒推为运行期值无法切片）仍受各自 LIMIT 截断并保留 019 的文件
内提示行。print_excel 同步改走分批收集（不再经过 _build_report_payload）。

断言：
  T1. 突破核心：REPORT_ROW_LIMIT=10 时 SQL 分页导出仍全量 25 行、无提示行
      （019 语义下此时会被截断为 10 行 + 提示行）。
  T2. 分批正确性：批 7（25 = 7+7+7+4 四批）导出行数 25、行序与单批一致
      （首 IN-E 末 IN-A、每单恰 5 行、无重复无遗漏）。
  T3. 模板打印同步突破：REPORT_ROW_LIMIT=10 + 批 7 时 print_excel 全量 25 行。
  T4. 空结果导出：仅表头、无数据行、无提示行、不异常。
  T5. 一致性：JSON 分页 total == 分批导出行数（导出 = 页面全集）。
"""
from __future__ import annotations

import io
import os
import sys
from collections import Counter
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
APP_DIR = ROOT / "app"
sys.path.insert(0, str(APP_DIR))
os.chdir(APP_DIR)

os.environ.setdefault("WMS_BOOTSTRAP_PASSWORD", "admin")
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["WMS_DATABASE_URI"] = "sqlite:///:memory:"
os.environ.setdefault("WMS_DEBUG", "0")

import app as app_module  # noqa: E402
from app import (InOrder, InOrderItem, Material, Unit, User,  # noqa: E402
                 Warehouse, db)

app_module.app.config["TESTING"] = True
app_module.app.config["WTF_CSRF_ENABLED"] = False

D1 = date(2026, 9, 1)


def _xlsx_rows(resp_bytes):
    """解析下载的 xlsx → 非空行列表（字符串化）。"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp_bytes))
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(v is not None and str(v).strip() != '' for v in row):
            rows.append(['' if v is None else str(v) for v in row])
    return rows


class TestBug20260907020:
    def setup_method(self):
        with app_module.app.app_context():
            db.drop_all()
            db.create_all()
            from werkzeug.security import generate_password_hash
            wh = Warehouse(code="WHA", name="仓库A", is_default=True, status="active")
            unit = Unit(code="PCS", name="个")
            user = User(username="admin",
                        password_hash=generate_password_hash("admin"),
                        role="admin", must_change_password=False)
            db.session.add_all([wh, unit, user])
            db.session.flush()
            material = Material(code="M001", name="电缆", spec="2.5mm",
                                unit_id=unit.id, price=2.0, stock=0.0)
            db.session.add(material)
            db.session.flush()
            self.wh_id = wh.id
            # 5 单 × 5 行 = 25 行明细；日期递增便于断言默认倒序行序
            for i in range(5):
                order = InOrder(order_no=f"IN-{chr(ord('A') + i)}",
                                warehouse="仓库A", status="completed",
                                date=D1 + timedelta(days=i))
                db.session.add(order)
                db.session.flush()
                for _ in range(5):
                    db.session.add(InOrderItem(
                        in_order_id=order.id, material_id=material.id,
                        quantity=float(i + 1), price=2.0, amount=2.0 * (i + 1)))
            db.session.commit()

    def _client(self):
        client = app_module.app.test_client()
        client.post("/login", data={"username": "admin", "password": "admin"})
        return client

    def test_T1_export_ignores_report_row_limit(self):
        """突破核心：REPORT_ROW_LIMIT=10 时导出仍全量 25 行、无提示行。"""
        old = app_module.REPORT_ROW_LIMIT
        app_module.REPORT_ROW_LIMIT = 10
        try:
            r = self._client().get(
                f"/report/api/in_detail?export=excel&warehouse_id={self.wh_id}")
        finally:
            app_module.REPORT_ROW_LIMIT = old
        assert r.status_code == 200
        rows = _xlsx_rows(r.data)
        data_rows = rows[1:]
        assert len(data_rows) == 25, (
            f"分批导出应突破 REPORT_ROW_LIMIT 上限全量 25 行，实际 {len(data_rows)}")
        joined = '\n'.join(cell for row in rows for cell in row)
        assert '导出截断' not in joined, "SQL 分页报表分批全量导出不应有截断提示行"

    def test_T2_batched_rows_concatenate_exactly(self):
        """分批正确：批 7 → 4 批拼 25 行，行序与单批全量一致、无重复无遗漏。"""
        old = app_module.REPORT_EXPORT_BATCH_SIZE
        app_module.REPORT_EXPORT_BATCH_SIZE = 7
        try:
            r = self._client().get(
                f"/report/api/in_detail?export=excel&warehouse_id={self.wh_id}")
        finally:
            app_module.REPORT_EXPORT_BATCH_SIZE = old
        assert r.status_code == 200
        rows = _xlsx_rows(r.data)
        data_rows = rows[1:]
        assert len(data_rows) == 25, f"分批拼接应恰好 25 行，实际 {len(data_rows)}"
        header = rows[0]
        no_idx = header.index('入库单号')
        nos = [row[no_idx] for row in data_rows]
        assert nos[0] == 'IN-E' and nos[-1] == 'IN-A', (
            f"分批拼接行序应保持默认日期倒序，实际首尾 {nos[0]}/{nos[-1]}")
        counts = Counter(nos)
        assert counts == {f"IN-{c}": 5 for c in "ABCDE"}, (
            f"每单应恰好 5 行无重复无遗漏，实际 {dict(counts)}")

    def test_T3_print_excel_also_batched_full(self):
        """模板打印同步走分批：REPORT_ROW_LIMIT=10 + 批 7 仍全量 25 行。"""
        old_limit = app_module.REPORT_ROW_LIMIT
        old_batch = app_module.REPORT_EXPORT_BATCH_SIZE
        app_module.REPORT_ROW_LIMIT = 10
        app_module.REPORT_EXPORT_BATCH_SIZE = 7
        try:
            r = self._client().get(
                f"/report/in_detail/print_excel?warehouse_id={self.wh_id}")
        finally:
            app_module.REPORT_ROW_LIMIT = old_limit
            app_module.REPORT_EXPORT_BATCH_SIZE = old_batch
        assert r.status_code == 200
        rows = _xlsx_rows(r.data)
        for i, row in enumerate(rows):
            if row[0] == '日期':
                data_rows = [x for x in rows[i + 1:] if not x[0].startswith('制单')]
                break
        else:
            raise AssertionError(f"未找到表头行，实际前 3 行：{rows[:3]}")
        assert len(data_rows) == 25, (
            f"模板打印应随分批导出全量 25 行，实际 {len(data_rows)}")

    def test_T4_empty_result_exports_header_only(self):
        """空结果：仅表头、无数据行、无提示行、不异常。"""
        r = self._client().get(
            f"/report/api/in_detail?export=excel&warehouse_id={self.wh_id}"
            f"&material_code=NOTEXIST")
        assert r.status_code == 200
        rows = _xlsx_rows(r.data)
        assert len(rows) == 1, f"空结果应仅含表头行，实际 {len(rows)} 行"
        assert '日期' in rows[0], f"唯一行应为表头，实际 {rows[0][:3]}"

    def test_T5_export_rows_equal_page_total(self):
        """一致性：JSON 分页 total == 分批导出行数（导出 = 页面全集）。"""
        j = self._client().get(
            f"/report/api/in_detail?warehouse_id={self.wh_id}").get_json()
        assert j['status'] == 'success' and j['total'] == 25
        r = self._client().get(
            f"/report/api/in_detail?export=excel&warehouse_id={self.wh_id}")
        rows = _xlsx_rows(r.data)
        assert len(rows) - 1 == j['total'], (
            f"导出行数 {len(rows) - 1} 应等于页面 total {j['total']}")
