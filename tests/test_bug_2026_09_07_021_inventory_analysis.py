# -*- coding: utf-8 -*-
"""BUG-2026-09-07-021 回归：库存经营分析报表（骨架 + ABC + 呆滞）。

背景：报表中心缺经营视角——只有流水类（出入库明细/台账）与静态库存
列表，没有按物料维度的库存状态预警、ABC 资金占用分类、呆滞识别。
本报表在 SQL_PAGED_REPORT_BUILDERS 协议内新增 inventory_analysis：
物料维度宽表（行集=筛选后物料，行数可控，与库存报表 015 同判定不适合
SQL 下沉），Python 全量构建 + 切片分页，019/020 的分批导出/模板打印
按协议自动接入。

口径（v1，用户裁定 021/022/023 合并为单原子动作，库龄/周转不做）：
- 库存状态：min_stock>0 才预警——stock<=0 缺货 / stock<=min_stock 低于
  安全库存 / 其余正常；建议补货量 = 补至 max_stock（未设则 min_stock）。
- ABC：库存金额降序累计占比（不含自身）<70%→A / <90%→B / 其余 C；
  同金额按物料 id 升序；金额合计 <=0 全 C（首项豁免：单物料 A 不是 C）。
- 呆滞：stock>0 且 idle_days >= stock_idle_days（系统设置默认 90）；
  idle_days = today - 最后出库日期（无出库按建档日期起算）；出库方向
  与 _is_inbound_transaction 严格同源（期初类型既非入也非出不参与）。

断言：
  T1. 宽表行=物料 + 仓库隔离：两仓环境行数=物料数、库存按所选仓净额。
  T2. 库存状态与建议补货：缺货/低于安全/正常、min_stock<=0 不预警。
  T3. ABC 分类：集成梯度 A..C + 纯函数边界（首项豁免/total=0 全 C）。
  T4. 呆滞：90 天边界（=90 呆滞、89 不呆滞）、stock=0 不标记、
      从未出库按建档起算、期初流水不参与出库方向。
  T5. 分页/排序/导出/模板打印接入（SQL_PAGED 协议 + 020 分批）。
  T6. hide_zero 过滤零库存行（与库存报表口径一致）。
  T7. summary 口径：count=物料数、quantity=净库存和、amount=金额和。
"""
from __future__ import annotations

import io
import os
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
APP_DIR = ROOT / "app"
sys.path.insert(0, str(APP_DIR))
os.chdir(APP_DIR)

os.environ.setdefault("WMS_BOOTSTRAP_PASSWORD", "admin")
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["WMS_DATABASE_URI"] = "sqlite:///:memory:"
os.environ.setdefault("WMS_DEBUG", "0")

import app as app_module  # noqa: E402
from app import (Material, StockTransaction, Supplier, Unit, User,  # noqa: E402
                 Warehouse, db)

app_module.app.config["TESTING"] = True
app_module.app.config["WTF_CSRF_ENABLED"] = False

TODAY = date.today()
NOW = datetime.now()


def _xlsx_rows(resp_bytes):
    """解析下载的 xlsx → 非空行列表（字符串化）。"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp_bytes))
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(v is not None and str(v).strip() != '' for v in row):
            rows.append(['' if v is None else str(v) for v in row])
    return rows


class TestBug20260907021:
    def setup_method(self):
        with app_module.app.app_context():
            db.drop_all()
            db.create_all()
            from werkzeug.security import generate_password_hash
            wh_a = Warehouse(code="WHA", name="仓库A", is_default=True,
                             status="active")
            wh_b = Warehouse(code="WHB", name="仓库B", is_default=False,
                             status="active")
            unit = Unit(code="PCS", name="个")
            user = User(username="admin",
                        password_hash=generate_password_hash("admin"),
                        role="admin", must_change_password=False)
            supplier = Supplier(code="SUP01", name="供应商一")
            db.session.add_all([wh_a, wh_b, unit, user, supplier])
            db.session.flush()
            self.wh_a_id, self.wh_b_id = wh_a.id, wh_b.id
            self.supplier_id = supplier.id

            # 物料矩阵：金额梯度 + 状态梯度 + 呆滞边界一体覆盖
            #   M1 价100 净10(仓A) 金额1000  min0        正常   出库@d-90  呆滞(=90)
            #   M2 价50  净4(仓A)  金额200   min5        低于安全 出库@d-89  否(89<90)
            #   M3 价30  净0(仓A)  金额0     min8 max20  缺货    出库@d-10  否(stock0)
            #   M4 价1   净3(仓A)  金额3     min0        正常    无出库     呆滞(建档d-100)
            #   M5 价5   净1(仓A)  金额5     min0        正常    仅期初流水 否(期初非出库)
            specs = [
                ("M1", "物料甲", 100.0, 0.0, 0.0),
                ("M2", "物料乙", 50.0, 5.0, 0.0),
                ("M3", "物料丙", 30.0, 8.0, 20.0),
                ("M4", "物料丁", 1.0, 0.0, 0.0),
                ("M5", "物料戊", 5.0, 0.0, 0.0),
            ]
            self.materials = {}
            for code, name, price, min_stock, max_stock in specs:
                m = Material(code=code, name=name, unit_id=unit.id,
                             price=price, min_stock=min_stock,
                             max_stock=max_stock, stock=0.0,
                             supplier_id=supplier.id)
                db.session.add(m)
                db.session.flush()
                self.materials[code] = m

            def _txn(code, ttype, qty, days_ago, wh_id):
                db.session.add(StockTransaction(
                    material_id=self.materials[code].id,
                    transaction_type=ttype, quantity=qty,
                    warehouse_id=wh_id,
                    created_at=NOW - timedelta(days=days_ago)))

            # 仓A 流水（净额 = 上面矩阵列）
            _txn("M1", "in", 100.0, 120, wh_a.id)
            _txn("M1", "out", -90.0, 90, wh_a.id)
            _txn("M2", "in", 104.0, 120, wh_a.id)
            _txn("M2", "out", -100.0, 89, wh_a.id)
            _txn("M3", "in", 10.0, 30, wh_a.id)
            _txn("M3", "out", -10.0, 10, wh_a.id)
            _txn("M4", "in", 3.0, 1, wh_a.id)
            _txn("M5", "opening", 1.0, 200, wh_a.id)  # 期初：非入非出
            # 仓B 仅 M1 一笔入 → 仓库隔离断言
            _txn("M1", "in", 2.0, 60, wh_b.id)

            # M4 从未出库 → 呆滞天数按建档日期起算（建档 d-100）
            self.materials["M4"].created_at = NOW - timedelta(days=100)
            db.session.commit()

    def _client(self):
        client = app_module.app.test_client()
        client.post("/login", data={"username": "admin", "password": "admin"})
        return client

    def _query(self, wh_id=None, **extra):
        params = {'warehouse_id': wh_id if wh_id is not None else self.wh_a_id}
        params.update(extra)
        qs = '&'.join(f"{k}={v}" for k, v in params.items())
        return self._client().get(f"/report/api/inventory_analysis?{qs}")

    def _rows_by_code(self, payload):
        return {row['code']: row for row in payload['data']}

    def test_T1_rows_are_material_level_and_warehouse_scoped(self):
        """宽表行=物料（不随仓库翻倍）；库存/金额按所选仓库净额隔离。"""
        j = self._query().get_json()
        assert j['status'] == 'success'
        assert j['total'] == 5, f"仓A 应 5 行（=物料数），实际 {j['total']}"
        rows = self._rows_by_code(j)
        assert rows['M1']['stock'] == 10.0, (
            f"仓A M1 净库存应 10（入100出90），实际 {rows['M1']['stock']}")
        assert rows['M1']['stock_value'] == 1000.0
        # 仓B 视角：行集仍=物料级（0 库存行保留——缺货预警依赖该行集语义，
        # 与库存报表默认不过滤零库存一致）；M1 净 2、金额 200，其余 0
        jb = self._query(wh_id=self.wh_b_id).get_json()
        assert jb['total'] == 5, f"仓B 行集应仍为物料级 5 行，实际 {jb['total']}"
        rows_b = self._rows_by_code(jb)
        assert rows_b['M1']['stock'] == 2.0 \
            and rows_b['M1']['stock_value'] == 200.0, (
            f"仓隔离：M1 仓B 净库存应 2、金额 200，"
            f"实际 {rows_b['M1']['stock']}/{rows_b['M1']['stock_value']}")
        assert rows_b['M2']['stock'] == 0.0, "M2 无仓B 流水应显示 0"
        # 表头字段：与 columns 定义一一对应（16 列宽表）
        assert [c['field'] for c in j['columns']] == [
            'code', 'name', 'spec', 'category', 'unit', 'supplier', 'stock',
            'price', 'stock_value', 'min_stock', 'stock_status', 'suggest_qty',
            'abc_class', 'last_out_date', 'idle_days', 'idle_flag']

    def test_T2_stock_status_and_suggest_qty(self):
        """缺货/低于安全/正常三态；min_stock<=0 不预警；补货量补至 max/min。"""
        rows = self._rows_by_code(self._query().get_json())
        assert rows['M3']['stock_status'] == '缺货', rows['M3']
        assert rows['M3']['suggest_qty'] == 20.0, "缺货应补至 max_stock=20"
        assert rows['M2']['stock_status'] == '低于安全库存', rows['M2']
        assert rows['M2']['suggest_qty'] == 1.0, "未设 max 应补至 min_stock=5-4=1"
        assert rows['M1']['stock_status'] == '正常', rows['M1']
        assert rows['M4']['stock_status'] == '正常', "min_stock=0 不应预警"
        assert rows['M4']['suggest_qty'] == 0.0
        assert rows['M5']['stock_status'] == '正常'
        # 展示字段：单价/安全库存原样透出
        assert rows['M2']['price'] == 50.0 and rows['M2']['min_stock'] == 5.0

    def test_T3_abc_classes_cumulative_ratio(self):
        """ABC：集成梯度 + 纯函数边界（首项豁免、total=0 全 C、同金额 id 序）。"""
        rows = self._rows_by_code(self._query().get_json())
        # 金额梯度 1000/200/0/3/5 → cum_before: 0→A, 1000→B, 1200→C...
        assert rows['M1']['abc_class'] == 'A', rows['M1']
        assert rows['M2']['abc_class'] == 'B', rows['M2']
        assert rows['M3']['abc_class'] == 'C'
        assert rows['M4']['abc_class'] == 'C'
        assert rows['M5']['abc_class'] == 'C'
        # 纯函数边界：金额 100/50/30/20/10/10（total 220）
        abc = app_module._abc_classes(
            {1: 100.0, 2: 50.0, 3: 30.0, 4: 20.0, 5: 10.0, 6: 10.0})
        assert abc == {1: 'A', 2: 'A', 3: 'A', 4: 'B', 5: 'C', 6: 'C'}, abc
        # 首项豁免：单物料 cum_before=0 <70% → A（而非 C）
        assert app_module._abc_classes({9: 100.0}) == {9: 'A'}
        # 金额合计 0 → 全 C
        assert app_module._abc_classes({1: 0.0, 2: 0.0}) == {1: 'C', 2: 'C'}

    def test_T4_idle_days_boundary_and_never_outbound(self):
        """呆滞：=90 达标、89 不达标、stock=0 不标记、建档起算、期初非出库。"""
        rows = self._rows_by_code(self._query().get_json())
        # M1 最后出库恰 90 天前且 stock>0 → 呆滞（>= 边界含等号）
        assert rows['M1']['idle_days'] == 90, rows['M1']
        assert rows['M1']['idle_flag'] == '呆滞'
        assert rows['M1']['last_out_date'] == (
            TODAY - timedelta(days=90)).isoformat()
        # M2 最后出库 89 天前 → 未达默认阈值 90
        assert rows['M2']['idle_days'] == 89 and rows['M2']['idle_flag'] == ''
        # M3 有出库但 stock=0 → 不标记（呆滞只针对占压库存）
        assert rows['M3']['last_out_date'] != '' \
            and rows['M3']['idle_flag'] == ''
        # M4 从未出库 → 按建档日期 d-100 起算，idle=100 ≥90 且 stock>0
        assert rows['M4']['last_out_date'] == ''
        assert rows['M4']['idle_days'] == 100, rows['M4']
        assert rows['M4']['idle_flag'] == '呆滞'
        # M5 仅期初流水：期初类型既非入也非出（与 _is_inbound_transaction
        # 同源）——不得被当作出库把 idle 拉到 200 天
        assert rows['M5']['last_out_date'] == ''
        assert rows['M5']['idle_days'] == 0, rows['M5']
        assert rows['M5']['idle_flag'] == ''

    def test_T5_pagination_sort_and_batched_export(self):
        """SQL_PAGED 协议接入：分页元数据、全量排序回退、020 分批导出/打印。"""
        j = self._query(page=1, page_size=2).get_json()
        assert j['total'] == 5 and len(j['data']) == 2
        assert j['total_pages'] == 3 and j['truncated'] is False
        # 表头排序：stock_value 降序 → M1(1000) 首行（全量排序回退）
        js = self._query(sort_field='stock_value',
                         sort_order='desc').get_json()
        assert [r['code'] for r in js['data']][:2] == ['M1', 'M2']
        # 020 分批导出：批 3 → 2 批全量 5 行、无截断提示
        old_batch = app_module.REPORT_EXPORT_BATCH_SIZE
        app_module.REPORT_EXPORT_BATCH_SIZE = 3
        try:
            r = self._client().get(
                f"/report/api/inventory_analysis?export=excel"
                f"&warehouse_id={self.wh_a_id}")
        finally:
            app_module.REPORT_EXPORT_BATCH_SIZE = old_batch
        assert r.status_code == 200
        xrows = _xlsx_rows(r.data)
        assert len(xrows) == 6, f"表头+5 行全量导出，实际 {len(xrows)} 行"
        joined = '\n'.join(cell for row in xrows for cell in row)
        assert '导出截断' not in joined
        # 模板打印同步接入（020 分批收集协议）
        rp = self._client().get(
            f"/report/inventory_analysis/print_excel?warehouse_id={self.wh_a_id}")
        assert rp.status_code == 200
        prows = _xlsx_rows(rp.data)
        assert len(prows) >= 6, "模板打印应含表头+5 行全量"

    def test_T6_hide_zero_filters_zero_stock_rows(self):
        """hide_zero=1 过滤零库存行（M3），口径与库存报表一致。"""
        j = self._query(hide_zero=1).get_json()
        assert j['total'] == 4, f"hide_zero 应剔除 M3 仅 4 行，实际 {j['total']}"
        codes = {row['code'] for row in j['data']}
        assert 'M3' not in codes and {'M1', 'M2', 'M4', 'M5'} == codes
        # 默认不过滤：零库存行保留（缺货预警依赖该行）
        j_all = self._query().get_json()
        assert j_all['total'] == 5

    def test_T7_summary_scope(self):
        """summary：count=物料数、quantity=净库存和、amount=金额和。"""
        j = self._query().get_json()
        s = j['summary']
        assert s['count'] == 5, f"count 应为物料数 5，实际 {s['count']}"
        assert abs(s['quantity'] - 18.0) < 1e-6, (
            f"净库存和 10+4+0+3+1=18，实际 {s['quantity']}")
        assert abs(s['amount'] - 1208.0) < 1e-6, (
            f"金额和 1000+200+0+3+5=1208，实际 {s['amount']}")
        assert j['summary_labels'] == \
            app_module.REPORT_DEFINITIONS['inventory_analysis']['summary_labels']
