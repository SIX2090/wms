# -*- coding: utf-8 -*-
"""BUG-2026-09-30-004 回归：/report/inout/print 与 /report/inout/export 的 N+1 查询锁。

现象（M1 性能基线，scripts/perf_export_baseline_5000.json）：
- 5000 物料/单时导出耗时 15.57s、内存峰值 45.3MB，而响应仅 172KB——
  耗时全在查询而非写入；
- 根因：`for order in .all(): for item in order.items:` 逐单懒加载 items，
  `order.supplier` / `item.material` 再逐行懒加载——每单 3 次额外查询，
  5000 单 ≈ 1.5 万次 SQL。

修复：selectinload 批量预加载（items→material 链、supplier），
查询数与单据数解耦（SQLAlchemy selectin 每批 500）。

回归锁（本文件）：
1. 查询计数锁——N 单导出期间的 SQL 数必须与 N 解耦（有界）；
   修复前 ~3N+ 条，修复后 ~15 条以内，阈值取 50（80 入库单 + 40 出库单）。
2. 内容正确性锁——导出 Excel 行数/字段值与造数一致，
   确保 eager loading 不改变导出内容（含 supplier.name 显示）。
"""
from __future__ import annotations

import io
import os
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
APP_DIR = ROOT / "app"
sys.path.insert(0, str(APP_DIR))
os.chdir(APP_DIR)

os.environ.setdefault("WMS_BOOTSTRAP_PASSWORD", "admin")
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["WMS_DATABASE_URI"] = "sqlite:///:memory:"
os.environ.setdefault("WMS_DEBUG", "0")

import app as app_module  # noqa: E402
from app import db  # noqa: E402

app_module.app.config["TESTING"] = True
app_module.app.config["WTF_CSRF_ENABLED"] = False

TODAY = date.today()

N_IN = 80    # 入库单数
N_OUT = 40   # 出库单数


def _reset_db():
    db.drop_all()
    db.create_all()


def _make_client():
    client = app_module.app.test_client()
    client.post(
        "/login",
        data={"username": "admin", "password": "admin"},
        content_type="application/x-www-form-urlencoded",
    )
    return client


def _seed_admin():
    from werkzeug.security import generate_password_hash
    from app import User
    db.session.add(User(username="admin", password_hash=generate_password_hash("admin"),
                        role="admin", must_change_password=False))
    db.session.commit()


class TestInOutExportNPlusOne:
    """selectinload 回归锁：查询数有界 + 导出内容不变。"""

    def teardown_method(self):
        # 清场：残留 120 单/物料脏数据会污染同进程后续未在 setup 重建库的
        # 测试文件（实证：test_feat_2026_09_24_001 test_t8_marker_rows_blank
        # 因本文件残留数据顺序依赖失败）。drop+create 留空库。
        with app_module.app.app_context():
            db.drop_all()
            db.create_all()

    def test_export_query_count_bounded_and_content_correct(self):
        from sqlalchemy import event
        from app import (InOrder, InOrderItem, OutOrder, OutOrderItem,
                         Material, Supplier, Warehouse)

        with app_module.app.app_context():
            _reset_db()
            _seed_admin()
            wh = Warehouse(code="WHA", name="仓库A", status="active", is_default=True)
            db.session.add(wh)
            db.session.flush()
            sup = Supplier(code="SUP-N1", name="鑫达供货")
            db.session.add(sup)
            db.session.flush()

            # 每单一独立物料（物料懒加载无法命中缓存，最坏场景），
            # 入库单挂 supplier（走 supplier.name 路径）
            for i in range(N_IN):
                m = Material(code=f"MP{i:04d}", name=f"物料{i}", stock=0)
                db.session.add(m)
                db.session.flush()
                o = InOrder(order_no=f"IN-N1-{i:04d}", warehouse="仓库A",
                            status="completed", date=TODAY, supplier_id=sup.id)
                db.session.add(o)
                db.session.flush()
                db.session.add(InOrderItem(in_order_id=o.id, material_id=m.id,
                                           quantity=3, price=2, amount=6))
            for i in range(N_OUT):
                m = Material(code=f"MO{i:04d}", name=f"出料{i}", stock=0)
                db.session.add(m)
                db.session.flush()
                o = OutOrder(order_no=f"OUT-N1-{i:04d}", warehouse="仓库A",
                             status="completed", date=TODAY)
                db.session.add(o)
                db.session.flush()
                db.session.add(OutOrderItem(out_order_id=o.id, material_id=m.id,
                                            quantity=2, price=5, amount=10))
            db.session.commit()
            wh_id = wh.id

        # 登录在计数窗口外（登录本身有 SQL，避免噪声）
        client = _make_client()

        counter = {"n": 0}

        def _count(conn, cursor, statement, parameters, context, executemany):
            counter["n"] += 1

        with app_module.app.app_context():
            event.listen(db.engine, "before_cursor_execute", _count)
            try:
                resp = client.get(f"/report/inout/export?warehouse_id={wh_id}")
            finally:
                event.remove(db.engine, "before_cursor_execute", _count)

        # --- 锁 1：查询数有界（修复前 80 单入库 ≈ 1+80+80+80 条 + 40 单出库
        #     ≈ 1+40+40 条，共 320+；修复后 ~15 条内。阈值 50 双向安全带） ---
        assert resp.status_code == 200, resp.status_code
        assert counter["n"] <= 50, (
            f"导出期间 SQL 查询数 {counter['n']} 超阈值 50——"
            f"N+1 懒加载疑似回归（N_IN={N_IN}, N_OUT={N_OUT}，"
            f"修复后预期 ~15 条，修复前 ~320 条）"
        )

        # --- 锁 2：内容正确性（eager loading 不得改变导出内容） ---
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.data))
        assert "入库统计" in wb.sheetnames and "领料统计" in wb.sheetnames

        ws_in = wb["入库统计"]
        rows_in = list(ws_in.iter_rows(values_only=True))
        assert rows_in[0][0] == "单据编号"
        data_in = rows_in[1:]
        assert len(data_in) == N_IN, f"入库行数 {len(data_in)} != {N_IN}"
        # spot check：字段齐全（编号/供应商/物料/数量/金额）——同日期排序不保证
        # 插入序，按 order_no 建索引后断言
        by_no_in = {r[0]: r for r in data_in}
        row0 = by_no_in["IN-N1-0000"]
        assert row0[2] == "鑫达供货", f"supplier.name 未随导出显示：{row0[2]!r}"
        assert row0[5] == "MP0000" and row0[6] == "物料0", row0
        assert row0[7] == 3 and row0[8] == 6, row0
        # 每行都带 supplier（supplier 批量预加载正确性的全集校验）
        assert all(r[2] == "鑫达供货" for r in data_in), "存在 supplier 缺失行"

        ws_out = wb["领料统计"]
        rows_out = list(ws_out.iter_rows(values_only=True))
        data_out = rows_out[1:]
        assert len(data_out) == N_OUT, f"出库行数 {len(data_out)} != {N_OUT}"
        by_no_out = {r[0]: r for r in data_out}
        row0o = by_no_out["OUT-N1-0000"]
        assert row0o[5] == "MO0000" and row0o[7] == 2 and row0o[8] == 10, row0o
