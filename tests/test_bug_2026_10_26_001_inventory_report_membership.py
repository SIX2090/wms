"""库存报表行集必须属于所选仓；保留真实零余额及既有库存归属规则。"""
from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

import app as m


@pytest.fixture(params=['0', '1'], ids=['without_locations', 'with_locations'])
def inventory_scope(request, monkeypatch):
    monkeypatch.setitem(m.app.config, 'TESTING', True)
    monkeypatch.setitem(m.app.config, 'WTF_CSRF_ENABLED', False)
    with m.app.app_context():
        m.db.drop_all()
        m.db.create_all()
        a = m.Warehouse(code='WA', name='甲仓', status='active', is_default=True)
        b = m.Warehouse(code='WB', name='乙仓', status='active')
        empty = m.Warehouse(code='WC', name='空仓', status='active')
        unit = m.Unit(code='PCS', name='件')
        supplier = m.Supplier(code='S1', name='供应商甲')
        user = m.User(username='scope-user', role='admin', must_change_password=False,
                      password_hash=generate_password_hash('qa-scope-only'))
        m.db.session.add_all([a, b, empty, unit, supplier, user])
        m.db.session.flush()
        mats = {}
        for code in ['A', 'SHARED', 'B', 'ZERO', 'NEGATIVE', 'LEGACY-NAME',
                     'LEGACY-CODE', 'SOURCE', 'CONFLICT', 'UNKNOWN', 'NEVER']:
            material = m.Material(code=code, name='物料-' + code, unit_id=unit.id,
                                  price=2, stock=9999,
                                  supplier_id=supplier.id if code == 'A' else None)
            m.db.session.add(material)
            m.db.session.flush()
            mats[code] = material.id
        source = m.InOrder(order_no='QA-IN-A', date=date(2026, 10, 26),
                           warehouse=a.name, status='completed', operator_id=user.id)
        m.db.session.add(source)
        m.db.session.flush()
        m.set_system_setting('location_management_enabled', request.param)

        def record(code, qty, wh_id, location, ref=None):
            if request.param == '1':
                m.db.session.add(m.LocationInventory(
                    material_id=mats[code], quantity=qty,
                    warehouse_id=wh_id, location=location or ''))
            else:
                m.db.session.add(m.StockTransaction(
                    material_id=mats[code], quantity=qty, warehouse_id=wh_id,
                    location=location, transaction_type='in' if qty >= 0 else 'out',
                    reference_type='in_order' if ref else None,
                    reference_id=ref, operator_id=user.id))

        record('A', 5, a.id, b.name)  # 明确 ID 优先于冲突的历史名称
        record('SHARED', 2, a.id, 'A-01')
        record('SHARED', 7, b.id, 'B-01')
        record('B', 11, b.id, 'B-01')
        record('ZERO', 4, a.id, 'A-01')
        record('ZERO', -4, a.id, 'A-02')
        record('NEGATIVE', -2, a.id, 'A-01')
        record('LEGACY-NAME', 3, None, a.name)
        record('LEGACY-CODE', 4, None, a.code)
        record('SOURCE', 6, None, None, source.id)
        record('CONFLICT', 55, b.id, a.name)
        record('UNKNOWN', 888, None, None)  # 无可确认归属，不猜到默认仓
        m.db.session.commit()
        result = dict(a=a.id, b=b.id, empty=empty.id, supplier=supplier.id,
                      user=user.id, mode=request.param,
                      expected_a={'A': 5, 'SHARED': 2, 'ZERO': 0, 'NEGATIVE': -2,
                                  'LEGACY-NAME': 3, 'LEGACY-CODE': 4},
                      expected_b={'SHARED': 7, 'B': 11, 'CONFLICT': 55})
        if request.param == '0':
            result['expected_a']['SOURCE'] = 6
    client = m.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(result['user'])
        session['_fresh'] = True
    result['client'] = client
    yield result
    with m.app.app_context():
        m.db.session.remove()


def _query(case, warehouse, **params):
    response = case['client'].get('/report/api/inventory', query_string={
        'warehouse_id': case[warehouse], 'page_size': 500, **params})
    assert response.status_code == 200
    data = response.get_json()
    assert data['status'] == 'success'
    return data


def _assert_rows(data, expected):
    assert {row['code']: row['stock'] for row in data['data']} == expected
    assert data['total'] == len(expected)
    assert data['summary'] == {'count': len(expected),
                               'quantity': sum(expected.values()),
                               'amount': 2 * sum(expected.values())}


def _xlsx_codes(response):
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    try:
        return [row[0] for row in workbook.active.iter_rows(values_only=True)
                if row and row[0] is not None]
    finally:
        workbook.close()


def test_switch_warehouse_changes_membership_and_keeps_own_zero(inventory_scope):
    case = inventory_scope
    for warehouse in ['a', 'b', 'a']:
        _assert_rows(_query(case, warehouse), case['expected_' + warehouse])


def test_empty_warehouse_has_no_other_warehouse_materials(inventory_scope):
    _assert_rows(_query(inventory_scope, 'empty'), {})


def test_hide_zero_only_removes_real_zero_and_keeps_negative(inventory_scope):
    expected = {k: v for k, v in inventory_scope['expected_a'].items() if v != 0}
    _assert_rows(_query(inventory_scope, 'a', hide_zero=1), expected)


def test_supplier_filter_intersects_warehouse_membership(inventory_scope):
    case = inventory_scope
    _assert_rows(_query(case, 'a', supplier_id=case['supplier']), {'A': 5})
    _assert_rows(_query(case, 'b', supplier_id=case['supplier']), {})


def test_material_search_cannot_reintroduce_other_warehouse_rows(inventory_scope):
    _assert_rows(_query(inventory_scope, 'b', material_code='物料-A'), {})
    _assert_rows(_query(inventory_scope, 'a', material_code='物料-A'), {'A': 5})


def test_pagination_sort_summary_and_export_share_full_scope(inventory_scope):
    case = inventory_scope
    collected = []
    expected = case['expected_a']
    for page in range(1, (len(expected) + 1) // 2 + 1):
        data = _query(case, 'a', page=page, page_size=2,
                      sort_field='stock', sort_order='desc')
        assert data['total'] == len(expected)
        assert data['summary']['quantity'] == sum(expected.values())
        collected.extend(data['data'])
    assert {r['code']: r['stock'] for r in collected} == expected
    assert [r['stock'] for r in collected] == sorted(expected.values(), reverse=True)
    response = case['client'].get('/report/api/inventory', query_string={
        'warehouse_id': case['a'], 'page': 2, 'page_size': 2,
        'sort_field': 'stock', 'sort_order': 'desc', 'export': 'excel'})
    codes = _xlsx_codes(response)
    assert codes[0] == '物料编码'
    assert codes[1:] == [r['code'] for r in collected]


def test_template_print_uses_same_warehouse_full_set(inventory_scope):
    case = inventory_scope
    response = case['client'].get('/report/inventory/print_excel', query_string={
        'warehouse_id': case['b'], 'page': 2, 'page_size': 1})
    codes = _xlsx_codes(response)
    assert set(codes).intersection({'A', 'ZERO', 'NEGATIVE', 'LEGACY-NAME',
                                  'LEGACY-CODE', 'SOURCE', 'UNKNOWN', 'NEVER'}) == set()
    assert set(case['expected_b']).issubset(codes)


def test_default_and_unknown_warehouse_do_not_leak_global_materials(inventory_scope):
    case = inventory_scope
    response = case['client'].get('/report/api/inventory', query_string={'page_size': 500})
    assert response.status_code == 200
    _assert_rows(response.get_json(), case['expected_a'])
    response = case['client'].get('/report/api/inventory', query_string={'warehouse_id': 999999})
    assert response.status_code == 200
    _assert_rows(response.get_json(), {})
    with m.app.app_context():
        m.Warehouse.query.update({'is_default': False})
        m.db.session.commit()
    response = case['client'].get('/report/api/inventory')
    assert response.status_code == 400
    assert response.get_json()['msg'] == '请选择仓库'


def test_analysis_preserves_explicit_master_list_for_shortage_warnings(inventory_scope):
    case = inventory_scope
    response = case['client'].get('/report/api/inventory_analysis', query_string={
        'warehouse_id': case['b'], 'page_size': 500})
    assert response.status_code == 200
    data = response.get_json()
    assert data['total'] == 11  # 经营分析原有缺货预警约定，不随静态库存报表改动
    rows = {row['code']: row['stock'] for row in data['data']}
    assert rows['NEVER'] == 0
    assert rows['SHARED'] == 7
    assert data['summary']['quantity'] == 73
