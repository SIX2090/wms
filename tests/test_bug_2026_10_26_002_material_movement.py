"""物料收发存：库存事实、业务日期、仓库隔离、期初及全集输出契约。"""
from datetime import date, datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse
import io

import pytest
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash
from sqlalchemy import event
import app as m

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(params=['0', '1'], ids=['locations_off', 'locations_on'])
def movement(request, monkeypatch, tmp_path):
    monkeypatch.setitem(m.app.config, 'TESTING', True)
    monkeypatch.setitem(m.app.config, 'WTF_CSRF_ENABLED', False)
    monkeypatch.setattr(m.app, 'static_folder', str(tmp_path / 'static'))
    with m.app.app_context():
        m.db.drop_all(); m.db.create_all()
        whs = [m.Warehouse(code='MV-W' + x, name=x + '仓', status='active',
                           is_default=x == 'A') for x in 'ABC']
        units = [m.Unit(code='PCS', name='件'), m.Unit(code='M', name='米')]
        user = m.User(username='movement-qa', role='admin', must_change_password=False,
                      password_hash=generate_password_hash('movement-fixture-only'))
        m.db.session.add_all([*whs, *units, user]); m.db.session.flush()
        a, b, empty = whs
        mats = {}
        for code in ['A', 'CARRY', 'GONE', 'SHARED', 'B', 'LEGACY', 'UNKNOWN', 'DRAFT', 'FUTURE']:
            material = m.Material(code='MV-' + code, name='收发存-' + code, spec='规格-' + code,
                                  unit_id=units[1 if code == 'SHARED' else 0].id,
                                  stock=99999, price=987)
            m.db.session.add(material); m.db.session.flush(); mats[code] = material.id
        docs = []
        for suffix, day, status in [('PRE', 1, 'completed'), ('IN', 10, 'completed'),
                                    ('FUTURE', 21, 'completed'), ('DRAFT', 12, 'pending')]:
            doc = m.InOrder(order_no='MV-IN-' + suffix, date=date(2026, 10, day),
                            warehouse=a.name, status=status, operator_id=user.id)
            m.db.session.add(doc); m.db.session.flush(); docs.append(doc)
        out = m.OutOrder(order_no='MV-OUT', date=date(2026, 10, 20), warehouse=a.name,
                         status='completed', operator_id=user.id)
        m.db.session.add(out); m.db.session.flush()
        # 草稿含真实明细，但没有库存事实，不得成为收发存。
        m.db.session.add(m.InOrderItem(in_order_id=docs[3].id,
                                      material_id=mats['DRAFT'], quantity=999, price=0, amount=0))

        def txn(code, qty, day, wh=a, location=None, reference_type=None, reference_id=None,
                kind=None, created=None):
            m.db.session.add(m.StockTransaction(
                material_id=mats[code], quantity=qty, warehouse_id=wh.id if wh else None,
                location=location, reference_type=reference_type, reference_id=reference_id,
                transaction_type=kind or ('in' if qty >= 0 else 'out'), operator_id=user.id,
                created_at=created or datetime(2026, 10, day, 12)))

        txn('A', 100, 25, reference_type='in_order', reference_id=docs[0].id)
        txn('A', 20, 25, reference_type='in_order', reference_id=docs[1].id)
        txn('A', -7, 1, reference_type='out_order', reference_id=out.id)
        txn('A', -4, 12, kind='transfer_out'); txn('A', 4, 12, wh=b, kind='transfer_in')
        txn('A', 3, 14, kind='adjustment_in'); txn('A', -2, 15, kind='after_sale_out')
        txn('A', 5, 16, kind='opening')  # 本期内建账也是本期库存增加，不能丢失
        txn('A', 123, 1, reference_type='in_order', reference_id=docs[2].id)
        txn('CARRY', 10, 2, kind='opening'); txn('GONE', 5, 2); txn('GONE', -5, 13)
        txn('SHARED', 2, 11); txn('SHARED', 90, 11, wh=b)
        txn('B', 50, 2, wh=b)
        txn('LEGACY', 3, 2, wh=None, location=a.name)
        txn('LEGACY', 4, 11, wh=None, location=a.code)
        txn('LEGACY', 6, 25, wh=None, reference_type='in_order', reference_id=docs[1].id)
        txn('UNKNOWN', 888, 11, wh=None)
        txn('B', 9, 11, wh=b, location=a.name)  # 明确仓库ID不被旧名称覆盖
        txn('FUTURE', 123, 1, reference_type='in_order', reference_id=docs[2].id)
        m.set_system_setting('location_management_enabled', request.param)
        m.db.session.commit()
        case = dict(a=a.id, b=b.id, empty=empty.id, user=user.id, mats=mats)
    client = m.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(case['user']); session['_fresh'] = True
    case['client'] = client
    yield case
    with m.app.app_context(): m.db.session.remove()


def query(case, **overrides):
    response = case['client'].get('/report/api/material_movement', query_string={
        'warehouse_id': case['a'], 'start_date': '2026-10-10', 'end_date': '2026-10-20',
        'page_size': 500, **overrides})
    assert response.status_code == 200, response.get_data(as_text=True)
    data = response.get_json(); assert data['status'] == 'success'
    return data


def by_code(data): return {row['material_code']: row for row in data['data']}


def workbook_rows(response):
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.data), read_only=True)
    try: return list(book.active.iter_rows(values_only=True))
    finally: book.close()


def test_default_summary_opens_material_view_and_keeps_daily(movement):
    page = movement['client'].get('/report/view/summary').get_data(as_text=True)
    assert 'const reportType = "material_movement"' in page
    assert '期初数量' in page and '期末结存' in page and '按日期汇总' in page
    assert 'name="supplier"' not in page and 'name="status"' not in page
    daily = movement['client'].get('/report/view/summary?summary_view=daily').get_data(as_text=True)
    assert 'const reportType = "summary"' in daily
    assert 'name="status"' in daily and 'name="summary_view"' in daily
    assert 'id="movementDateScope"' in daily  # 不限日期经日报切回后不能被默认本月覆盖
    with m.app.app_context():
        columns = {col['field']: col['title'] for col in m._summary_columns()}
        assert columns['in_count'] == '入库明细行数'
        assert columns['out_count'] == '领料明细行数'


def test_period_balances_use_business_date_and_all_stock_movements(movement):
    data = query(movement); rows = by_code(data)
    assert set(rows) == {'MV-A', 'MV-CARRY', 'MV-GONE', 'MV-SHARED', 'MV-LEGACY'}
    a = rows['MV-A']
    assert [a[k] for k in ['opening_quantity', 'in_quantity', 'out_quantity', 'balance_quantity']] == [100, 28, 13, 115]
    assert [rows['MV-CARRY'][k] for k in ['opening_quantity', 'in_quantity', 'out_quantity', 'balance_quantity']] == [10, 0, 0, 10]
    assert rows['MV-LEGACY']['opening_quantity'] == 3
    assert rows['MV-LEGACY']['in_quantity'] == 10
    for row in rows.values():
        assert row['opening_quantity'] + row['in_quantity'] - row['out_quantity'] == pytest.approx(row['balance_quantity'])
    assert not any('amount' in col['field'] for col in data['columns'])


def test_warehouses_and_ambiguous_legacy_records_stay_isolated(movement):
    b = by_code(query(movement, warehouse_id=movement['b']))
    assert set(b) == {'MV-A', 'MV-B', 'MV-SHARED'}
    assert b['MV-A']['in_quantity'] == 4
    assert b['MV-B']['opening_quantity'] == 50 and b['MV-B']['in_quantity'] == 9
    empty = query(movement, warehouse_id=movement['empty'])
    assert empty['total'] == 0 and empty['data'] == [] and empty['summary']['unit_totals'] == []


def test_unit_totals_are_separate_and_match_all_material_rows(movement):
    data = query(movement)
    totals = {row['unit']: row for row in data['summary']['unit_totals']}
    assert set(totals) == {'件', '米'}
    assert data['summary']['count'] == data['total'] == 5
    for unit, total in totals.items():
        for key in ['opening_quantity', 'in_quantity', 'out_quantity', 'balance_quantity']:
            assert total[key] == pytest.approx(sum(r[key] for r in data['data'] if r['unit'] == unit))
    assert totals['米']['balance_quantity'] == 2


def test_hide_zero_and_material_keyword_filter(movement):
    data = query(movement, hide_zero=1)
    assert 'MV-GONE' not in by_code(data) and data['total'] == 4
    only = query(movement, material_code='MV-CARRY')
    assert set(by_code(only)) == {'MV-CARRY'}
    assert only['summary']['unit_totals'][0]['opening_quantity'] == 10


def test_pagination_sort_export_and_template_print_use_same_full_set(movement):
    expected = query(movement, sort_field='balance_quantity', sort_order='desc')
    collected = []
    for page in range(1, 4):
        part = query(movement, page=page, page_size=2, sort_field='balance_quantity', sort_order='desc')
        assert part['summary'] == expected['summary']; collected.extend(part['data'])
    assert collected == expected['data']
    params = dict(warehouse_id=movement['a'], start_date='2026-10-10', end_date='2026-10-20', page=2, page_size=1)
    export = workbook_rows(movement['client'].get('/report/api/material_movement', query_string={**params, 'export': 'excel'}))
    assert set(row[0] for row in export[1:]) == set(by_code(expected))
    printed = workbook_rows(movement['client'].get('/report/material_movement/print_excel', query_string=params))
    assert set(by_code(expected)).issubset({row[0] for row in printed})


def test_material_link_opens_same_warehouse_and_period_ledger(movement):
    row = by_code(query(movement))['MV-A']
    url = urlparse(row['ledger_url']); args = parse_qs(url.query)
    assert url.path == '/report/view/ledger'
    assert args['warehouse_id'] == [str(movement['a'])]
    assert args['material_code'] == ['MV-A']
    assert args['material_id'] == [str(movement['mats']['A'])]
    assert args['start_date'] == ['2026-10-10'] and args['end_date'] == ['2026-10-20']
    response = movement['client'].get('/report/api/ledger?' + url.query)
    assert response.status_code == 200
    totals = [r for r in response.get_json()['data'] if r['reference_type'] == '本期合计']
    assert totals[0]['balance_quantity'] == 115


def test_linked_ledger_does_not_include_similarly_named_material(movement):
    with m.app.app_context():
        extra = m.Material(code='MV-A-EXTRA', name='相似物料', unit_id=1, stock=70)
        m.db.session.add(extra); m.db.session.flush()
        m.db.session.add(m.StockTransaction(material_id=extra.id, warehouse_id=movement['a'],
                                            quantity=70, transaction_type='in', created_at=datetime(2026, 10, 12)))
        m.db.session.commit()
    row = by_code(query(movement))['MV-A']
    url = urlparse(row['ledger_url'])
    data = movement['client'].get('/report/api/ledger?' + url.query).get_json()
    assert {r['material_code'] for r in data['data']} == {'MV-A'}
    page = movement['client'].get(row['ledger_url']).get_data(as_text=True)
    assert 'id="ledgerMaterialId"' in page


def test_invalid_warehouse_and_date_range_fail_closed(movement):
    assert query(movement, warehouse_id=999999)['total'] == 0
    response = movement['client'].get('/report/api/material_movement', query_string={
        'warehouse_id': movement['a'], 'start_date': '2026-10-20', 'end_date': '2026-10-10'})
    assert response.status_code == 400
    with m.app.app_context():
        m.Warehouse.query.update({'is_default': False}); m.db.session.commit()
    assert movement['client'].get('/report/api/material_movement').status_code == 400


def test_unbounded_dates_balance_from_start_of_history(movement):
    row = by_code(query(movement, start_date='', end_date=''))['MV-A']
    assert [row[k] for k in ['opening_quantity', 'in_quantity', 'out_quantity', 'balance_quantity']] == [0, 251, 13, 238]


def test_report_query_does_not_write_inventory(movement):
    statements = []
    with m.app.app_context():
        engine = m.db.engine
        def before(_conn, _cursor, sql, _params, _context, _many):
            if sql.lstrip().upper().startswith(('INSERT', 'UPDATE', 'DELETE', 'CREATE', 'ALTER')):
                statements.append(sql)
        event.listen(engine, 'before_cursor_execute', before)
    try: query(movement)
    finally: event.remove(engine, 'before_cursor_execute', before)
    assert statements == []


def test_more_than_ledger_limit_is_never_silently_truncated(movement):
    with m.app.app_context():
        m.db.session.execute(m.StockTransaction.__table__.insert(), [dict(
            material_id=movement['mats']['A'], warehouse_id=movement['empty'],
            quantity=.01, transaction_type='in', created_at=datetime(2026, 10, 12)) for _ in range(50003)])
        m.db.session.commit()
    data = query(movement, warehouse_id=movement['empty'])
    assert data['total'] == 1 and not data['truncated']
    assert data['data'][0]['in_quantity'] == pytest.approx(500.03)
    assert data['data'][0]['balance_quantity'] == pytest.approx(500.03)


def test_template_center_registers_and_can_create_material_report_template(movement):
    from routes.print_template_center import _code_options, _report_columns_for_blank
    with m.app.app_context():
        options = {item['code']: item['label'] for item in _code_options()['report']}
        assert 'report_material_movement' in options
        assert options['report_material_movement'] != options['report_summary']
        title, columns = _report_columns_for_blank('material_movement')
        assert '收发存' in title
        assert {col['field'] for col in columns}.issuperset({'opening_quantity', 'balance_quantity'})
    response = movement['client'].post('/print_templates/create_blank', data={
        'name': '隔离收发存模板', 'target_type': 'report', 'target_code': 'report_material_movement'})
    assert response.status_code == 302
    with m.app.app_context():
        template = m.ExcelPrintTemplate.query.filter_by(name='隔离收发存模板').one()
        from print_fill import template_file_abspath
        path = Path(template_file_abspath(template.excel_template_path, m.app.static_folder)).resolve()
        assert path.is_file() and ROOT not in path.parents
