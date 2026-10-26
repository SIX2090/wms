"""仓库月报：业务月份、仓库行集、连续结存及末月汇总。"""
from datetime import date, datetime
from urllib.parse import urlparse, parse_qs
import io

import pytest
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash
from sqlalchemy import event
import app as m


@pytest.fixture(params=['0', '1'])
def monthly(request, monkeypatch, tmp_path):
    monkeypatch.setitem(m.app.config, 'TESTING', True)
    monkeypatch.setitem(m.app.config, 'WTF_CSRF_ENABLED', False)
    monkeypatch.setattr(m.app, 'static_folder', str(tmp_path / 'static'))
    with m.app.app_context():
        m.db.drop_all(); m.db.create_all()
        warehouses = [m.Warehouse(code='MONTH-' + x, name=x + '月报仓', status='active',
                                  is_default=x == 'A') for x in 'ABC']
        units = [m.Unit(code='PC', name='件'), m.Unit(code='MT', name='米')]
        user = m.User(username='monthly-qa', role='admin', must_change_password=False,
                      password_hash=generate_password_hash('monthly-fixture-only'))
        m.db.session.add_all([*warehouses, *units, user]); m.db.session.flush()
        a, b, empty = warehouses
        mats = {}
        for code in ['A', 'B', 'CARRY', 'ZERO', 'METRE', 'LEGACY', 'UNKNOWN', 'FUTURE', 'FREE']:
            material = m.Material(code='MONTH-' + code, name='月报物料' + code, stock=99999,
                                  price=99, unit_id=units[code == 'METRE'].id)
            m.db.session.add(material); m.db.session.flush(); mats[code] = material.id
        def txn(code, qty, day, wh=a, location=None, ref=None):
            m.db.session.add(m.StockTransaction(material_id=mats[code], quantity=qty,
                transaction_type='in' if qty >= 0 else 'out', warehouse_id=wh.id if wh else None,
                location=location, created_at=datetime.fromisoformat(day),
                reference_type='in_order' if ref else None, reference_id=ref))
        # 8月期末100，9月120-30=90，10月90+5-10=85；录入时间跨月。
        txn('A', 100, '2026-08-10'); txn('A', 20, '2026-09-10')
        txn('A', -30, '2026-09-15'); txn('A', -10, '2026-10-15')
        doc = m.InOrder(order_no='MONTH-BACKDATED', date=date(2026, 10, 12),
                        warehouse=a.name, status='completed', operator_id=user.id)
        free = m.InOrder(order_no='MONTH-FREE', date=date(2026, 10, 13),
                         warehouse=a.name, status='completed', operator_id=user.id)
        m.db.session.add_all([doc, free]); m.db.session.flush()
        m.db.session.add_all([m.InOrderItem(in_order_id=doc.id, material_id=mats['A'], quantity=5,
                                           price=2, amount=10),
                              m.InOrderItem(in_order_id=free.id, material_id=mats['FREE'], quantity=2,
                                           price=0, amount=0)])
        txn('A', 5, '2026-11-02', ref=doc.id); txn('FREE', 2, '2026-11-02', ref=free.id)
        txn('B', 777, '2026-09-01', wh=b)
        txn('CARRY', 9, '2026-08-01'); txn('ZERO', 4, '2026-09-01'); txn('ZERO', -4, '2026-09-02')
        txn('METRE', 10, '2026-08-01'); txn('METRE', -1, '2026-10-01')
        txn('LEGACY', 3, '2026-08-01', wh=None, location=a.code)
        txn('LEGACY', 2, '2026-10-01', wh=None, location=a.name)
        txn('UNKNOWN', 888, '2026-09-01', wh=None)
        txn('FUTURE', 999, '2026-11-01')
        m.set_system_setting('location_management_enabled', request.param)
        m.db.session.commit()
        case = dict(a=a.id, b=b.id, empty=empty.id, user=user.id, mats=mats)
    client = m.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(case['user']); session['_fresh'] = True
    case['client'] = client
    yield case
    with m.app.app_context(): m.db.session.remove()


def query(case, **changes):
    response = case['client'].get('/report/api/warehouse_monthly', query_string={
        'warehouse_id': case['a'], 'start_date': '2026-09-01', 'end_date': '2026-10-31',
        'page_size': 500, **changes})
    assert response.status_code == 200, response.get_data(as_text=True)
    return response.get_json()


def rows_for(data, code): return {r['month']: r for r in data['data'] if r['material_code'] == 'MONTH-' + code}


def test_warehouse_membership_and_empty_warehouse(monthly):
    codes = {r['material_code'] for r in query(monthly)['data']}
    assert codes == {'MONTH-' + x for x in ['A', 'CARRY', 'ZERO', 'METRE', 'LEGACY', 'FREE']}
    assert {r['material_code'] for r in query(monthly, warehouse_id=monthly['b'])['data']} == {'MONTH-B'}
    assert query(monthly, warehouse_id=monthly['empty'])['total'] == 0


def test_build_warehouse_monthly_report(monthly):
    rows = rows_for(query(monthly), 'A')
    assert [rows['2026-09'][k] for k in ['beginning_quantity','in_quantity','out_quantity','ending_quantity']] == [100,20,30,90]
    assert [rows['2026-10'][k] for k in ['beginning_quantity','in_quantity','out_quantity','ending_quantity']] == [90,5,10,85]
    assert rows['2026-09']['ending_quantity'] == rows['2026-10']['beginning_quantity']
    assert rows['2026-10']['in_amount'] == 10


def test_period_summary_never_sums_monthly_balances_or_mixed_units(monthly):
    data = query(monthly)
    totals = {v['unit']: v for v in data['summary']['unit_totals']}
    assert [totals['件'][k] for k in ['beginning_quantity','in_quantity','out_quantity','ending_quantity']] == [112,33,44,101]
    assert [totals['米'][k] for k in ['beginning_quantity','in_quantity','out_quantity','ending_quantity']] == [10,0,1,9]
    for unit, total in totals.items():
        rows = [r for r in data['data'] if r['unit'] == unit]
        assert total['beginning_quantity'] == sum(r['beginning_quantity'] for r in rows if r['month'] == '2026-09')
        assert total['ending_quantity'] == sum(r['ending_quantity'] for r in rows if r['month'] == '2026-10')
        assert total['in_quantity'] == sum(r['in_quantity'] for r in rows)
        assert total['out_quantity'] == sum(r['out_quantity'] for r in rows)


def test_month_inputs_and_midmonth_dates_mean_whole_months(monthly):
    full = query(monthly)
    assert query(monthly, start_date='2026-09-20', end_date='2026-10-02')['data'] == full['data']
    assert query(monthly, start_date='', end_date='', start_month='2026-09', end_month='2026-10')['data'] == full['data']


def test_zero_month_omitted_but_carry_forward_preserved(monthly):
    data = query(monthly)
    assert set(rows_for(data, 'ZERO')) == {'2026-09'}
    assert set(rows_for(data, 'CARRY')) == {'2026-09', '2026-10'}
    assert all(r['ending_quantity'] == 9 for r in rows_for(data, 'CARRY').values())
    hidden = query(monthly, hide_zero='1')
    assert not rows_for(hidden, 'ZERO')
    assert hidden['summary']['unit_totals']  # totals use the visible full row set


def test_zero_actual_price_does_not_fall_back_to_current_price(monthly):
    with m.app.app_context():
        m.StockTransaction.query.filter_by(material_id=monthly['mats']['FREE']).update({'created_at':datetime(2026,10,13)})
        m.db.session.commit()
    row = rows_for(query(monthly), 'FREE')['2026-10']
    assert row['in_amount'] == 0


def test_invalid_month_and_warehouse_fail_closed(monthly):
    assert query(monthly, warehouse_id=999999)['total'] == 0
    for params in [{'start_month':'2026-13'}, {'start_month':'2026-11','end_month':'2026-09'}]:
        response = monthly['client'].get('/report/api/warehouse_monthly', query_string={'warehouse_id':monthly['a'], **params})
        assert response.status_code == 400
    with m.app.app_context(): m.Warehouse.query.update({'is_default':False}); m.db.session.commit()
    assert monthly['client'].get('/report/api/warehouse_monthly').status_code == 400


def test_pagination_sorting_and_full_summary(monthly):
    full = query(monthly)
    first = query(monthly, page_size=2, sort_field='ending_quantity', sort_order='desc')
    pages = [query(monthly,page_size=2,page=i,sort_field='ending_quantity',sort_order='desc') for i in range(1,first['total_pages']+1)]
    joined = [r for p in pages for r in p['data']]
    assert len(joined) == len(full['data']) == first['total']
    assert [r['ending_quantity'] for r in joined] == sorted([r['ending_quantity'] for r in joined], reverse=True)
    assert all(p['summary'] == full['summary'] for p in pages)


def test_export_and_template_print_use_full_rows(monthly):
    expected = query(monthly)['data']
    for url in ['/report/api/warehouse_monthly','/report/warehouse_monthly/print_excel']:
        response = monthly['client'].get(url,query_string={'warehouse_id':monthly['a'],'start_month':'2026-09','end_month':'2026-10','export':'excel','page_size':1})
        assert response.status_code == 200
        book = load_workbook(io.BytesIO(response.data),read_only=True)
        try: values=list(book.active.iter_rows(values_only=True))
        finally: book.close()
        for row in expected: assert any(row['month'] in cells and row['material_code'] in cells for cells in values)


def test_monthly_link_matches_ledger(monthly):
    row = rows_for(query(monthly), 'A')['2026-10']
    args=parse_qs(urlparse(row['ledger_url']).query)
    assert args['material_id'] == [str(monthly['mats']['A'])]
    assert args['start_date']==['2026-10-01'] and args['end_date']==['2026-10-31']
    data=monthly['client'].get(row['ledger_url'].replace('/view/','/api/')).get_json()
    assert next(r for r in data['data'] if r['reference_type']=='本期合计')['balance_quantity']==85


def test_page_month_filters_and_unit_cards(monthly):
    page=monthly['client'].get('/report/view/warehouse_monthly').get_data(as_text=True)
    assert 'type="month"' in page and 'id="start_month"' in page and 'id="end_month"' in page
    assert 'id="summary-beginning_quantity"' in page and 'id="summary-ending_quantity"' in page
    assert 'data-month-range="previous"' in page and 'data-date-range="today"' not in page
    assert 'monthly-stock-table' in page
    data=query(monthly)
    assert not {'remark','transaction_count'}.intersection(c['field'] for c in data['columns'])
    assert 'remark' in data['data'][0] and 'transaction_count' in data['data'][0]


def test_month_query_is_readonly(monthly):
    writes=[]
    with m.app.app_context(): engine=m.db.engine
    def before(_c,_cursor,sql,_params,_context,_many):
        if sql.lstrip().upper().startswith(('INSERT','UPDATE','DELETE','ALTER','CREATE')): writes.append(sql)
    event.listen(engine,'before_cursor_execute',before)
    try: query(monthly)
    finally: event.remove(engine,'before_cursor_execute',before)
    assert writes==[]


def test_monthly_never_truncates_at_ledger_limit(monthly):
    with m.app.app_context():
        m.db.session.execute(m.StockTransaction.__table__.insert(), [dict(
            material_id=monthly['mats']['A'], warehouse_id=monthly['empty'],
            quantity=.01, transaction_type='in', created_at=datetime(2026,10,12)) for _ in range(50003)])
        m.db.session.commit()
    data=query(monthly, warehouse_id=monthly['empty'])
    assert data['total']==1 and not data['truncated']
    assert data['data'][0]['ending_quantity']==pytest.approx(500.03)
