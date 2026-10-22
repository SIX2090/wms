"""Run reachable inbound/outbound list actions with real API and binary downloads."""
import base64
from functools import lru_cache
from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess

import pytest

from tests.test_frontend_remaining_page_transport import callback

ROOT = Path(__file__).resolve().parents[1]
ACTIONS = ['in_batch_complete','in_batch_delete','in_print','in_export','in_copy','in_import','out_batch_complete','out_batch_delete','out_single_complete','out_single_delete','out_single_revert','out_import']
CASES = [(a,s) for a in ACTIONS for s in ['success','business','http','html','invalid_json','null','invalid_body','expired401','expired419','network']]
CASES += [(a,'cancel') for a in ACTIONS if not a.endswith('import')]
CASES += [(a,s) for a in ACTIONS if 'batch_' in a or a in ['in_print','in_export'] for s in ['empty','over_limit']]
CASES += [('in_print','multiple'),('in_import','empty'),('out_import','empty'),('in_copy','fallback'),('in_export','fallback'),('in_export','octet'),('in_print','missing_redirect'),('in_batch_complete','warning'),('out_batch_complete','warning')]


@lru_cache(maxsize=1)
def workbook_bytes():
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(['单据编号','数量'])
    book.active.append(['IN-FIXTURE',3])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def app_function(name):
    source = (ROOT / 'app/static/js/app.js').read_text(encoding='utf-8')
    start = source.index('function ' + name + '(')
    end = source.find('\nfunction ',start + 1)
    return source[start:end] if end >= 0 else source[start:]


@pytest.mark.parametrize('action,scenario', CASES)
def test_list_actions_preserve_protocol_confirmation_navigation_and_real_download(action,scenario):
    page = 'in_order' if action.startswith('in_') else 'out_order'
    name = 'submitDocumentImport' if action.endswith('import') else 'copyInOrder' if action=='in_copy' else 'singleOutOrderAction' if 'single_' in action else 'batchInOrderAction' if page=='in_order' else 'batchOutOrderAction'
    handler = callback(page,'function '+name+'(')
    program = r'''
const vm=require('node:vm'),assert=require('node:assert/strict'),p=JSON.parse(require('node:fs').readFileSync(0,'utf8'));
let sent=[],messages=[],timers=[],reloads=0,confirms=0,downloads=[],revoked=[],created=[];let scenario=p.scenario;
class FormData{constructor(form){this.owner=form.id;}}
const file={files:p.scenario==='empty'?[]:[{name:'fixture.xlsx'}],value:'fixture.xlsx',disabled:false};
const form={id:p.page==='in_order'?'inOrderImportForm':'outOrderImportForm',action:'/import/'+p.page};
const values=p.scenario==='empty'?[]:p.scenario==='over_limit'?Array.from({length:101},(_,i)=>String(i+1)):p.action==='in_print'?(p.scenario==='multiple'?['7','8']:['7','7']):['7','7','8'];
const bytes=Buffer.from(p.xlsx,'base64');const urlAPI={createObjectURL(blob){created.push(blob);return 'blob:fixture-'+created.length;},revokeObjectURL(u){revoked.push(u);}};
const ctx={FormData,Headers,Blob,URL:urlAPI,document:{querySelector:()=>({content:'fixture-csrf'}),querySelectorAll:s=>{assert.equal(s,'#'+(p.page==='in_order'?'inOrderTable':'outOrderTable')+' .check-item:checked');return values.map(value=>({value}));},getElementById:id=>id.endsWith('File')?file:form,
createElement:tag=>{assert.equal(tag,'a');return {href:'',download:'',click(){downloads.push({href:this.href,name:this.download});},remove(){}};},body:{appendChild(){}}},showToast:(m,t)=>messages.push({m,t}),showConfirm:()=>{confirms++;return Promise.resolve(p.scenario!=='cancel');},confirm:()=>false,setTimeout:fn=>timers.push(fn),location:{href:'original',reload(){reloads++;}},
fetch:(u,o)=>{sent.push({u,o});const s=scenario;const status=s==='http'?400:s==='html'?503:s==='expired401'?401:s==='expired419'?419:200;const binary=p.action==='in_export'&&['success','fallback','octet'].includes(s);
const body=s==='null'?null:s==='invalid_body'?'unexpected':{status:['business','http'].includes(s)?'error':s==='warning'?'warning':'success',msg:'fixture explanation',id:17,redirect_url:s==='fallback'?undefined:'/in_order/17?copied=1',redirect:s==='missing_redirect'?undefined:'/in_order/7/print',data:{mustPreserve:true}};
return s==='network'?Promise.reject(Error('offline')):Promise.resolve({ok:status===200,status,headers:new Headers({'Content-Type':binary?(s==='octet'?'application/octet-stream':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'):s==='html'?'text/html':'application/json','Content-Disposition':s==='fallback'?'':"attachment; filename*=UTF-8''%E5%85%A5%E5%BA%93%E5%8D%95_%E5%AF%BC%E5%87%BA.xlsx"}),
json:()=>binary?Promise.reject(SyntaxError('xlsx is not JSON')):['html','invalid_json'].includes(s)?Promise.reject(SyntaxError('invalid JSON')):Promise.resolve(body),blob:()=>{assert(binary,'error bodies must never be downloaded');return Promise.resolve(new Blob([bytes]));}});}};
ctx.window=ctx;vm.createContext(ctx);vm.runInContext(p.api,ctx);vm.runInContext(p.helpers+'\n'+p.selected+'\n'+p.handler,ctx);
function actionUrl(){if(p.action==='in_print')return '/in_order/batch_print';if(p.action==='in_export')return '/in_order/batch_export';return '/'+p.page+'/'+(p.action.includes('single_')?'7/'+p.action.split('single_')[1]:p.action.includes('batch_')?'batch_'+p.action.split('batch_')[1]:'');}
async function invoke(){if(p.action.endsWith('import'))ctx[p.name](form.id,p.page==='in_order'?'inOrderImportFile':'outOrderImportFile');else if(p.action==='in_copy')ctx[p.name](7);else ctx[p.name](actionUrl(),'empty selection','confirm');await new Promise(r=>setImmediate(r));await new Promise(r=>setImmediate(r));}
(async()=>{await invoke();
if(['empty','cancel','over_limit','multiple'].includes(p.scenario)){assert.equal(sent.length,0);assert.equal(downloads.length,0);assert.equal(reloads,0);if(p.scenario!=='cancel'&&!p.action.endsWith('import'))assert(messages.some(m=>m.t==='warning'));if(p.scenario==='multiple')assert(messages.some(m=>m.m.includes('1张')));return;}
assert.equal(sent.length,1);const req=sent[0],h=new Headers(req.o.headers);assert.equal(req.o.method,'POST');assert.equal(req.o.credentials,'same-origin');assert.equal(h.get('X-CSRFToken'),'fixture-csrf');assert.equal(h.get('X-Requested-With'),'XMLHttpRequest');
assert.equal(req.u,p.action.endsWith('import')?form.action:p.action==='in_copy'?'/in_order/7/copy':actionUrl());
if(p.action.endsWith('import')){assert(req.o.body instanceof FormData);assert.equal(req.o.body.owner,form.id);assert(!h.has('Content-Type'));assert.equal(file.value,'');assert.equal(file.disabled,false);}
else if(p.action==='in_copy')assert.deepEqual(JSON.parse(req.o.body),{});else if(p.action.includes('batch_')||p.action==='in_export'||p.action==='in_print')assert.deepEqual(JSON.parse(req.o.body),{ids:p.action==='in_print'?['7']:['7','8']},'dedupe detail rows and preserve ids without force');
else assert.equal(req.o.body,undefined,'single operations keep empty request body');
const ok=['success','fallback','octet'].includes(p.scenario);
if(ok){
 if(p.action==='in_print'){assert.equal(ctx.location.href,'/in_order/7/print');assert.equal(reloads,0);}
 else if(p.action==='in_export'){assert.equal(downloads.length,1);assert.equal(created.length,1);assert.equal(downloads[0].name,p.scenario==='fallback'?'入库单批量导出.xlsx':'入库单_导出.xlsx');assert.deepEqual(Buffer.from(await created[0].arrayBuffer()),bytes,'download actual workbook bytes');assert.equal(reloads,0);timers.forEach(fn=>fn());assert.deepEqual(revoked,['blob:fixture-1']);process.stdout.write(JSON.stringify({download:p.xlsx}));}
 else {assert(messages.some(m=>m.t==='success'));timers.forEach(fn=>fn());if(p.action==='in_copy')assert.equal(ctx.location.href,p.scenario==='fallback'?'/in_order/17':'/in_order/17?copied=1');else assert.equal(reloads,1);}
}else{assert(messages.some(m=>m.t==='danger'));assert.equal(reloads,0);assert.equal(ctx.location.href,'original');assert.equal(downloads.length,0);assert.equal(created.length,0);if(['business','http','warning'].includes(p.scenario))assert(messages.some(m=>m.m.includes('fixture explanation')));
 if(p.scenario==='html'&&p.action!=='in_export')assert(messages.some(m=>m.m.includes('503')));
 scenario='success';await invoke();assert.equal(sent.length,2,'failure permits retry');if(p.action==='in_export')assert.equal(downloads.length,1);else if(p.action==='in_print')assert.equal(ctx.location.href,'/in_order/7/print');else assert.equal(messages.at(-1).t,'success');}
})().catch(e=>{process.stderr.write(e.stack);process.exitCode=1;});
'''
    node=shutil.which('node')
    assert node,'Node required for actual API and download verification'
    payload=dict(action=action,scenario=scenario,page=page,name=name,handler=handler,
                 selected=callback(page,'function selectedOrderIds('),helpers=app_function('filenameFromDisposition')+'\n'+app_function('downloadBlob'),
                 api=(ROOT/'app/static/js/api.js').read_text(encoding='utf-8'),xlsx=base64.b64encode(workbook_bytes()).decode())
    result=subprocess.run([node,'-e',program],input=json.dumps(payload),capture_output=True,text=True,encoding='utf-8',timeout=15)
    assert result.returncode==0,result.stdout+result.stderr
    if result.stdout:
        from openpyxl import load_workbook
        downloaded=load_workbook(BytesIO(base64.b64decode(json.loads(result.stdout)['download'])))
        assert downloaded.active['A2'].value=='IN-FIXTURE'
        assert downloaded.active['B2'].value==3


def test_lists_do_not_create_a_copy_button_for_an_unreachable_outbound_handler():
    out=(ROOT/'app/templates/out_order.html').read_text(encoding='utf-8')
    assert 'copyOutOrder' not in out
    detail=(ROOT/'app/templates/out_order_detail.html').read_text(encoding='utf-8')
    assert 'onclick="copyOrder(' in detail
    source=(ROOT/'app/templates/in_order.html').read_text(encoding='utf-8')
    assert '打印已选（单张）' in source
    assert '多单合并打印暂未支持' in source
