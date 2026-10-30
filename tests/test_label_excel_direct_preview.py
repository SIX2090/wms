"""Actual saved Excel labels preview and print in the unified label selector."""
import json
import re
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app import app, db, ExcelPrintTemplate, Material, User


@pytest.fixture()
def excel_label_case(tmp_path, monkeypatch):
    monkeypatch.setitem(app.config, 'TESTING', True)
    monkeypatch.setitem(app.config, 'WTF_CSRF_ENABLED', False)
    monkeypatch.setattr(app, 'static_folder', str(tmp_path))
    folder = tmp_path / 'uploads/print_templates'
    folder.mkdir(parents=True)
    path = folder / 'label.xlsx'
    book = Workbook()
    sheet = book.active
    sheet['A1'] = '{item.code}'
    sheet['B1'] = '{item.brand}'
    sheet['C1'] = '{img_qrcode:item.code}'
    sheet['A2'] = '{item.name}'
    sheet.merge_cells('A2:B2')
    sheet['A3'] = '{item.spec}'
    sheet.merge_cells('A3:B3')
    sheet.merge_cells('C1:C3')
    for row, height in [(1, 20), (2, 40), (3, 40)]:
        sheet.row_dimensions[row].height = height
    sheet['A2'].font = Font(size=16, bold=True)
    sheet['A2'].alignment = Alignment(wrap_text=True)
    sheet.print_area = 'A1:C3'
    sheet['D5'] = 'excluded by print area'
    book.save(path)
    book.close()
    with app.app_context():
        assert str(db.engine.url) == 'sqlite:///:memory:'
        db.session.remove()
        db.drop_all()
        db.create_all()
        user = User(username='excel-direct', role='admin', password_hash='unused', must_change_password=False)
        materials = [Material(code='EX / A&B', brand='品牌甲', name='<b>名称甲</b>', spec='规格甲'),
                     Material(code='EX-002', brand='品牌乙', name='名称乙', spec='规格乙')]
        template = ExcelPrintTemplate(name='Excel三行二维码', target_type='label', target_code='material_label',
                                      excel_template_path='/static/uploads/print_templates/label.xlsx')
        db.session.add_all([user, template, *materials])
        db.session.commit()
        uid, tid, mids = user.id, template.id, ','.join(str(m.id) for m in materials)
    client = app.test_client()
    with client.session_transaction() as session:
        session['_user_id'], session['_fresh'] = str(uid), True
    yield client, tid, mids, path
    with app.app_context():
        assert Material.query.count() == 2
        db.session.remove()
        db.drop_all()


def test_build_excel_label_preview(excel_label_case):
    client, tid, mids, path = excel_label_case
    before = path.read_bytes()
    response = client.get('/label/excel_preview', query_string={
        'ids': mids, 'template_id': tid, 'qr_base_url': 'https://wms.example'})
    assert response.status_code == 200
    data = response.json
    assert data['status'] == 'success'
    assert len(data['labels']) == 2
    first = data['labels'][0]
    assert (first['width'], first['height']) == (90, 50)
    texts = [cell['text'] for cell in first['cells']]
    assert 'EX / A&B' in texts and '品牌甲' in texts and '<b>名称甲</b>' in texts and '规格甲' in texts
    assert 'excluded by print area' not in texts
    name = next(cell for cell in first['cells'] if cell['text'] == '<b>名称甲</b>')
    assert 'bold' in name['css'] and 'pre-wrap' in name['css']
    assert name['width'] > first['cells'][0]['width']
    image = next(cell['image'] for cell in first['cells'] if cell.get('image'))
    assert image.startswith('data:image/png;base64,')
    from services.material_label import material_label_qr
    with app.test_request_context('/'):
        expected = material_label_qr('EX / A&B', 'https://wms.example')
    assert expected['qrcode'] == 'https://wms.example/mobile/material?code=EX+%2F+A%26B'
    assert image == 'data:image/png;base64,' + expected['qrcode_png']
    assert path.read_bytes() == before
    assert not response.headers.get('Content-Disposition')
    assert '名称乙' in [cell['text'] for cell in data['labels'][1]['cells']]


def test_excel_preview_never_falls_back_for_missing_or_wrong_template(excel_label_case):
    client, tid, mids, path = excel_label_case
    assert app.test_client().get('/label/excel_preview').status_code == 302
    assert client.get('/label/excel_preview', query_string={'ids': mids, 'template_id': 99999}).status_code == 404
    path.unlink()
    response = client.get('/label/excel_preview', query_string={'ids': mids, 'template_id': tid})
    assert response.status_code == 400 and response.json['status'] == 'error'
    assert '模板文件' in response.json['msg']


def test_excel_formula_is_visible_failure_instead_of_wrong_print(excel_label_case):
    from openpyxl import load_workbook
    client, tid, mids, path = excel_label_case
    book = load_workbook(path)
    book.active['A1'] = '=1+1'
    book.save(path)
    book.close()
    response = client.get('/label/excel_preview', query_string={'ids': mids, 'template_id': tid})
    assert response.status_code == 400 and '公式' in response.json['msg']


def test_excel_preview_keeps_embedded_images_and_enforces_label_scope(excel_label_case):
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image
    client, tid, mids, path = excel_label_case
    logo_path = path.parent / 'logo.png'
    Image.new('RGB', (12, 10), 'black').save(logo_path)
    book = load_workbook(path)
    book.active.add_image(XLImage(logo_path), 'B1')
    book.save(path)
    book.close()
    response = client.get('/label/excel_preview', query_string={'ids': mids, 'template_id': tid})
    assert response.status_code == 200
    assert len(response.json['labels'][0]['images']) == 1
    assert response.json['labels'][0]['images'][0]['src'].startswith('data:image/png;base64,')
    assert client.get('/label/excel_preview', query_string={'ids': '', 'template_id': tid}).status_code == 400
    assert client.get('/label/excel_preview', query_string={'ids': mids, 'template_id': tid,
                                                         'qr_base_url': 'javascript:bad'}).status_code == 400
    with app.app_context():
        db.session.get(ExcelPrintTemplate, tid).target_type = 'report'
        db.session.commit()
    assert client.get('/label/excel_preview', query_string={'ids': mids, 'template_id': tid}).status_code == 404


def test_unified_selector_switches_excel_and_back_and_blocks_stale_print(excel_label_case):
    from tests.test_batch_label_template_loading_runtime import Nodes, PROGRAM
    client, _, mids, _ = excel_label_case
    html = client.get('/label/batch_print', query_string={'ids': mids}).text
    # Only one template selector belongs to the main print toolbar.
    toolbar = re.search(r'<div id="labelPrintToolbar"[^>]*>(.*?)</div>', html, re.S)[1]
    assert [attrs.get('id') for tag, attrs in Nodes(toolbar).nodes if tag == 'select'] == ['templateSelect']
    export = re.search(r'<details id="excelExportOptions"[^>]*>(.*?)</details>', html, re.S)[1]
    assert 'id="excelTemplateSelect"' in export
    branch = r'''
 if(p.action==='unified_excel'){
  await finish(excel,{status:'success',templates:[{id:19,name:'Excel三行二维码',has_file:true}]});
  await finish(list,{status:'success',templates:[]});
  assert(ids.templateSelect.textContent.includes('Excel三行二维码'));
  ctx.loadExcelTemplates();await finish(excel,{status:'success',templates:[{id:19,name:'Excel三行二维码',has_file:true}]});
  assert.equal(all(ids.templateSelect,'option').filter(node=>node.value==='excel:19').length,1,'refresh must not duplicate options');
  select('excel:19');
  assert.equal(ids.selectedLabelTemplate.value,'excel:19','QR address update preserves selected template');
  const preview=calls.find(call=>call.url.startsWith('/label/excel_preview?')).url;
  assert(preview.includes('template_id=19') && preview.includes('ids=1%2C2'));
  assert.equal(ctx.printLabels(),false);
  const data={status:'success',labels:[{width:90,height:50,scale:1,content_width:100,content_height:40,cells:[{text:'<b>real Excel name</b>',css:'font-weight:bold',left:0,top:0,width:100,height:40}],images:[]}]};
  await finish(preview,data);
  assert.equal(ids.labelsContainer.getAttribute('data-wms-print-state'),'ready');
  assert(ids.labelsContainer.textContent.includes('<b>real Excel name</b>'));
  assert.equal(ctx.printLabels(),true);assert.equal(printed,1);assert.equal(opened.length,0);
  ctx.loadTemplates();await finish(list,{status:'success',templates:[]});await finish(preview,data);
  assert.equal(ids.templateSelect.value,'excel:19','refresh keeps the selected Excel template');
  select('');assert(ids.labelsContainer.textContent.includes('<b>名称甲</b>'));
  select('excel:19');select('');await finish(preview,data);
  assert(ids.labelsContainer.textContent.includes('<b>名称甲</b>'),'late Excel response cannot replace default');
  select('excel:19');await finish(preview,{status:'error',msg:'模板文件损坏'},'business');
  assert.equal(ids.labelsContainer.getAttribute('data-wms-print-state'),'error');
  assert.equal(ctx.printLabels(),false);assert.equal(printed,1);assert.equal(opened.length,0);
  return;
 }
'''
    program = PROGRAM.replace(" if(p.action==='token')", branch + " if(p.action==='token')")
    root = Path(__file__).resolve().parents[1]
    payload = dict(action='unified_excel', nodes=Nodes(html).nodes,
                   scripts=[body for attrs, body in re.findall(r'<script([^>]*)>(.*?)</script>', html, re.S) if 'src=' not in attrs],
                   api=(root / 'app/static/js/api.js').read_text(encoding='utf-8'))
    result = subprocess.run([shutil.which('node'), '-e', program], input=json.dumps(payload),
                            capture_output=True, text=True, encoding='utf-8', timeout=20)
    assert result.returncode == 0, result.stderr
