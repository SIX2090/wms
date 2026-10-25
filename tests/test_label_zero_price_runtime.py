"""BUG-2026-10-25-007: zero prices survive actual label render and Excel GETs."""
import io
import importlib.util
import json
import re
import shutil
import subprocess
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace as NS

import pytest
from jinja2 import ChoiceLoader, DictLoader, Environment, FileSystemLoader
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
_harness_spec = importlib.util.spec_from_file_location(
    'zero_price_batch_harness', ROOT / 'tests/test_batch_label_template_loading_runtime.py')
_harness = importlib.util.module_from_spec(_harness_spec)
_harness_spec.loader.exec_module(_harness)
Nodes, BATCH_PROGRAM = _harness.Nodes, _harness.PROGRAM


def _values(html):
    return re.findall(r'<span class="cell-value">(.*?)</span>', html)


@pytest.mark.parametrize('price,expected', [
    (0, '0'), (0.0, '0.0'), (Decimal('0'), '0'),
    (None, ''), ('', ''), ('0', '0'), (12.5, '12.5'),
])
def test_actual_regular_template_preserves_zero_and_empty_price(price, expected):
    env = Environment(loader=ChoiceLoader([
        DictLoader({'base.html': '{% block content %}{% endblock %}'}),
        FileSystemLoader(ROOT / 'app/templates'),
    ]), autoescape=True)
    env.filters['from_json'] = json.loads
    template = NS(id=4, name='零价测试', width=60, height=40, rows=1, cols=1,
                  layout=json.dumps({'0-0': {'field': 'price'}}))
    html = env.get_template('print_label.html').render(
        template=template, templates=[template],
        materials=[NS(id=7, code='ZERO', name='测试物料', price=price)])
    assert _values(html) == [expected]


@pytest.mark.parametrize('price,expected', [
    (0, '¥0'), (0.0, '¥0'), ('0', '¥0'),
    (None, ''), ('', ''), ('missing', ''), (12.5, '¥12.5'),
])
def test_complete_batch_page_and_http_client_render_print_time_price(price, expected):
    material = {'id': 7, 'code': 'ZERO', 'name': '测试物料'}
    if price != 'missing':
        material['price'] = price
    env = Environment(loader=FileSystemLoader(ROOT / 'app/templates'), autoescape=True)
    html = env.get_template('print_batch_labels.html').render(
        materials=[material], materials_data=[material],
        request=NS(args={}), current_user=NS(is_authenticated=True), config={},
        url_for=lambda *args, **kwargs: '/fixture')
    scripts = [text for attrs, text in re.findall(r'<script([^>]*)>(.*?)</script>', html, re.S)
               if 'src=' not in attrs]
    # Reuse the established DOM/HTTP harness, adding only this independent scenario.
    branch = r'''
 if(p.action==='zero_price'){
  await finish(excel,{status:'success',templates:[]});
  await finish(list,{status:'success',templates:[{id:'A',name:'Price',width:60,height:40}]});
  await finish(detail('A'),{status:'success',template:{id:'A',width:60,height:40,rows:1,cols:1,layout:{cells:[{row:0,col:0,field:'price'}]}}});
  const card=ids.labelsContainer.children[0];
  assert.equal(ids.labelsContainer.children.length,1);
  assert.equal(all(card,'td')[0].textContent,p.expected,'actual printed price');
  const originalPrint=ctx.print;
  ctx.print=()=>{
   assert.equal(all(ids.labelsContainer.children[0],'td')[0].textContent,p.expected,'price at window.print');
   originalPrint();
  };
  assert.equal(ctx.printLabels(),true);
  assert.equal(printed,1);
  assert.equal(calls.length,3);
  for(const call of calls){assert.equal(call.options.method,'GET');assert.equal(call.options.credentials,'same-origin');}
  return;
 }
'''
    program = BATCH_PROGRAM.replace(" if(p.action==='token'){", branch + " if(p.action==='token'){")
    assert program != BATCH_PROGRAM
    payload = dict(action='zero_price', expected=expected, nodes=Nodes(html).nodes, scripts=scripts,
                   api=(ROOT / 'app/static/js/api.js').read_text(encoding='utf-8'))
    node = shutil.which('node')
    assert node, 'Node is required for the complete label page runtime'
    result = subprocess.run([node, '--max-old-space-size=64', '-e', program],
                            input=json.dumps(payload), capture_output=True, text=True,
                            encoding='utf-8', timeout=20)
    assert result.returncode == 0, result.stdout + result.stderr


@pytest.fixture
def label_price_gets(tmp_path, monkeypatch):
    from app import app, db, User, Material, LabelTemplate, ExcelPrintTemplate

    monkeypatch.setitem(app.config, 'TESTING', True)
    monkeypatch.setitem(app.config, 'WTF_CSRF_ENABLED', False)
    static = tmp_path / 'static'
    template_dir = static / 'uploads' / 'print_templates'
    template_dir.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.append(['Code', 'Price'])
    workbook.active.append(['{item.code}', '{item.price}'])
    workbook.save(template_dir / 'zero-price.xlsx')
    monkeypatch.setattr(app, 'static_folder', str(static))
    with app.app_context():
        assert str(db.engine.url) == 'sqlite:///:memory:'
        db.session.remove()
        db.drop_all()
        db.create_all()
        user = User(username='label-zero-fixture', role='admin',
                    password_hash='unused', must_change_password=False)
        materials = [Material(code='ZERO', name='零价', price=Decimal('0')),
                     Material(code='NONE', name='未填写', price=1),
                     Material(code='POS', name='正价', price=12.5)]
        template = LabelTemplate(name='零价', width=60, height=40, rows=1, cols=1,
                                 layout=json.dumps({'0-0': {'field': 'price'}}))
        excel = ExcelPrintTemplate(name='零价', target_type='label', target_code='material_label',
                                   template_type='excel', is_default=True,
                                   excel_template_path='/static/uploads/print_templates/zero-price.xlsx')
        db.session.add_all([user, *materials, template, excel])
        db.session.commit()
        # Set after INSERT so the model's default=0 does not replace explicit None.
        materials[1].price = None
        db.session.commit()
        ids = [material.id for material in materials]
        user_id, template_id, excel_id = user.id, template.id, excel.id
        before = [(material.id, material.price) for material in materials]
    client = app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(user_id)
        session['_fresh'] = True
    yield client, ids, template_id, excel_id
    with app.app_context():
        after = [(db.session.get(Material, mid).id, db.session.get(Material, mid).price)
                 for mid in ids]
        assert after == before, 'label GETs must not change prices'
        db.session.remove()
        db.drop_all()


def test_actual_regular_print_get_keeps_zero_and_none_distinct(label_price_gets):
    client, ids, template_id, _ = label_price_gets
    response = client.get(f'/label_template/{template_id}/print?ids=' + ','.join(map(str, ids)))
    assert response.status_code == 200
    assert _values(response.get_data(as_text=True)) == ['0.0', '', '12.5']


def test_actual_batch_get_serializes_zero_without_replacing_empty_price(label_price_gets):
    client, ids, _, _ = label_price_gets
    response = client.get('/label/batch_print?ids=' + ','.join(map(str, ids)))
    assert response.status_code == 200
    binding = re.search(r'var MATERIALS\s*=\s*(.*?);', response.get_data(as_text=True))
    assert binding
    data = json.loads(binding.group(1))
    assert [(row['code'], row['price']) for row in data] == [
        ('ZERO', '0.0'), ('NONE', ''), ('POS', '12.5')]


def test_actual_excel_get_fills_zero_in_real_workbook(label_price_gets):
    client, ids, _, excel_id = label_price_gets
    response = client.get('/label/batch_print_excel?ids=' + ','.join(map(str, ids))
                          + f'&template_id={excel_id}')
    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    workbook = load_workbook(io.BytesIO(response.data))
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows == [('Code', 'Price'), ('ZERO', '0.0'), ('NONE', None), ('POS', '12.5')]


def test_existing_preview_remains_consistent_with_actual_print(label_price_gets):
    client, ids, template_id, _ = label_price_gets
    response = client.get(f'/label_template/{template_id}/preview?material_id={ids[0]}')
    assert response.status_code == 200
    assert _values(response.get_data(as_text=True)) == ['0.0']
