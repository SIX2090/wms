# -*- coding: utf-8 -*-
"""ARCH-OS-IMPORT 回归：期初库存"导入导出模板 / 批量导入 / 粘贴导入"补齐。

历史问题（用户反馈"导入导出模板、粘贴导入根本没法用"）：
- "导入导出模板"按钮只弹提示，不生成/下载模板文件；
- "批量导入"跳 /batch_import?type=opening_stock，但批量导入页无 opening_stock
  卡片，后端 /opening_stock/import 仅重定向 stub，无真实导入逻辑；
- "粘贴导入"用 prompt() 单行输入，纯前端 push 不入库，且走原生 fetch。

修复后要求：
- GET /opening_stock/import/template 返回可用 xlsx 模板（表头 + 示例行）；
- POST /opening_stock/import 解析 xlsx 并经 batch_save 同一校验/入账路径写入
  OpeningStock（复用 _apply_opening_stock_balance），示例行/缺列/不存在仓/重复键
  正确跳过并提示，含 require_role('warehouse') 角色校验；
- 前端"导入导出模板"按钮指向模板下载、粘贴导入改多行 modal 并经 batch_save 入库。
"""
from __future__ import annotations

import io
import os
import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
APP_DIR = ROOT / "app"
sys.path.insert(0, str(APP_DIR))

os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ.setdefault("WMS_BOOTSTRAP_PASSWORD", "admin")
os.environ["WMS_DEBUG"] = "0"

from werkzeug.security import generate_password_hash  # noqa: E402

import app as app_module  # noqa: E402
from app import (  # noqa: E402
    Material, OpeningStock, Unit, User, Warehouse, db,
)


def _login(client):
    page = client.get("/login").get_data(as_text=True)
    token = re.search(r'name="csrf_token".*?value="([^"]+)"', page)
    client.post("/login", data={
        "username": "admin", "password": "admin",
        "csrf_token": token.group(1) if token else "",
    })


@pytest.fixture()
def client():
    app_module.app.config["WTF_CSRF_ENABLED"] = False
    app_module.app.config["TESTING"] = True
    with app_module.app.app_context():
        db.drop_all()
        db.create_all()
        if not User.query.filter_by(username="admin").first():
            db.session.add(User(
                username="admin",
                password_hash=generate_password_hash("admin"),
                role="admin", must_change_password=False,
            ))
        db.session.add_all([
            Unit(code="TAO", name="套"),
            Warehouse(code="WH001", name="项目仓", status="active", is_default=True),
        ])
        db.session.commit()
        db.session.add(Material(code="M-0001", name="轴承6204", stock=0, price=25.5, unit_id=1))
        db.session.commit()
    c = app_module.app.test_client()
    _login(c)
    yield c


def _make_xlsx(rows, headers=("仓库编码", "物料编码", "物料名称", "规格", "单位", "数量", "单价", "备注")):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_import_template_downloads_xlsx(client):
    """导入导出模板按钮对应的后端路由返回真实 xlsx（表头 + 示例行）。"""
    rv = client.get("/opening_stock/import/template")
    assert rv.status_code == 200
    assert "spreadsheetml" in rv.headers.get("Content-Type", "")
    assert len(rv.data) > 1000
    # 用 openpyxl 读回校验表头
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(rv.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert "物料编码" in header and "数量" in header and "仓库编码" in header


def test_excel_import_writes_opening_stock(client):
    """Excel 批量导入经统一校验写入台账；示例行与不存在仓跳过。"""
    xlsx = _make_xlsx([
        ["WH001", "M-0001", "示例-轴承6204", "", "套", "999", "1", "示例行应跳过"],
        ["WH001", "M-0001", "轴承6204", "", "套", "100", "25.50", "有效行"],
        ["WH-NOEXIST", "M-0001", "轴承6204", "", "套", "50", "10", "不存在仓应跳过"],
    ])
    rv = client.post("/opening_stock/import", data={
        "file": (xlsx, "t.xlsx"),
    }, content_type="multipart/form-data")
    assert rv.status_code == 200
    data = rv.get_json()
    assert data["status"] == "success", data
    assert data["imported"] == 1
    # 示例行属模板说明，静默忽略不计 skipped；仅"不存在仓"计入 skipped
    assert data["skipped"] == 1
    with app_module.app.app_context():
        m = Material.query.filter_by(code="M-0001").first()
        recs = OpeningStock.query.filter_by(material_id=m.id).all()
        assert len(recs) == 1
        assert abs(recs[0].quantity - 100) < 1e-6
        assert abs(recs[0].price - 25.5) < 1e-6


def test_excel_import_missing_required_column_rejected(client):
    """缺"物料编码"或"数量"列时返回明确错误。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["仓库编码", "物料名称"])  # 缺 物料编码/数量
    ws.append(["WH001", "轴承6204"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    rv = client.post("/opening_stock/import", data={
        "file": (buf, "bad.xlsx"),
    }, content_type="multipart/form-data")
    assert rv.status_code == 400
    data = rv.get_json()
    assert data["status"] == "error"
    assert "物料编码" in data["msg"] or "数量" in data["msg"]


def test_excel_import_missing_warehouse_skipped(client):
    """期初建账仓库必填：未指定仓库的行报错跳过，不静默默认仓回落。"""
    xlsx = _make_xlsx([
        ["", "M-0001", "轴承6204", "", "套", "30", "5", "缺仓库"],
    ])
    rv = client.post("/opening_stock/import", data={
        "file": (xlsx, "t.xlsx"),
    }, content_type="multipart/form-data")
    assert rv.status_code == 200
    data = rv.get_json()
    assert data["status"] == "success"
    assert data["imported"] == 0
    assert data["skipped"] == 1
    assert any("仓库" in e for e in data["errors"])


def test_frontend_buttons_point_to_real_capability():
    """前端"导入导出模板"指向模板下载、粘贴导入改 modal、不再用 prompt。"""
    src = (APP_DIR / "templates" / "opening_stock.html").read_text(encoding="utf-8")
    assert "/opening_stock/import/template" in src, "模板按钮应指向模板下载路由"
    assert "pasteImportModal" in src, "粘贴导入应使用多行 modal"
    assert "prompt(" not in src, "粘贴导入不应再用 prompt() 单行输入"
    assert "showImportHint" not in src, "导入导出模板不应再是提示占位"


# ---------------------------------------------------------------------------
# FIX-OS-DATE-001：新增期初库存单日期不可改
# 根因：表头 <input type="date"> 无 id，collectItems() 不传 date，saveDocument()
# 从不读取表头日期；粘贴导入 / Excel 导入同样丢弃日期。后端逐行
# _parse_opening_stock_date(item.get('date')) 恒得 None → 一律落"今天"。
# ---------------------------------------------------------------------------


def test_batch_save_honours_date(client):
    """POST /opening_stock/batch_save 带 date 时按指定日期落库（新增分支）。"""
    with app_module.app.app_context():
        m = Material.query.filter_by(code="M-0001").first()
        wh = Warehouse.query.filter_by(code="WH001").first()
        mid, wid = m.id, wh.id
    rv = client.post("/opening_stock/batch_save", json={"items": [
        {"material_id": mid, "warehouse_id": wid, "quantity": 100,
         "price": 25.5, "date": "2026-01-15"},
    ]})
    assert rv.status_code == 200
    assert rv.get_json()["status"] == "success"
    with app_module.app.app_context():
        rec = OpeningStock.query.filter_by(material_id=mid, warehouse_id=wid).first()
        assert rec is not None
        assert rec.date.isoformat() == "2026-01-15", f"实际日期={rec.date}"


def test_batch_save_update_changes_date(client):
    """已存在记录再次保存时，日期随请求更新（update 分支同样写 date）。"""
    with app_module.app.app_context():
        m = Material.query.filter_by(code="M-0001").first()
        wh = Warehouse.query.filter_by(code="WH001").first()
        mid, wid = m.id, wh.id
    client.post("/opening_stock/batch_save", json={"items": [
        {"material_id": mid, "warehouse_id": wid, "quantity": 100,
         "price": 25.5, "date": "2026-01-15"},
    ]})
    client.post("/opening_stock/batch_save", json={"items": [
        {"material_id": mid, "warehouse_id": wid, "quantity": 120,
         "price": 25.5, "date": "2026-03-31"},
    ]})
    with app_module.app.app_context():
        rec = OpeningStock.query.filter_by(material_id=mid, warehouse_id=wid).first()
        assert rec.date.isoformat() == "2026-03-31", f"实际日期={rec.date}"
        assert abs(rec.quantity - 120) < 1e-6


def test_batch_save_missing_date_falls_back_today(client):
    """不传 date 时回落当天（向后兼容：不破坏既有调用方）。"""
    import datetime as _dt
    with app_module.app.app_context():
        m = Material.query.filter_by(code="M-0001").first()
        wh = Warehouse.query.filter_by(code="WH001").first()
        mid, wid = m.id, wh.id
    client.post("/opening_stock/batch_save", json={"items": [
        {"material_id": mid, "warehouse_id": wid, "quantity": 10, "price": 1},
    ]})
    with app_module.app.app_context():
        rec = OpeningStock.query.filter_by(material_id=mid, warehouse_id=wid).first()
        assert rec.date == _dt.date.today()


def test_excel_template_has_date_column(client):
    """Excel 模板含"日期"列，供用户指定建账日期。"""
    rv = client.get("/opening_stock/import/template")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(rv.data))
    header = [c.value for c in wb.active[1]]
    assert "日期" in header, f"模板表头应含日期列，实际={header}"


def test_excel_import_honours_date(client):
    """Excel 导入按"日期"列建账（支持字符串与日期对象两种单元格类型）。"""
    import datetime as _dt
    xlsx = _make_xlsx(
        [
            ["WH001", "M-0001", "轴承6204", "", "套", "100", "25.50", "字符串日期", "2026-05-06"],
            ["WH001", "M-0001", "轴承6204", "", "套", "9", "1", "示例", "2026-05-06"],
        ],
        headers=("仓库编码", "物料编码", "物料名称", "规格", "单位", "数量", "单价", "备注", "日期"),
    )
    rv = client.post("/opening_stock/import", data={"file": (xlsx, "t.xlsx")},
                     content_type="multipart/form-data")
    assert rv.status_code == 200
    assert rv.get_json()["status"] == "success"
    with app_module.app.app_context():
        m = Material.query.filter_by(code="M-0001").first()
        rec = OpeningStock.query.filter_by(material_id=m.id).first()
        assert rec is not None
        assert rec.date.isoformat() == "2026-05-06", f"实际日期={rec.date}"
        assert isinstance(rec.date, _dt.date)


def test_frontend_date_wiring():
    """前端日期链路：表头日期输入有 id 并被 saveDocument/粘贴导入读取。"""
    src = (APP_DIR / "templates" / "opening_stock.html").read_text(encoding="utf-8")
    assert 'id="headerDocDate"' in src, "表头日期输入应具备 id 供 JS 读取"
    assert "getHeaderDocDate" in src, "应定义 getHeaderDocDate() 读取表头日期"
    assert "item.date = headerDate" in src, "saveDocument 应把表头日期注入明细行"
    assert 'id="pasteDocDate"' in src, "粘贴导入 modal 应提供建账日期输入"
    assert "parsePasteLines(text, document.getElementById('pasteDocDate')" in src, \
        "confirmPasteImport 应把 modal 日期传入解析器"
    assert "date: rowDate || defaultDate" in src, "解析器应回填行日期或默认日期"
