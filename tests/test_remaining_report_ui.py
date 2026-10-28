"""2026-10-28 browser observations: report scope, labels, source and print layout."""
from datetime import date, datetime
from pathlib import Path
import io
import shutil
import subprocess

from bs4 import BeautifulSoup
from openpyxl import load_workbook, Workbook
import pytest

import app as m


@pytest.fixture
def report_case(monkeypatch, tmp_path):
    monkeypatch.setitem(m.app.config, 'TESTING', True)
    monkeypatch.setitem(m.app.config, 'WTF_CSRF_ENABLED', False)
    monkeypatch.setattr(m.app, 'static_folder', str(tmp_path))
    with m.app.app_context():
        m.db.session.remove()
        m.db.drop_all()
        m.db.create_all()
        user = m.User(username='report-qa', role='admin', must_change_password=False,
                      password_hash='isolated-only')
        a = m.Warehouse(code='R-A', name='甲仓', status='active')
        b = m.Warehouse(code='R-B', name='乙仓', status='active')
        unit = m.Unit(code='R-PCS', name='个')
        supplier = m.Supplier(code='R-S', name='测试供应商')
        m.db.session.add_all([user, a, b, unit, supplier])
        m.db.session.flush()
        mat = m.Material(code='R-M', name='测试物料', unit_id=unit.id, brand=None)
        m.db.session.add(mat)
        m.db.session.commit()
        ids = dict(user=user.id, a=a.id, b=b.id, mat=mat.id, supplier=supplier.id)
    client = m.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(ids['user'])
        session['_fresh'] = True
    yield dict(client=client, root=tmp_path, **ids)
    with m.app.app_context():
        m.db.session.remove()


def test_opening_editor_uses_own_toolbar(report_case):
    response = report_case['client'].get('/opening_stock/add')
    soup = BeautifulSoup(response.data, 'html.parser')
    assert response.status_code == 200
    assert soup.select_one('.opening-toolbar')
    assert 'wms-no-global-actionbar' in soup.body.get('class', [])


def test_outbound_report_labels_cover_sales_and_requisition():
    definition = m._get_report_definition('out_detail')
    columns = {c['field']: c['title'] for c in m._out_detail_columns()}
    assert definition['title'] == '出库明细报表'
    assert definition['summary_labels']['quantity'] == '出库数量'
    assert columns['order_no'] == '出库单号'
    assert columns['customer'] == '客户/部门'


@pytest.mark.parametrize('kind', ['in_detail', 'out_detail'])
def test_detail_report_has_readable_horizontal_layout(report_case, kind):
    response = report_case['client'].get('/report/view/' + kind,
                                        query_string={'warehouse_id': report_case['a']})
    soup = BeautifulSoup(response.data, 'html.parser')
    assert response.status_code == 200
    table = soup.select_one('#reportTable.detail-report-table')
    assert table and table.parent.get('class') and 'table-responsive' in table.parent['class']


@pytest.mark.parametrize('state,received,expected', [('pending', 2, 2), ('completed', 1, 1),
                                                    ('completed', 3, 0)])
def test_home_receipts_use_completed_facts_not_push_counter(report_case, state, received, expected):
    c = report_case
    with m.app.app_context():
        po = m.PurchaseOrder(order_no='PO-R', supplier_id=c['supplier'], status='pending')
        m.db.session.add(po)
        m.db.session.flush()
        item = m.PurchaseOrderItem(purchase_order_id=po.id, material_id=c['mat'],
                                   quantity=2, received_quantity=2, price=1)
        m.db.session.add(item)
        m.db.session.flush()
        receipt = m.InOrder(order_no='IN-R', status=state, warehouse='甲仓',
                            source_purchase_order_id=po.id, business_type='采购入库')
        m.db.session.add(receipt)
        m.db.session.flush()
        m.db.session.add(m.InOrderItem(in_order_id=receipt.id, material_id=c['mat'],
                                      source_purchase_order_item_id=item.id, quantity=received, price=1, amount=received))
        m.db.session.commit()
        summary = m.build_purchase_order_todo_summary()
        assert summary['remaining_qty'] == 0  # Existing unpushed definition stays intact.
        assert summary.get('pending_receipt_qty') == expected
        assert summary['open_count'] == 1
    html = c['client'].get('/').get_data(as_text=True)
    assert '当前没有待入库采购单。' not in html
    assert '待收货数量 %.2f' % expected in html


def test_home_legacy_header_receipts_and_revert_are_read_only(report_case):
    c = report_case
    with m.app.app_context():
        po = m.PurchaseOrder(order_no='PO-R', supplier_id=c['supplier'], status='partial')
        m.db.session.add(po)
        m.db.session.flush()
        m.db.session.add(m.PurchaseOrderItem(purchase_order_id=po.id, material_id=c['mat'],
                                            quantity=5, received_quantity=5, price=1))
        receipt = m.InOrder(order_no='IN-R', status='completed', warehouse='',
                            source_purchase_order_id=po.id, business_type='采购入库')
        m.db.session.add(receipt)
        m.db.session.flush()
        m.db.session.add(m.InOrderItem(in_order_id=receipt.id, material_id=c['mat'], quantity=2,
                                      price=1, amount=2))
        m.db.session.commit()
        assert m.build_purchase_order_todo_summary().get('pending_receipt_qty') == 3
        assert m.InOrder.query.count() == 1 and m.StockTransaction.query.count() == 0
        receipt.status = 'pending'
        m.db.session.commit()
        assert m.build_purchase_order_todo_summary().get('pending_receipt_qty') == 5


@pytest.mark.parametrize('invalid', ['missing', 'foreign_warehouse', 'foreign_material'])
def test_ledger_opening_missing_or_mismatched_source_has_no_wrong_link(report_case, invalid):
    c = report_case
    with m.app.test_request_context('/'):
        doc = m.OpeningStockDoc(doc_no='QS-FOREIGN', warehouse_id=c['b'] if invalid == 'foreign_warehouse' else c['a'])
        m.db.session.add(doc)
        m.db.session.flush()
        line = m.OpeningStock(doc_id=doc.id, material_id=c['mat'], warehouse_id=doc.warehouse_id)
        if invalid == 'foreign_material':
            other = m.Material(code='OTHER', name='另一物料')
            m.db.session.add(other)
            m.db.session.flush()
            line.material_id = other.id
        m.db.session.add(line)
        m.db.session.flush()
        source = 99999 if invalid == 'missing' else line.id
        m.db.session.add(m.StockTransaction(material_id=c['mat'], warehouse_id=c['a'],
            location='甲仓', transaction_type='opening_stock', quantity=3,
            reference_type='opening_stock', reference_id=source, created_at=datetime(2026, 10, 1)))
        m.db.session.commit()
        _, rows, _ = m._build_ledger_report(dict(warehouse_id=c['a'], warehouse='甲仓',
            warehouse_code='R-A', start_date=None, end_date=None, material_code='R-M',
            sort_field='', sort_order='asc'))
        actual = next(row for row in rows if row.get('reference_type') == '期初库存')
        assert actual['reference_no'] == '期初记录 #%d' % source
        assert actual['reference_url'] == ''


@pytest.mark.parametrize('legacy', [False, True, 'mixed'])
def test_ledger_opening_line_resolves_document_or_friendly_legacy(report_case, legacy):
    c = report_case
    with m.app.test_request_context('/'):
        doc = m.OpeningStockDoc(doc_no='QS-R', warehouse_id=c['b'] if legacy == 'mixed' else c['a'], date=date(2026, 10, 1))
        m.db.session.add(doc)
        m.db.session.flush()
        line = m.OpeningStock(material_id=c['mat'], warehouse_id=c['a'],
                              doc_id=None if legacy is True else doc.id, quantity=3)
        m.db.session.add(line)
        m.db.session.flush()
        m.db.session.add(m.StockTransaction(material_id=c['mat'], warehouse_id=c['a'],
            location='甲仓', transaction_type='opening_stock', quantity=3,
            reference_type='opening_stock', reference_id=line.id, created_at=datetime(2026, 10, 1)))
        m.db.session.commit()
        filters = dict(warehouse_id=c['a'], warehouse='甲仓', warehouse_code='R-A',
                       start_date=None, end_date=None, material_code='R-M', sort_field='', sort_order='asc')
        _, rows, summary = m._build_ledger_report(filters)
        actual = next(row for row in rows if row.get('reference_type') == '期初库存')
        assert actual['reference_no'] == ('期初记录 #%d' % line.id if legacy is True else 'QS-R')
        assert actual['reference_url'] == ('' if legacy is True else '/opening_stock/%d' % doc.id)
        assert actual['balance_quantity'] == 3
        _, other, _ = m._build_ledger_report({**filters, 'warehouse_id': c['b'],
                                             'warehouse': '乙仓', 'warehouse_code': 'R-B'})
        assert not any(r.get('reference_no') == 'QS-R' for r in other)


@pytest.mark.parametrize('business', ['采购入库', '产品入库', '其他入库', '销售退货入库'])
def test_inbound_print_html_and_fallback_excel_follow_business(report_case, business):
    c = report_case
    with m.app.app_context():
        order = m.InOrder(order_no='IN-R', business_type=business, warehouse='甲仓', status='pending')
        m.db.session.add(order)
        m.db.session.flush()
        m.db.session.add(m.InOrderItem(in_order_id=order.id, material_id=c['mat'], quantity=2, price=0, amount=0))
        m.db.session.add(m.InOrderPrintTemplate(name='默认', template_type='excel', is_default=True))
        m.db.session.commit()
        oid = order.id
    html = c['client'].get('/in_order/%d/print' % oid).get_data(as_text=True)
    assert '<div class="doc-title">%s单</div>' % business in html
    assert '>None<' not in html
    response = c['client'].get('/in_order/%d/print_excel' % oid)
    wb = load_workbook(io.BytesIO(response.data))
    assert wb.active['A1'].value == business + '单'
    assert wb.active['D2'].value == business + '单号：IN-R'
    wb.close()


@pytest.mark.parametrize('managed', [True, False])
def test_inbound_managed_excel_retitle_is_in_memory_and_custom_untouched(report_case, managed):
    c = report_case
    name = 'builtin_in_order_default_v2.xlsx' if managed else 'custom.xlsx'
    path = c['root'] / name
    wb = Workbook()
    wb.active.append(['采购入库单'])
    wb.active['D2'] = '采购入库单号：{order.order_no}'
    wb.active['A3'] = '{item.material.code}'
    wb.save(path)
    before = path.read_bytes()
    with m.app.app_context():
        order = m.InOrder(order_no='PI-R', business_type='产品入库', status='pending', warehouse='甲仓')
        m.db.session.add(order)
        m.db.session.flush()
        m.db.session.add(m.InOrderItem(in_order_id=order.id, material_id=c['mat'], quantity=2, price=0, amount=0))
        m.db.session.add(m.InOrderPrintTemplate(name='模板', template_type='excel', is_default=True,
                                                excel_template_path='/static/' + name))
        m.db.session.commit()
        oid = order.id
    response = c['client'].get('/in_order/%d/print_excel' % oid)
    wb = load_workbook(io.BytesIO(response.data))
    assert wb.active['A1'].value == ('产品入库单' if managed else '采购入库单')
    assert path.read_bytes() == before
    wb.close()


def test_preview_prints_warehouse_and_month_scope_inside_paper(report_case, monkeypatch):
    c = report_case
    monkeypatch.setattr(m, '_collect_report_export_rows', lambda *a: (
        [{'field': 'code', 'title': '编码'}], [{'code': 'R-M'}], False, 1))
    response = c['client'].get('/report/warehouse_monthly/print_preview', query_string={
        'warehouse_id': c['b'], 'start_month': '2026-09', 'end_month': '2026-10'})
    assert response.status_code == 200
    soup = BeautifulSoup(response.data, 'html.parser')
    scope = soup.select_one('.sheet .print-scope')
    assert scope and '乙仓' in scope.text and 'R-B' in scope.text
    assert '2026-09' in scope.text and '2026-10' in scope.text


def test_builtin_preview_footer_spans_columns_and_columns_fit_paper():
    from services.report_print_preview import build_report_print_preview
    cols = [{'field': 'f%d' % i, 'title': '字段%d' % i} for i in range(9)]
    sheets = build_report_print_preview('月报', cols, [{'f%d' % i: i for i in range(9)}])
    footer = next(cell for row in sheets[0]['rows'] for cell in row['cells']
                  if cell['text'].startswith('打印日期：'))
    assert footer['colspan'] >= 3
    template = (Path(__file__).resolve().parents[1] / 'app/templates/report_print_preview.html').read_text('utf-8')
    assert 'width:{{ width }}px' not in template
    assert 'sheet.width_percentages' in template


@pytest.mark.parametrize('scenario', ['outside', 'inside', 'management', 'escape', 'preview',
                                     'download', 'missing', 'invalid', 'standard'])
def test_report_print_menu_actual_handlers(scenario):
    node = shutil.which('node')
    if not node:
        pytest.skip('Node.js required for menu handler runtime checks')
    root = Path(__file__).resolve().parents[1]
    result = subprocess.run([node, str(root / 'tests/frontend_report_print_menu.cjs'), str(root), scenario],
                            capture_output=True, text=True, encoding='utf-8', timeout=30)
    assert result.returncode == 0, result.stderr
