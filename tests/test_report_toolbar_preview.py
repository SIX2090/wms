"""Real-data report template preview is read-only, scoped and faithful to the chosen layout."""
from __future__ import annotations

from datetime import date

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import app as m

_REAL_COLLECT = m._collect_report_export_rows


@pytest.fixture
def preview_case(monkeypatch, tmp_path):
    monkeypatch.setitem(m.app.config, 'TESTING', True)
    monkeypatch.setitem(m.app.config, 'WTF_CSRF_ENABLED', False)
    monkeypatch.setattr(m.app, 'static_folder', str(tmp_path))
    with m.app.app_context():
        m.db.session.remove()
        m.db.drop_all()
        m.db.create_all()
        user = m.User(username='preview', role='viewer', must_change_password=False,
                      password_hash='isolated-test-only')
        warehouse = m.Warehouse(code='PREVIEW', name='预览仓', status='active')
        m.db.session.add_all([user, warehouse])
        m.db.session.commit()
        ids = {'user': user.id, 'warehouse': warehouse.id}
    rows = [{'code': 'ROW-%02d' % n, 'name': '物料%d' % n,
             'quantity': n, 'amount': n * 7.25} for n in range(25)]
    columns = [{'field': 'code', 'title': '编码'}, {'field': 'name', 'title': '名称'},
               {'field': 'quantity', 'title': '数量'},
               {'field': 'amount', 'title': '金额', 'type': 'money'}]
    monkeypatch.setattr(m, '_collect_report_export_rows',
                        lambda kind, filters: (columns, rows, False, len(rows)))
    monkeypatch.setattr(m, 'price_dimension_enabled', lambda: True)
    client = m.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(ids['user'])
        session['_fresh'] = True
    yield dict(client=client, ids=ids, rows=rows, root=tmp_path)
    with m.app.app_context():
        m.db.session.remove()


def _get(case, **params):
    return case['client'].get('/report/in_detail/print_preview', query_string={
        'warehouse_id': case['ids']['warehouse'], **params})


def _template(case, *, name='定制报表', target='report_in_detail', formula=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells('A1:B1')
    sheet['A1'] = name
    sheet['A1'].font = Font(bold=True, color='008000', size=18)
    sheet['A1'].fill = PatternFill('solid', fgColor='FFEEDD')
    sheet.column_dimensions['A'].width = 22
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = '9'
    sheet.page_margins.left = .4
    sheet['A2'], sheet['B2'] = '{item.code}', '{item.name}'
    sheet['A3'] = '=SUM(A2:A2)' if formula else '{print_date}'
    path = case['root'] / ('preview-%s.xlsx' % target)
    workbook.save(path)
    with m.app.app_context():
        record = m.ExcelPrintTemplate(name=name, target_type='report',
                    target_code=target, excel_template_path='/static/' + path.name)
        m.db.session.add(record)
        m.db.session.commit()
        return record.id


def test_route_requires_login_and_warehouse(preview_case):
    anon = m.app.test_client()
    assert anon.get('/report/in_detail/print_preview').status_code in (302, 401)
    response = preview_case['client'].get('/report/in_detail/print_preview')
    assert response.status_code == 400
    assert '请选择仓库' in response.get_data(as_text=True)
    assert _get(preview_case, warehouse_id=99999).status_code == 400


def test_current_filter_all_rows_escaped_and_no_writes(preview_case, monkeypatch):
    case = preview_case
    case['rows'][-1]['name'] = '<script>malicious()</script>'
    with m.app.app_context():
        before = m.ExcelPrintTemplate.query.count()
    before_files = list(case['root'].rglob('*'))
    response = _get(case, page=9, page_size=1)
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert 'ROW-00' in html and 'ROW-24' in html
    assert '&lt;script&gt;malicious()&lt;/script&gt;' in html
    assert '<script>malicious()' not in html
    assert 'window.print()' in html
    assert '当前筛选共 25 行' in html
    assert list(case['root'].rglob('*')) == before_files
    with m.app.app_context():
        assert m.ExcelPrintTemplate.query.count() == before


def test_actual_selected_template_styles_merges_paper_and_values(preview_case):
    template_id = _template(preview_case)
    html = _get(preview_case, template_id=template_id).get_data(as_text=True)
    assert '定制报表' in html and 'ROW-24' in html
    assert 'colspan="2"' in html
    assert '#008000' in html and '#FFEEDD' in html
    assert 'A4 landscape' in html
    assert 'width:297mm;min-height:210mm' in html
    assert '10.16mm' in html
    assert '{item.' not in html


@pytest.mark.parametrize('chosen', ['99999', 'bad', '0', '-2', '9' * 30])
def test_invalid_explicit_template_never_falls_back(preview_case, chosen):
    response = _get(preview_case, template_id=chosen)
    assert response.status_code == 400
    assert 'ROW-00' not in response.get_data(as_text=True)


def test_wrong_report_template_and_missing_file_rejected(preview_case):
    wrong = _template(preview_case, target='report_inventory')
    assert _get(preview_case, template_id=wrong).status_code == 400
    chosen = _template(preview_case)
    (preview_case['root'] / 'preview-report_in_detail.xlsx').unlink()
    assert _get(preview_case, template_id=chosen).status_code == 400


def test_formula_is_not_silently_printed_as_a_result(preview_case):
    chosen = _template(preview_case, formula=True)
    response = _get(preview_case, template_id=chosen)
    assert response.status_code == 400
    assert '公式' in response.get_data(as_text=True)
    assert 'SUM(' not in response.get_data(as_text=True)


def test_truncated_data_refused_instead_of_partial_print(preview_case, monkeypatch):
    monkeypatch.setattr(m, '_collect_report_export_rows',
                        lambda *a: ([{'field': 'code', 'title': '编码'}],
                                    [{'code': 'PARTIAL'}], True, 50001))
    response = _get(preview_case)
    assert response.status_code == 400
    assert '缩小筛选范围' in response.get_data(as_text=True)
    assert 'PARTIAL' not in response.get_data(as_text=True)


def test_price_disabled_cannot_be_reintroduced_by_custom_template(preview_case, monkeypatch):
    monkeypatch.setattr(m, 'price_dimension_enabled', lambda: False)
    case = preview_case
    case['rows'][0]['amount'] = 87654321.25
    sheet = Workbook()
    sheet.active.append(['{item.code}', '{item.amount}', '{order.total_amount}'])
    path = case['root'] / 'money.xlsx'
    sheet.save(path)
    with m.app.app_context():
        record = m.ExcelPrintTemplate(name='金额模板', target_type='report',
            target_code='report_in_detail', excel_template_path='/static/money.xlsx')
        m.db.session.add(record)
        m.db.session.commit()
        chosen = record.id
    response = _get(case, template_id=chosen)
    assert response.status_code == 200
    assert '87654321' not in response.get_data(as_text=True)
    blocked = case['client'].get('/report/purchase_price_analysis/print_preview',
              query_string={'warehouse_id': case['ids']['warehouse']})
    assert blocked.status_code == 400


def test_service_rejects_oversized_template(preview_case):
    from services.report_print_preview import build_report_print_preview
    sheet = Workbook()
    sheet.active.cell(2001, 1, 'too large')
    path = preview_case['root'] / 'huge.xlsx'
    sheet.save(path)
    with pytest.raises(ValueError, match='规模过大'):
        build_report_print_preview('报表', [{'field': 'code', 'title': '编码'}],
            [{'code': 'A'}], template_path=path)


def test_real_report_over_500_rows_sort_filter_and_warehouse(preview_case, monkeypatch):
    """Use the real paged builder, force multiple batches, retain all matching rows only."""
    case = preview_case
    monkeypatch.setattr(m, '_collect_report_export_rows', _REAL_COLLECT)
    monkeypatch.setattr(m, 'REPORT_EXPORT_BATCH_SIZE', 100)
    with m.app.app_context():
        unit = m.Unit(code='PREVIEW-UNIT', name='件')
        other = m.Warehouse(code='OTHER', name='另仓', status='active')
        m.db.session.add_all([unit, other])
        m.db.session.flush()
        material = m.Material(code='MATCH', name='命中物料', unit_id=unit.id)
        omitted = m.Material(code='SKIP', name='不匹配物料', unit_id=unit.id)
        m.db.session.add_all([material, omitted])
        m.db.session.flush()
        for warehouse, count in [('预览仓', 503), ('另仓', 2)]:
            for index in range(count):
                order = m.InOrder(order_no='%s-%04d' % ('MATCH' if warehouse == '预览仓' else 'FOREIGN', index),
                    warehouse=warehouse, date=date(2026, 10, 26), status='completed')
                m.db.session.add(order)
                m.db.session.flush()
                m.db.session.add(m.InOrderItem(in_order_id=order.id, material_id=material.id,
                                    quantity=1, price=2, amount=2))
        draft = m.InOrder(order_no='DRAFT-NOT-PRINTED', warehouse='预览仓',
            date=date(2026, 10, 25), status='pending')
        m.db.session.add(draft)
        m.db.session.flush()
        m.db.session.add(m.InOrderItem(in_order_id=draft.id, material_id=omitted.id,
                                     quantity=1, price=2, amount=2))
        m.db.session.commit()
    response = _get(case, page=9, page_size=1, material_code='MATCH', status='completed',
                    start_date='2026-10-26', end_date='2026-10-26', sort_field='order_no', sort_order='desc')
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert '当前筛选共 503 行' in html
    assert html.count('MATCH-') == 503
    assert html.index('MATCH-0502') < html.index('MATCH-0000')
    assert 'FOREIGN-' not in html and 'DRAFT-NOT-PRINTED' not in html


def test_print_area_and_footer_merges_follow_expanded_rows(preview_case):
    chosen = _template(preview_case)
    path = preview_case['root'] / 'preview-report_in_detail.xlsx'
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet['A3'] = '打印日期：{print_date}'
    sheet.merge_cells('A3:B3')
    sheet.print_area = 'A1:B3'
    workbook.save(path)
    response = _get(preview_case, template_id=chosen)
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert 'ROW-24' in html and '打印日期：' in html
    assert html.count('colspan="2"') == 2


def test_stale_default_uses_memory_builtin_but_explicit_bad_default_errors(preview_case):
    chosen = _template(preview_case)
    with m.app.app_context():
        record = m.db.session.get(m.ExcelPrintTemplate, chosen)
        record.is_default = True
        m.db.session.commit()
    (preview_case['root'] / 'preview-report_in_detail.xlsx').unlink()
    response = _get(preview_case)
    assert response.status_code == 200
    assert '系统内置模板' in response.get_data(as_text=True)
    assert _get(preview_case, template_id=chosen).status_code == 400


def test_oversized_print_area_is_rejected_without_allocating_it(preview_case):
    chosen = _template(preview_case)
    path = preview_case['root'] / 'preview-report_in_detail.xlsx'
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    workbook.active.print_area = 'A1:ZZ999999'
    workbook.save(path)
    response = _get(preview_case, template_id=chosen)
    assert response.status_code == 400
    assert '打印区域过大' in response.get_data(as_text=True)


def test_empty_query_never_prints_template_sample_rows(preview_case):
    chosen = _template(preview_case)
    path = preview_case['root'] / 'preview-report_in_detail.xlsx'
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet['A3'] = 'SAMPLE-DO-NOT-PRINT'
    sheet['A4'] = '{print_date}'
    workbook.save(path)
    preview_case['rows'].clear()
    response = _get(preview_case, template_id=chosen)
    assert response.status_code == 200
    assert 'SAMPLE-DO-NOT-PRINT' not in response.get_data(as_text=True)


def test_request_does_not_issue_any_database_mutation(preview_case):
    from sqlalchemy import event
    statements = []
    def collect(connection, cursor, statement, parameters, context, executemany):
        statements.append(statement.lstrip().split()[0].upper())
    with m.app.app_context():
        engine = m.db.engine
        event.listen(engine, 'before_cursor_execute', collect)
        try:
            response = _get(preview_case)
        finally:
            event.remove(engine, 'before_cursor_execute', collect)
    assert response.status_code == 200
    assert not {'INSERT', 'UPDATE', 'DELETE', 'CREATE', 'DROP', 'ALTER'} & set(statements)


def test_numeric_values_keep_precision_and_custom_number_format(preview_case):
    from services.report_print_preview import _display_value
    workbook = Workbook()
    cell = workbook.active['A1']
    cell.value = 12345678.125
    assert _display_value(cell) == '12345678.125'
    cell.number_format = '#,##0.00'
    assert _display_value(cell) == '12,345,678.12'
    cell.value = .123
    cell.number_format = '0.0%'
    assert _display_value(cell) == '12.3%'


def test_projected_expansion_is_checked_before_filling(preview_case, monkeypatch):
    from services import report_print_preview as service
    workbook = Workbook()
    workbook.active.append(['{item.code}'] * 100)
    path = preview_case['root'] / 'wide.xlsx'
    workbook.save(path)
    def no_fill(*args):
        pytest.fail('Projected-size rejection must happen before filling')
    monkeypatch.setattr(service, '_fill_sheet', no_fill)
    with pytest.raises(ValueError, match='展开后的打印内容过大'):
        service.build_report_print_preview('宽报表', [{'field': 'code', 'title': '编码'}],
            [{'code': 'A'}] * 2000, template_path=path)


def test_visible_report_data():
    from services.report_print_preview import visible_report_data
    columns = [{'field': 'quantity'}, {'field': 'amount', 'type': 'money'}, {'field': 'unit_price'}]
    rows = [{'quantity': 7, 'amount': 21, 'unit_price': 3}]
    visible_columns, visible_rows = visible_report_data(columns, rows, include_prices=False)
    assert visible_columns == [{'field': 'quantity'}]
    assert visible_rows == [{'quantity': 7}]
    assert rows[0] == {'quantity': 7, 'amount': 21, 'unit_price': 3}
    assert len(columns) == 3
    assert visible_report_data(columns, rows, include_prices=True) == (columns, rows)


def test_build_report_print_preview(tmp_path, monkeypatch):
    from services.report_print_preview import build_report_print_preview
    monkeypatch.chdir(tmp_path)
    sheets = build_report_print_preview('实际领料', [{'field': 'quantity', 'title': '数量'}],
                                       [{'quantity': 12345678.125}])
    values = [cell['text'] for sheet in sheets for row in sheet['rows'] for cell in row['cells']]
    assert '实际领料' in values
    assert '12345678.125' in values
    assert sheets[0]['page_width'] in (210, 297)
    assert not list(tmp_path.iterdir())
