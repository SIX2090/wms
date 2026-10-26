"""Sales report outputs share page collectors, current filters and selected templates."""
from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import event

import app as m


@pytest.fixture
def sales_case(monkeypatch, tmp_path):
    monkeypatch.setitem(m.app.config, 'TESTING', True)
    monkeypatch.setitem(m.app.config, 'WTF_CSRF_ENABLED', False)
    monkeypatch.setattr(m.app, 'static_folder', str(tmp_path))
    with m.app.app_context():
        m.db.session.remove()
        m.db.drop_all()
        m.db.create_all()
        wh = m.Warehouse(code='SA', name='销售甲仓', status='active')
        other = m.Warehouse(code='SB', name='销售乙仓', status='active')
        user = m.User(username='sales-print', role='viewer', must_change_password=False,
                      password_hash='isolated-only')
        customer = m.Customer(code='C-A', name='客户甲')
        employee = m.Employee(code='E-A', name='业务员甲')
        unit = m.Unit(code='PCS', name='件')
        m.db.session.add_all([wh, other, user, customer, employee, unit])
        m.db.session.flush()
        mat = m.Material(code='SP-A', name='测试销售物料', unit_id=unit.id)
        m.db.session.add(mat)
        m.db.session.flush()
        orders = []
        for index, warehouse in enumerate([wh, other]):
            order = m.SalesOrder(order_no='SO-%s' % ('A' if not index else 'FOREIGN'),
                date=date.today(), warehouse_id=warehouse.id, warehouse=warehouse.name,
                customer_id=customer.id, salesperson_id=employee.id,
                status='confirmed', shipment_status='partial', project_no='P-A',
                contract_no='CONTRACT-A', project_name='工程甲', total_amount=113,
                untaxed_amount=100, tax_amount=13, shipped_amount=22.6)
            m.db.session.add(order)
            m.db.session.flush()
            item = m.SalesOrderItem(sales_order_id=order.id, material_id=mat.id,
                quantity=10, price=11.3, amount=113, shipped_quantity=2,
                tax_included_amount=113, untaxed_amount=100, tax_amount=13, tax_rate=.13)
            m.db.session.add(item)
            m.db.session.flush()
            outbound = m.OutOrder(order_no='OUT-%d' % index, date=date.today(),
                warehouse=warehouse.name, customer=customer.name,
                business_type='销售出库', status='completed', source_sales_order_id=order.id)
            m.db.session.add(outbound)
            m.db.session.flush()
            m.db.session.add(m.OutOrderItem(out_order_id=outbound.id, material_id=mat.id,
                source_sales_order_item_id=item.id, quantity=2, price=11.3, amount=22.6))
            orders.append(order.id)
        m.set_system_setting('price_dimension_enabled', '1')
        m.db.session.commit()
        ids = dict(warehouse=wh.id, other=other.id, user=user.id, material=mat.id,
                   customer=customer.id, employee=employee.id, order=orders[0])
    client = m.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'], session['_fresh'] = str(ids['user']), True
    yield dict(client=client, ids=ids, root=tmp_path)
    with m.app.app_context():
        m.db.session.remove()


def _get(case, kind='summary', mode='print_preview', **params):
    return case['client'].get('/sales/%s/%s' % (kind, mode), query_string={
        'warehouse_id': case['ids']['warehouse'],
        **({'scope': 'customer'} if kind == 'summary' else {}), **params})


def test_sales_report_template_preview(sales_case):
    response = _get(sales_case)
    assert response.status_code == 200
    assert '客户甲' in response.get_data(as_text=True)
    assert m.app.test_client().get('/sales/summary/print_preview').status_code in (302, 401)
    assert sales_case['client'].get('/sales/summary/print_preview?scope=customer').status_code == 400
    assert _get(sales_case, scope='').status_code == 400
    assert _get(sales_case, scope='bad').status_code == 400
    assert _get(sales_case, kind='bad').status_code == 400


def test_sales_report_template_excel(sales_case):
    response = _get(sales_case, mode='print_excel')
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    text = str(list(workbook.active.values))
    assert '客户甲' in text
    assert '{item.' not in text


@pytest.mark.parametrize('scope', ['customer', 'material', 'salesperson', 'detail'])
def test_collect_sales_report_print(sales_case, scope):
    from services.sales_report_print import collect_sales_report_print
    with m.app.test_request_context('/sales/summary/print_preview', query_string={
        'warehouse_id': sales_case['ids']['warehouse'], 'scope': scope}):
        output = collect_sales_report_print('summary')
        assert output['target_code'] == 'report_sales_summary_' + scope
        assert len(output['rows']) == 1
        assert 'SO-FOREIGN' not in str(output['rows'])


@pytest.mark.parametrize('kind', ['reconciliation', 'outflow', 'trend', 'execution', 'price_analysis'])
def test_six_report_outputs_include_current_rows(sales_case, kind):
    response = _get(sales_case, kind=kind)
    assert response.status_code == 200, response.get_data(as_text=True)
    html = response.get_data(as_text=True)
    assert '当前筛选共 1 行' in html
    assert 'SO-FOREIGN' not in html and 'OUT-1' not in html
    excel = _get(sales_case, kind=kind, mode='print_excel')
    assert excel.status_code == 200
    assert excel.mimetype.endswith('spreadsheetml.sheet')


def test_resolve_sales_report_print_template(sales_case):
    from services.sales_report_print import resolve_sales_report_print_template
    with m.app.test_request_context('/?template_id=99999'):
        with pytest.raises(ValueError):
            resolve_sales_report_print_template('report_sales_summary_customer')
    workbook = Workbook()
    workbook.active['A1'] = '选中客户模板'
    workbook.active['A2'] = '{item.name}'
    workbook.active['A3'] = '{print_date}'
    path = sales_case['root'] / 'custom.xlsx'
    workbook.save(path)
    with m.app.app_context():
        record = m.ExcelPrintTemplate(name='定制客户模板', target_type='report',
            target_code='report_sales_summary_customer', excel_template_path='/static/custom.xlsx')
        m.db.session.add(record)
        m.db.session.commit()
        chosen = record.id
    response = _get(sales_case, template_id=chosen)
    assert response.status_code == 200
    assert '选中客户模板' in response.get_data(as_text=True)
    assert _get(sales_case, scope='material', template_id=chosen).status_code == 400
    path.unlink()
    assert _get(sales_case, template_id=chosen).status_code == 400


def test_render_sales_report_excel(sales_case):
    from services.sales_report_print import render_sales_report_excel
    output = render_sales_report_excel('测试报表', [{'field': 'name', 'title': '名称'}],
                                     [{'name': '真实名称'}])
    assert '真实名称' in str(list(load_workbook(output).active.values))


def test_print_filters_and_full_detail_not_page(sales_case):
    with m.app.app_context():
        m.db.session.add_all([m.SalesOrderItem(sales_order_id=sales_case['ids']['order'],
            material_id=sales_case['ids']['material'], quantity=1, price=1, amount=1)
            for _ in range(55)])
        m.db.session.commit()
    html = _get(sales_case, scope='detail', material_code='SP-A', page=2, per_page=1).get_data(as_text=True)
    assert '当前筛选共 56 行' in html
    filtered = _get(sales_case, scope='detail', customer_id=999999)
    assert '当前筛选共 0 行' in filtered.get_data(as_text=True)
    for params in [dict(salesperson_id=99999), dict(status='draft'),
                   dict(shipment_status='shipped'), dict(project_no='NOT-MATCH')]:
        assert '当前筛选共 0 行' in _get(sales_case, scope='detail', **params).get_data(as_text=True)


def test_existing_price_policy_and_read_only(sales_case):
    with m.app.app_context():
        m.set_system_setting('price_dimension_enabled', '0')
        m.db.session.commit()
        engine = m.db.engine
    assert _get(sales_case, kind='price_analysis').status_code == 400
    assert _get(sales_case, kind='price_analysis', mode='print_excel').status_code == 400
    statements = []
    def collect(connection, cursor, statement, parameters, context, executemany):
        statements.append(statement.lstrip().split()[0].upper())
    before_files = list(sales_case['root'].rglob('*'))
    event.listen(engine, 'before_cursor_execute', collect)
    try:
        assert _get(sales_case).status_code == 200
        assert _get(sales_case, mode='print_excel').status_code == 200
    finally:
        event.remove(engine, 'before_cursor_execute', collect)
    assert not {'INSERT', 'UPDATE', 'DELETE', 'CREATE', 'ALTER', 'DROP'} & set(statements)
    assert list(sales_case['root'].rglob('*')) == before_files


@pytest.mark.parametrize('code', ['report_sales_summary_customer', 'report_sales_summary_material',
    'report_sales_summary_salesperson', 'report_sales_summary_detail', 'report_sales_reconciliation',
    'report_sales_outflow', 'report_sales_trend', 'report_sales_execution', 'report_sales_price_analysis'])
def test_sales_scopes_registered_for_template_center(sales_case, code):
    response = sales_case['client'].get('/doc_print_templates/' + code + '.json')
    assert response.status_code == 200
    assert response.get_json()['target_type'] == 'report'


def test_sales_summary_context(sales_case):
    from services.sales_report_print import sales_summary_context
    with m.app.test_request_context(query_string={'warehouse_id': sales_case['ids']['warehouse']}):
        context = sales_summary_context()
        assert context['by_customer'][0]['amount'] == 113
        assert context['by_material'][0]['quantity'] == 10
        assert context['by_salesperson'][0]['name'] == '业务员甲'


def test_sales_reconciliation_context(sales_case):
    from services.sales_report_print import sales_reconciliation_context
    with m.app.test_request_context(query_string={'warehouse_id': sales_case['ids']['warehouse']}):
        row = sales_reconciliation_context()['rows'][0]
        assert row['ok'] and row['quantity_diff'] == 0
        assert row['shipped_qty'] == row['completed_qty'] == 2


def test_sales_outflow_context(sales_case):
    from services.sales_report_print import sales_outflow_context
    with m.app.test_request_context(query_string={'warehouse_id': sales_case['ids']['warehouse']}):
        context = sales_outflow_context()
        assert context['total_quantity'] == 2
        assert context['rows'][0]['tax_amount'] == 2.6


def test_sales_trend_context(sales_case):
    from services.sales_report_print import sales_trend_context
    with m.app.test_request_context(query_string={'warehouse_id': sales_case['ids']['warehouse']}):
        row = sales_trend_context()['rows'][0]
        assert row['amount'] == 113 and row['orders'] == 1


def test_sales_execution_context(sales_case):
    from services.sales_report_print import sales_execution_context
    with m.app.test_request_context(query_string={'warehouse_id': sales_case['ids']['warehouse']}):
        row = sales_execution_context()['rows'][0]
        assert row['remaining_quantity'] == 8
        assert row['execution_rate'] == 20


def test_sales_price_context(sales_case):
    from services.sales_report_print import sales_price_context
    with m.app.test_request_context(query_string={'warehouse_id': sales_case['ids']['warehouse']}):
        row = sales_price_context()['rows'][0]
        assert row['customers'] == 1 and row['avg_price'] == 11.3


def test_outflow_applied_filter_matches_shared_collector(sales_case):
    assert '当前筛选共 0 行' in _get(sales_case, kind='outflow', search='NO-MATCH').get_data(as_text=True)
    assert '当前筛选共 0 行' in _get(sales_case, kind='outflow', customer='别的客户').get_data(as_text=True)
    assert _get(sales_case, kind='outflow', date_start='wrong-date').status_code == 400


def test_stale_default_fallback_invalid_id_and_output_limits(sales_case):
    with m.app.app_context():
        record = m.ExcelPrintTemplate(name='失效默认', target_type='report',
            target_code='report_sales_summary_customer', is_default=True,
            excel_template_path='/static/missing.xlsx')
        m.db.session.add(record)
        m.db.session.commit()
        chosen = record.id
    assert _get(sales_case).status_code == 200
    assert _get(sales_case, template_id=chosen).status_code == 400
    for raw in ['bad', '-1', '0', '9' * 30]:
        assert _get(sales_case, template_id=raw).status_code == 400
    from services.sales_report_print import render_sales_report_excel
    with pytest.raises(ValueError, match='5000'):
        render_sales_report_excel('超限', [{'field': 'code', 'title': '编码'}], [{'code': 'A'}] * 5001)


@pytest.mark.parametrize('extra', [dict(customer_id=99999), dict(salesperson_id=99999),
    dict(status='draft'), dict(shipment_status='shipped'), dict(project_no='NO-MATCH'),
    dict(material_code='NO-MATCH')])
def test_export_sales_report_applied_filters(sales_case, extra):
    response = sales_case['client'].get('/sales/report/export', query_string={
        'warehouse_id': sales_case['ids']['warehouse'], **extra})
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    rows = list(workbook.active.values)
    assert len(rows) == 1  # Header only: the existing export must not leak unmatched orders.
    assert workbook.sheetnames == ['销售订单执行']


@pytest.mark.parametrize('path', ['/sales/report', '/sales/reconciliation', '/sales/outflow_report',
    '/sales/trend_report', '/sales/execution_report', '/sales/price_analysis'])
def test_shared_collectors_preserve_existing_pages(sales_case, path):
    response = sales_case['client'].get(path, query_string={'warehouse_id': sales_case['ids']['warehouse']})
    assert response.status_code == 200
    assert 'SO-FOREIGN' not in response.get_data(as_text=True)
