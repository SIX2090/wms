"""BUG-2026-10-25-012: stock query and both Excel branches share effective filters."""
from pathlib import Path
from datetime import date
from html import unescape
from urllib.parse import parse_qs, urlparse
import io
import json
import re
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
import app as w
import doc_print_excel as dpe


def read_list(response):
    assert response.status_code==200, response.status_code
    html=response.get_data(as_text=True)
    body=re.search(r'<tbody>(.*?)</tbody>',html,re.S).group(1)
    rows=[]
    for row in re.findall(r'<tr\b.*?</tr>',body,re.S):
        cells=[unescape(re.sub('<[^>]+>','',cell)).strip()
               for cell in re.findall(r'<td\b[^>]*>(.*?)</td>',row,re.S)]
        if len(cells)>7 and cells[1].startswith(('AUD-','BULK-')):
            rows.append((cells[1],float(cells[7])))
    return html,rows


def read_excel(response,templated):
    assert response.status_code==200,response.status_code
    assert response.data[:2]==b'PK'
    rows=[]
    for row in load_workbook(io.BytesIO(response.data)).active.iter_rows(values_only=True):
        offset=1 if templated else 0
        if len(row)>offset and isinstance(row[offset],str) and row[offset].startswith(('AUD-','BULK-')):
            rows.append((row[offset],float(row[6 if templated else 5])))
    return rows


@pytest.fixture(scope='module')
def case(tmp_path_factory):
    from werkzeug.security import generate_password_hash
    static=tmp_path_factory.mktemp('stock-query-print-static')
    changes=pytest.MonkeyPatch()
    changes.setattr(w.app,'static_folder',str(static))
    changes.setitem(w.app.config,'TESTING',True)
    changes.setitem(w.app.config,'WTF_CSRF_ENABLED',False)
    history=None
    try:
        with w.app.app_context():
            assert str(w.db.engine.url)=='sqlite:///:memory:'
            w.db.drop_all();w.db.create_all()
            w.set_system_setting('location_management_enabled','0')
            w.set_system_setting('inventory_alert_enabled','1')
            user=w.User(username='stock-filter-probe',role='admin',must_change_password=False,
                        password_hash=generate_password_hash('fixture-only-stock-query'))
            whs=[w.Warehouse(code='AUD-W'+x,name='审计'+x+'仓',status='active',is_default=x=='A') for x in 'ABCD']
            cats=[w.MaterialCategory(code='AUD-C1',name='紧固&件'),w.MaterialCategory(code='AUD-C2',name='常规分类'),
                  w.MaterialCategory(code='AUD-C3',name='跨页分类')]
            units=[w.Unit(code='BOX',name='箱'),w.Unit(code='PCS',name='件')]
            suppliers=[w.Supplier(code='AUD-S1',name='独有供应商'),w.Supplier(code='AUD-S2',name='其他供应商')]
            w.db.session.add_all([user,*whs,*cats,*units,*suppliers]);w.db.session.flush()
            mats={}
            for code,cat,unit,supplier,stock,minimum,reorder,price in [
                ('C',cats[0],units[1],suppliers[1],11,0,0,6),
                ('A',cats[0],units[0],suppliers[0],63,3,7,2),
                ('B',cats[1],units[1],suppliers[1],11,3,5,4)]:
                m=w.Material(code='AUD-M'+code,name='审计物料'+code,spec='SPEC-'+code,
                             purpose='用途&Alpha' if code=='A' else '常规用途',category_id=cat.id,
                             unit_id=unit.id,supplier_id=supplier.id,stock=stock,min_stock=minimum,
                             reorder_point=reorder,price=price)
                w.db.session.add(m);w.db.session.flush();mats[code]=m
            a,b=whs[:2]
            source=w.InOrder(order_no='AUD-HISTORY-A',date=date(2026,10,25),status='completed',warehouse=a.name)
            w.db.session.add(source);w.db.session.flush()
            for code,wh,qty in [('A',a,2),('A',b,8),('B',a,10),('B',b,1),('C',b,6)]:
                w.db.session.add(w.StockTransaction(material_id=mats[code].id,transaction_type='in',quantity=qty,warehouse_id=wh.id,location=wh.name))
            w.db.session.add_all([
                w.StockTransaction(material_id=mats['A'].id,transaction_type='in',quantity=3,warehouse_id=None,location=None,reference_type='in_order',reference_id=source.id),
                w.StockTransaction(material_id=mats['C'].id,transaction_type='in',quantity=5,warehouse_id=None,location=a.code),
                w.StockTransaction(material_id=mats['A'].id,transaction_type='in',quantity=50,warehouse_id=None,location=None)])
            for n in range(1,59):
                m=w.Material(code=f'BULK-{n:03}',name=f'跨页物料{n:03}',spec=f'SPEC-{n:03}',
                             category_id=cats[2].id,unit_id=units[1].id,supplier_id=suppliers[1].id,
                             stock=2*n+103,min_stock=0,reorder_point=0,price=10+n)
                w.db.session.add(m);w.db.session.flush()
                w.db.session.add_all([
                    w.StockTransaction(material_id=m.id,transaction_type='in',quantity=n+3,warehouse_id=a.id,location=a.name),
                    w.StockTransaction(material_id=m.id,transaction_type='in',quantity=n+100,warehouse_id=b.id,location=b.name)])
            w.db.session.commit()
            ids={'a':a.id,'b':b.id,'category':cats[0].id,'bulk':cats[2].id}
            history=[(t.id,t.warehouse_id,t.location,t.quantity) for t in w.StockTransaction.query.filter(w.StockTransaction.warehouse_id.is_(None)).all()]
        client=w.app.test_client()
        assert client.post('/login',data={'username':'stock-filter-probe','password':'fixture-only-stock-query'}).status_code==302
        yield client,ids
    finally:
        try:
            with w.app.app_context():
                try:
                    if history is not None:
                        assert history==[(t.id,t.warehouse_id,t.location,t.quantity) for t in w.StockTransaction.query.filter(w.StockTransaction.warehouse_id.is_(None)).all()]
                finally:
                    w.db.session.remove()
                    w.db.drop_all()
                    w.db.create_all()
                    w.db.session.remove()
        finally:
            changes.undo()


@pytest.fixture(params=['legacy','explicit_template','default_template'])
def printing(case,request):
    client,ids=case
    with w.app.app_context():
        w.ExcelPrintTemplate.query.delete();w.db.session.commit()
        tid=None
        if request.param!='legacy':
            path=dpe._builtin_template_abspath(w.app.static_folder,'stock_query')
            tpl=w.ExcelPrintTemplate(name='隔离库存模板',target_type='list',target_code='stock_query',
                template_type='excel',excel_template_path='/static/uploads/print_templates/'+Path(path).name,
                is_default=request.param=='default_template')
            w.db.session.add(tpl);w.db.session.commit();tid=tpl.id
    def print_rows(query):
        params=dict(query)
        if request.param=='explicit_template':params['template_id']=tid
        return read_excel(client.get('/report/stock/print',query_string=params),request.param!='legacy')
    try:
        yield client,ids,print_rows,request.param
    finally:
        with w.app.app_context():
            w.ExcelPrintTemplate.query.delete()
            w.db.session.commit()
            w.db.session.remove()


FILTERS=[('exact',{'search':'AUD-MA'}),('purpose',{'search':'用途&Alpha'}),
         ('supplier',{'search':'独有供应商'}),('unit',{'search':'箱'}),('spec',{'search':'SPEC-B'}),
         ('name',{'search':'审计物料C'}),('category',{'category_id':'CATEGORY'}),
         ('low',{'stock_filter':'low'}),('normal',{'stock_filter':'normal','search':'AUD-'}),
         ('combined',{'search':'AUD-MA','category_id':'CATEGORY','stock_filter':'low'}),
         ('empty',{'search':'NO-SUCH-MATERIAL'}),
         ('code_desc',{'sort':'code','order':'desc'}),('name_desc',{'sort':'name','order':'desc'}),
         ('spec_desc',{'sort':'spec','order':'desc'}),('price_desc',{'sort':'price','order':'desc'}),
         ('minimum_desc',{'sort':'min_stock','order':'desc'})]


def all_list_rows(client,query):
    rows=[]
    for page in range(1,10):
        _,part=read_list(client.get('/stock_query',query_string={**query,'page':page,'per_page':20}))
        rows.extend(part)
        if len(part)<20:return rows
    raise AssertionError('Fixture unexpectedly exceeds known 61 materials')


@pytest.mark.parametrize('name,filters',FILTERS,ids=[x[0] for x in FILTERS])
def test_build_stock_material_query(printing,name,filters):
    client,ids,print_rows,branch=printing
    query={'warehouse_id':ids['a'],**{k:ids['category'] if v=='CATEGORY' else v for k,v in filters.items()}}
    expected=all_list_rows(client,query)
    if name=='empty':assert expected==[]
    if name in ('exact','purpose','supplier','unit','combined'):assert expected==[('AUD-MA',5.0)]
    if name=='low':assert expected==[('AUD-MA',5.0)]
    assert print_rows(query)==expected, f'{name}/{branch}: printed rows must equal filtered list full set/order'


def test_print_58_matching_rows_across_three_pages(printing):
    client,ids,print_rows,_=printing
    query={'warehouse_id':ids['a'],'category_id':ids['bulk'],'sort':'code','order':'desc'}
    expected=all_list_rows(client,query)
    assert len(expected)==58
    assert expected[0]==('BULK-058',61.0) and expected[-1]==('BULK-001',4.0)
    actual=print_rows({**query,'page':2,'per_page':20})
    assert actual==expected
    assert sum(quantity for _,quantity in actual)==sum(quantity for _,quantity in expected)==1885


def test_other_warehouse_filter_keeps_its_own_values(printing):
    client,ids,print_rows,_=printing
    query={'warehouse_id':ids['b'],'stock_filter':'low'}
    expected=all_list_rows(client,query)
    assert expected==[('AUD-MB',1.0)]
    assert print_rows(query)==expected


def test_unfiltered_print_keeps_resolved_historical_quantities(printing):
    _,ids,print_rows,_=printing
    a=dict(print_rows({'warehouse_id':ids['a']}));b=dict(print_rows({'warehouse_id':ids['b']}))
    assert [a['AUD-MA'],a['AUD-MB'],a['AUD-MC']]==[5,10,5]
    assert [b['AUD-MA'],b['AUD-MB'],b['AUD-MC']]==[8,1,6]
    assert a['AUD-MA']+b['AUD-MA']==13  # material.stock is63, unresolved50 stays outside warehouses.


def test_get_stock_query_filters(case):
    client,ids=case
    query={'warehouse_id':ids['a'],'search':'用途&Alpha','category_id':ids['category'],
           'stock_filter':'low','sort':'code','order':'desc','page':1,'per_page':20,'embedded':1}
    html,rows=read_list(client.get('/stock_query',query_string=query))
    assert rows==[('AUD-MA',5.0)]
    href=unescape(re.search(r'id="stockPrintBtn"\s+href="([^"]+)"',html).group(1))
    args=parse_qs(urlparse(href).query)
    for key in ('warehouse_id','search','category_id','stock_filter','sort','order'):
        assert args.get(key)==[str(query[key])],f'effective {key} missing from real button href'
    assert not {'page','per_page','csrf_token','embedded'}.intersection(args)


def test_actual_template_change_preserves_effective_href_filters(case):
    node=shutil.which('node')
    if not node:pytest.skip('Node required for actual template change runtime')
    client,ids=case
    query={'warehouse_id':ids['a'],'search':'用途&Alpha','category_id':ids['category'],
           'stock_filter':'low','sort':'code','order':'desc'}
    html,_=read_list(client.get('/stock_query',query_string=query))
    href=unescape(re.search(r'id="stockPrintBtn"\s+href="([^"]+)"',html).group(1))
    source=(ROOT/'app/templates/stock_query.html').read_text(encoding='utf-8')
    code=re.search(r'<script>\s*(\(function \(\) \{.*?\}\)\(\);)\s*</script>',source,re.S).group(1)
    script=r'''
const vm=require('vm'),assert=require('node:assert/strict'),params=JSON.parse(process.argv[1]);
let change;const select={value:'',appendChild(){},addEventListener(event,fn){assert.equal(event,'change');change=fn}};
const btn={href:params.href,getAttribute(){return this.href},setAttribute(key,val){assert.equal(key,'href');this.href=val}};
const document={getElementById(id){return id==='stockTemplateSelect'?select:btn},createElement(){return {}}};
vm.runInNewContext(params.code,{document,window:{location:{origin:'http://isolated.test'}},URL,fetch:async()=>({json:async()=>({status:'success',templates:[]})})});
for(const value of ['78','']){
 select.value=value;change.call(select);const url=new URL(btn.href,'http://isolated.test');
 for(const [key,expected]of Object.entries(params.query))assert.equal(url.searchParams.get(key),String(expected),'effective '+key+' lost');
 assert.equal(url.searchParams.get('template_id'),value||null);
}
'''
    run=subprocess.run([node,'--max-old-space-size=64','-e',script,json.dumps({'code':code,'href':href,'query':query},ensure_ascii=False)],
                       capture_output=True,text=True,encoding='utf-8',timeout=20)
    assert run.returncode==0,run.stderr


def test_all_warehouse_print_stays_disabled(case):
    client,_=case
    html,_=read_list(client.get('/stock_query?warehouse_id=0'))
    assert re.search(r'id="stockPrintBtn".*?aria-disabled="true"',html,re.S)


def test_invalid_warehouse_remains_rejected(case):
    client,_=case
    assert client.get('/stock_query?warehouse_id=999999').status_code==400
    assert client.get('/report/stock/print?warehouse_id=999999').status_code==400
