"""Inventory preview uses full current filters and the existing stock template namespace."""
from datetime import date
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import event
from werkzeug.datastructures import MultiDict

import app as m


@pytest.fixture
def stock_preview(monkeypatch, tmp_path):
    monkeypatch.setitem(m.app.config, 'TESTING', True)
    monkeypatch.setitem(m.app.config, 'WTF_CSRF_ENABLED', False)
    monkeypatch.setattr(m.app, 'static_folder', str(tmp_path))
    with m.app.app_context():
        assert str(m.db.engine.url) == 'sqlite:///:memory:'
        m.db.session.remove(); m.db.drop_all(); m.db.create_all()
        m.set_system_setting('location_management_enabled', '0')
        m.set_system_setting('inventory_alert_enabled', '1')
        m.set_system_setting('price_dimension_enabled', '1')
        user = m.User(username='stock-preview', role='viewer', must_change_password=False,
                      password_hash='isolated-fixture-only')
        a = m.Warehouse(code='SP-A', name='预览甲仓', status='active')
        b = m.Warehouse(code='SP-B', name='预览乙仓', status='active')
        m.db.session.add_all([user, a, b]); m.db.session.flush()
        history = m.InOrder(order_no='SP-HISTORY', date=date(2026, 10, 26),
                            status='completed', warehouse=a.name)
        m.db.session.add(history); m.db.session.flush()
        for n in range(1, 59):
            material = m.Material(code='SP-%03d' % n, name='实物%d' % n,
                stock=9999, price=123.45, min_stock=2, reorder_point=3)
            m.db.session.add(material); m.db.session.flush()
            m.db.session.add_all([
                m.StockTransaction(material_id=material.id, quantity=n,
                    transaction_type='in', warehouse_id=a.id, location=a.name),
                m.StockTransaction(material_id=material.id, quantity=100+n,
                    transaction_type='in', warehouse_id=b.id, location=b.name)])
            if n == 1:
                m.db.session.add_all([
                    m.StockTransaction(material_id=material.id, quantity=2,
                        transaction_type='in', warehouse_id=None, location=None,
                        reference_type='in_order', reference_id=history.id),
                    m.StockTransaction(material_id=material.id, quantity=50,
                        transaction_type='in', warehouse_id=None, location=None)])
        workbook = Workbook(); sheet = workbook.active; sheet.title = '库存查询'
        sheet['A1'] = '甲版真实库存'; sheet['A2'] = '{item.code}'; sheet['B2'] = '{item.stock}'
        sheet['C2'] = '{item.price}'; sheet['D2'] = '{item.amount}'
        sheet['A3'] = '合计'; sheet['B3'] = '{total_quantity}'
        sheet.page_setup.orientation = 'portrait'; workbook.save(tmp_path / 'stock.xlsx')
        template = m.ExcelPrintTemplate(name='甲版', target_type='list', target_code='stock_query',
            template_type='excel', excel_template_path='/static/stock.xlsx', is_default=True)
        foreign = m.ExcelPrintTemplate(name='别的表', target_type='report', target_code='report_stock',
            template_type='excel', excel_template_path='/static/stock.xlsx')
        m.db.session.add_all([template, foreign]); m.db.session.commit()
        ids = dict(user=user.id, a=a.id, b=b.id, template=template.id, foreign=foreign.id)
    client = m.app.test_client()
    with client.session_transaction() as session:
        session['_user_id'] = str(ids['user']); session['_fresh'] = True
    yield dict(client=client, ids=ids, root=tmp_path)
    with m.app.app_context():
        m.db.session.remove(); m.db.drop_all(); m.db.create_all(); m.db.session.remove()


def _get(case, download=False, **params):
    return case['client'].get('/report/stock/' + ('print' if download else 'print_preview'),
        query_string={'warehouse_id': case['ids']['a'], **params})


def test_report_stock_print_preview(stock_preview):
    case = stock_preview
    response = _get(case, sort='code', order='desc', page=2, per_page=20)
    html = response.get_data(as_text=True)
    assert response.status_code == 200
    assert '当前筛选共 58 行' in html and '甲版真实库存' in html
    assert html.index('SP-058') < html.index('SP-001')
    assert '1713' in html  # 1..58 plus historical source-resolved 2, not material.stock.
    assert '9999' not in html and 'window.print()' in html
    assert response.cache_control.no_store


def test_collect_stock_print_data(stock_preview):
    from services.stock_print_preview import collect_stock_print_data
    case = stock_preview
    with m.app.app_context():
        a, materials, quantities = collect_stock_print_data(MultiDict({
            'warehouse_id':str(case['ids']['a']), 'search':'SP-001'}))
        b, other, other_quantities = collect_stock_print_data(MultiDict({
            'warehouse_id':str(case['ids']['b']), 'search':'SP-001'}))
        assert [x.code for x in materials] == [x.code for x in other] == ['SP-001']
        assert quantities[materials[0].id] == 3
        assert other_quantities[other[0].id] == 101
        assert quantities[materials[0].id] + other_quantities[other[0].id] == 104


@pytest.mark.parametrize('params,codes', [
    ({'search':'SP-001'}, ['SP-001']),
    ({'stock_filter':'low'}, ['SP-001', 'SP-002', 'SP-003']),
    ({'search':'missing'}, [])])
def test_preview_filters_equal_download(stock_preview, params, codes):
    case = stock_preview
    html = _get(case, **params).get_data(as_text=True)
    response = _get(case, download=True, **params)
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data)).active
    rows = [row for row in sheet.iter_rows(values_only=True)
            if isinstance(row[0], str) and row[0].startswith('SP-')]
    assert [row[0] for row in rows] == codes
    for n in range(1, 59):
        assert (('SP-%03d' % n) in html) == (('SP-%03d' % n) in codes)
    assert '当前筛选共 %d 行' % len(rows) in html


@pytest.mark.parametrize('bad', ['0', '-1', 'abc', '999999', 'foreign', 'missing', 'outside'])
def test_resolve_stock_print_template(stock_preview, bad):
    case = stock_preview
    if bad == 'foreign':
        bad = str(case['ids']['foreign'])
    elif bad in ('missing', 'outside'):
        with m.app.app_context():
            template = m.db.session.get(m.ExcelPrintTemplate, case['ids']['template'])
            template.excel_template_path = '/static/no-such.xlsx' if bad == 'missing' else '../outside.xlsx'
            m.db.session.commit()
        bad = str(case['ids']['template'])
    for download in (False, True):
        response = _get(case, download=download, template_id=bad)
        assert response.status_code == 400
        assert response.data[:2] != b'PK'


@pytest.mark.parametrize('template_mode', [None, '1'])
def test_stock_print_rows(stock_preview, template_mode):
    case = stock_preview
    for include_prices in (False, True):
        with m.app.app_context():
            m.set_system_setting('price_dimension_enabled', '1' if include_prices else '0')
            m.db.session.commit()
        html = _get(case, search='SP-001').get_data(as_text=True)
        response = _get(case, download=True, search='SP-001',
            **({'template_mode': template_mode} if template_mode else {}))
        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.data)).active
        assert sheet['C2'].value == (123.45 if include_prices else None)
        assert sheet['D2'].value == (370.35 if include_prices else None)
        assert ('123.45' in html) is include_prices
        assert ('370.35' in html) is include_prices
        assert sheet['B2'].value == 3


def test_readonly_preview_without_template_never_creates_files_or_mutates_stock(stock_preview):
    case = stock_preview
    with m.app.app_context():
        m.ExcelPrintTemplate.query.delete(); m.db.session.commit()
        engine = m.db.engine
    before = {path.relative_to(case['root']): path.read_bytes() for path in case['root'].rglob('*') if path.is_file()}
    statements = []
    def capture(_conn, _cursor, statement, _parameters, _context, _executemany):
        if statement.lstrip().split(' ', 1)[0].upper() in ('INSERT','UPDATE','DELETE','CREATE','ALTER'):
            statements.append(statement)
    event.listen(engine, 'before_cursor_execute', capture)
    try:
        response = _get(case, search='SP-001')
    finally:
        event.remove(engine, 'before_cursor_execute', capture)
    assert response.status_code == 200 and not statements
    assert before == {path.relative_to(case['root']): path.read_bytes() for path in case['root'].rglob('*') if path.is_file()}
    assert '系统内置模板' in response.get_data(as_text=True)


def test_selected_stock_template_changes_actual_layout(stock_preview):
    case = stock_preview
    workbook = load_workbook(case['root'] / 'stock.xlsx')
    workbook.active['A1'] = '乙版横向'
    workbook.active.title = '自建工作表'
    workbook.active.page_setup.orientation = 'landscape'
    workbook.save(case['root'] / 'second.xlsx')
    with m.app.app_context():
        template = m.ExcelPrintTemplate(name='乙版', target_type='list', target_code='stock_query',
            template_type='excel', excel_template_path='/static/second.xlsx')
        m.db.session.add(template); m.db.session.commit(); template_id = template.id
    html = _get(case, template_id=template_id, search='SP-001', warehouse_id=case['ids']['b']).get_data(as_text=True)
    assert '乙版横向' in html and '甲版真实库存' not in html
    assert 'landscape' in html and '101' in html
    sheet = load_workbook(io.BytesIO(_get(case, download=True, template_mode=1, template_id=template_id,
        search='SP-001', warehouse_id=case['ids']['b']).data)).active
    assert sheet['A1'].value == '乙版横向' and sheet['B2'].value == 101


def test_stock_print_columns():
    from services.stock_print_preview import stock_print_columns
    assert [column['field'] for column in stock_print_columns()] == [
        'warehouse', 'code', 'brand', 'name', 'spec', 'unit', 'stock', 'price', 'amount']


@pytest.mark.parametrize('prices', [True, False])
def test_build_stock_excel_print(stock_preview, prices):
    case = stock_preview
    with m.app.app_context():
        m.ExcelPrintTemplate.query.delete()
        m.set_system_setting('price_dimension_enabled', '1' if prices else '0')
        m.db.session.commit()
    before = list(case['root'].rglob('*'))
    html = _get(case, search='SP-001').get_data(as_text=True)
    response = _get(case, download=True, template_mode=1, search='SP-001')
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data)).active
    values = [value for row in sheet.iter_rows(values_only=True) for value in row if value is not None]
    assert '库存查询（预览甲仓）' in values
    assert 'SP-001' in values and 3 in values
    assert ('单价' in values) is prices and ('单价' in html) is prices
    assert (123.45 in values) is prices and ('123.45' in html) is prices
    legacy = _get(case, download=True, search='SP-001')
    assert 'stock_report.xlsx' in legacy.headers['Content-Disposition']
    assert '分类' in [cell.value for cell in load_workbook(io.BytesIO(legacy.data)).active[1]]
    assert list(case['root'].rglob('*')) == before


@pytest.mark.parametrize('warehouse', [0, 999999, 'bad'])
def test_specific_active_warehouse_required(stock_preview, warehouse):
    for download in (False, True):
        assert _get(stock_preview, download=download, warehouse_id=warehouse).status_code == 400
    anon = m.app.test_client()
    assert anon.get('/report/stock/print_preview').status_code in (302, 401)


def test_missing_and_inactive_warehouse_rejected(stock_preview):
    case = stock_preview
    for path in ('/report/stock/print', '/report/stock/print_preview'):
        assert case['client'].get(path).status_code == 400
    with m.app.app_context():
        m.db.session.get(m.Warehouse, case['ids']['a']).status = 'inactive'
        m.db.session.commit()
    for download in (False, True):
        assert _get(case, download=download).status_code == 400


@pytest.mark.parametrize('problem', ['broken', 'dimension', 'expanded', 'filesize', 'sheets'])
def test_template_download_rejects_unsafe_workbooks_without_writes(stock_preview, problem):
    case = stock_preview
    path = case['root'] / 'stock.xlsx'
    if problem == 'broken':
        path.write_bytes(b'invalid workbook')
    elif problem == 'filesize':
        path.write_bytes(b'0' * (10 * 1024 * 1024 + 1))
    else:
        workbook = Workbook(); sheet = workbook.active
        sheet['A1'] = '{item.code}'
        if problem == 'dimension':
            sheet.cell(2000, 100, 'oversized matrix')
        elif problem == 'expanded':
            sheet.cell(1450, 100, 'before expansion fits; 58 rows exceed 150000 cells')
        else:
            for n in range(10):
                workbook.create_sheet('more-%s' % n)
        workbook.save(path); workbook.close()
    before_files = {p.relative_to(case['root']): p.read_bytes() for p in case['root'].rglob('*') if p.is_file()}
    with m.app.app_context():
        engine = m.db.engine
    writes = []
    def capture(_conn, _cursor, statement, _parameters, _context, _executemany):
        if statement.lstrip().split(' ', 1)[0].upper() in ('INSERT', 'UPDATE', 'DELETE', 'CREATE', 'ALTER'):
            writes.append(statement)
    event.listen(engine, 'before_cursor_execute', capture)
    try:
        response = _get(case, download=True, template_mode='1')
    finally:
        event.remove(engine, 'before_cursor_execute', capture)
    assert response.status_code == 400
    assert response.is_json
    assert any(text in str(response.get_json()) for text in ('模板', '打印内容过大', 'Excel 文件'))
    assert 'Traceback' not in response.get_data(as_text=True)
    assert not writes
    assert before_files == {p.relative_to(case['root']): p.read_bytes() for p in case['root'].rglob('*') if p.is_file()}


def test_stock_template_download_row_cap_before_loading_workbook(stock_preview, monkeypatch):
    from services.stock_print_preview import build_stock_excel_print
    import openpyxl
    def unexpected_load(*_args, **_kwargs):
        raise AssertionError('Oversized rows must be rejected before loading a workbook')
    monkeypatch.setattr(openpyxl, 'load_workbook', unexpected_load)
    with pytest.raises(ValueError, match='5000'):
        build_stock_excel_print('库存', [{}] * 5001)
