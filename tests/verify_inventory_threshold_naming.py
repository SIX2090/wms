# -*- coding: utf-8 -*-
"""库存阈值命名统一（AI-CI-GREEN-005）回归测试。

背景：`inventory_alert_enabled` 打开后，「安全库存」相关概念在系统里长期存在
**四种叫法**，容易在改动时写错字段、或在模板/接口间错位：

    数据库列            对外名称        出现位置
    min_stock           最低库存        界面、Excel、通知、报表
    max_stock           最大库存        （唯一，无歧义）
    reorder_point       **安全库存**    界面表单、Excel 表头、预警页
    safety_stock        （不是列！）    仅 API 响应 / Excel 导出 / 报表字典，
                                        值为 max(reorder_point, min_stock)

本测试锁定「对外统一叫法」，防止旧叫法回流：

T1. 一律不存在「最小库存」这个界面叫法（历史残留，已统一为「最低库存」）。
T2. 一律不存在「再订货点 / 再订购点」这个界面叫法（已统一为「安全库存」）；
    但在 Excel 表头**匹配**逻辑里必须保留为兼容别名，否则老模板导不进来。
T3. 两份物料导入模板（/material/download_template 与 /export/template/material）
    表头必须完全一致——它们是重复实现，任一漂移都会造成"两个入口模板不同"。
T4. 安全库存字段必须能正确往返：填「安全库存」→ 落 reorder_point；
    导出再导入后值不变（防止把 reorder_point 写进 min_stock 这类错位）。
T5. 界面文案与数据库列名的映射，必须在 Material 模型 docstring 里成文，
    避免下一个人重新猜。
"""
from __future__ import annotations

import io
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
APP_DIR = ROOT / "app"
sys.path.insert(0, str(APP_DIR))
os.chdir(APP_DIR)

os.environ.setdefault("WMS_BOOTSTRAP_PASSWORD", "admin")
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["WMS_DATABASE_URI"] = "sqlite:///:memory:"
os.environ.setdefault("WMS_DEBUG", "0")
os.environ.setdefault("WMS_SKIP_AUTO_UPDATE", "1")

import app as app_module  # noqa: E402
from app import db, Material, User  # noqa: E402

app_module.app.config["TESTING"] = True
app_module.app.config["WTF_CSRF_ENABLED"] = False


# ---------- T1/T2：源码级叫法扫描 ----------

# 只扫「会展示给用户」的位置：界面文案、通知正文、模板示例。
# 注释/兼容别名列表不算——那是有意保留的历史兼容。
TEMPLATE_DIR = APP_DIR / "templates"


def _user_facing_files():
    files = list(TEMPLATE_DIR.glob("*.html"))
    files.append(APP_DIR / "notifications.py")
    return files


def test_t1_no_legacy_min_stock_label():
    """T1：界面文案不得再出现「最小库存」（标准叫法是「最低库存」）。"""
    offenders = []
    for f in _user_facing_files():
        text = f.read_text(encoding="utf-8")
        if "最小库存" in text:
            offenders.append(f.name)
    assert not offenders, (
        f"以下文件仍在用旧叫法「最小库存」，应统一为「最低库存」：{offenders}"
    )


def test_t2_no_legacy_reorder_point_label_but_keep_alias():
    """T2：界面不再叫「再订货点」，但 Excel 表头匹配必须保留兼容别名。"""
    offenders = []
    for f in _user_facing_files():
        text = f.read_text(encoding="utf-8")
        if "再订货点" in text or "再订购点" in text:
            offenders.append(f.name)
    assert not offenders, (
        f"以下界面仍在用旧叫法「再订货点/再订购点」，应统一为「安全库存」：{offenders}"
    )
    # 兼容别名必须留在导入匹配逻辑里，否则老模板无法导入
    importer = (APP_DIR / "routes" / "material.py").read_text(encoding="utf-8")
    assert "再订货点" in importer and "再订购点" in importer, (
        "物料导入的表头匹配丢失了「再订货点/再订购点」兼容别名，"
        "会导致历史模板导入失败"
    )


def test_t3_two_material_templates_agree():
    """T3：两份物料导入模板的表头必须逐列一致。"""
    from openpyxl import load_workbook

    with app_module.app.app_context():
        db.drop_all()
        db.create_all()
        from werkzeug.security import generate_password_hash
        db.session.add(User(
            username="admin",
            password_hash=generate_password_hash("admin"),
            role="admin",
            must_change_password=False,
        ))
        app_module.set_system_setting("inventory_alert_enabled", "1")
        db.session.commit()

    client = app_module.app.test_client()
    client.post(
        "/login",
        data={"username": "admin", "password": "admin"},
        content_type="application/x-www-form-urlencoded",
    )

    urls = ["/material/download_template", "/export/template/material"]
    headers_seen = {}
    for url in urls:
        resp = client.get(url)
        assert resp.status_code == 200, f"{url} -> {resp.status_code}"
        ws = load_workbook(io.BytesIO(resp.data)).active
        headers_seen[url] = [c.value for c in ws[1]]

    assert headers_seen[urls[0]] == headers_seen[urls[1]], (
        "两份物料导入模板表头不一致：\n"
        f"  {urls[0]}: {headers_seen[urls[0]]}\n"
        f"  {urls[1]}: {headers_seen[urls[1]]}\n"
        "两处是重复实现，必须同步。"
    )
    # 并且确实包含统一后的叫法
    assert "最低库存" in headers_seen[urls[0]], "模板缺少「最低库存」列"
    assert "安全库存" in headers_seen[urls[0]], "模板缺少「安全库存」列"


def test_t4_safety_stock_round_trip_and_column_mapping():
    """T4：安全库存填值 → 落 reorder_point；导出再导入值不变。"""
    from openpyxl import load_workbook, Workbook

    with app_module.app.app_context():
        db.drop_all()
        db.create_all()
        from werkzeug.security import generate_password_hash
        if not User.query.filter_by(username="admin").first():
            db.session.add(User(
                username="admin",
                password_hash=generate_password_hash("admin"),
                role="admin",
                must_change_password=False,
            ))
        app_module.set_system_setting("inventory_alert_enabled", "1")
        db.session.commit()

    client = app_module.app.test_client()
    client.post(
        "/login",
        data={"username": "admin", "password": "admin"},
        content_type="application/x-www-form-urlencoded",
    )

    # ① 表单提交：最低库存=7，安全库存=21
    resp = client.post(
        "/material/add",
        data={
            "code": "NAMING-RT-1", "name": "命名往返测试", "price": "3",
            "min_stock": "7", "max_stock": "200",
            "reorder_point": "21", "alert_days": "30",
        },
        content_type="application/x-www-form-urlencoded",
        follow_redirects=False,
    )
    assert resp.status_code in (200, 302), f"/material/add -> {resp.status_code}"

    with app_module.app.app_context():
        m = Material.query.filter_by(code="NAMING-RT-1").first()
        assert m is not None, "物料未落库"
        assert m.min_stock == 7, f"最低库存应为 7，实际 {m.min_stock}"
        assert m.reorder_point == 21, (
            f"安全库存应落到 reorder_point=21，实际 {m.reorder_point}"
            "（若为 7，说明写进了 min_stock，字段错位）"
        )

    # ② 导出 → 再导入，值不得变化
    resp = client.get("/material/export")
    assert resp.status_code == 200, f"/material/export -> {resp.status_code}"
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert "安全库存" in header, f"导出缺少「安全库存」列：{header}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = client.post(
        "/material/import",
        data={"file": (buf, "roundtrip.xlsx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code in (200, 302), f"/material/import -> {resp.status_code}"

    with app_module.app.app_context():
        m2 = Material.query.filter_by(code="NAMING-RT-1").first()
        assert m2 is not None, "回导后物料丢失"
        assert m2.min_stock == 7, f"回导后最低库存变了：{m2.min_stock}"
        assert m2.reorder_point == 21, f"回导后安全库存变了：{m2.reorder_point}"


def test_t5_naming_convention_is_documented():
    """T5：命名映射必须成文在 Material 模型 docstring 里。"""
    src = (APP_DIR / "models" / "master_data.py").read_text(encoding="utf-8")
    assert "库存阈值命名约定" in src, "Material 模型缺少「库存阈值命名约定」说明段"
    # 三个关键映射都要写到
    for token in ("min_stock", "reorder_point", "safety_stock"):
        assert token in src, f"命名约定说明缺少 {token}"
    assert "安全库存" in src, "命名约定说明未点明 reorder_point 对外叫「安全库存」"


if __name__ == "__main__":
    for name, fn in sorted(list(globals().items())):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"PASS {name}")
    print("ALL PASSED")
